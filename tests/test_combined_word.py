from pathlib import Path
import tempfile
import unittest

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH

from media_pipeline.combined_word import export, qa_word


class CombinedWordTests(unittest.TestCase):
    def test_three_sections_and_file_order_are_exported(self):
        inventory = {
            "videos": [{"source_name": "视频.mp4", "sha256": "a" * 64}],
            "courseware": [{"source_name": "课件.pdf", "sha256": "b" * 64,
                            "extension": ".pdf"}],
        }
        corpus = {"files": [
            {"source_name": "视频.mp4", "source_sha256": "a" * 64, "kind": "video",
             "sections": {"ASR 结果": ["语音"], "OCR 结果": ["画面"],
                          "GPT 融合校对结果": ["融合"]}},
            {"source_name": "课件.pdf", "source_sha256": "b" * 64, "kind": "courseware",
             "sections": {"ASR 结果": ["不适用：该文件不含音轨。"], "OCR 结果": ["课件"],
                          "GPT 融合校对结果": ["融合课件"]}},
        ]}
        with tempfile.TemporaryDirectory() as directory:
            target = Path(directory) / "result.docx"
            export(corpus, inventory, target)
            qa = qa_word(target, corpus)
            self.assertEqual(qa["file_heading_count"], 2)
            self.assertEqual(qa["section_heading_count"], 6)
            self.assertTrue(target.is_file())
            document = Document(target)
            self.assertIn("视频.mp4", [p.text for p in document.paragraphs])

    def test_multiline_ocr_is_left_aligned_to_avoid_stretched_glyphs(self):
        inventory = {
            "videos": [],
            "courseware": [{"source_name": "课件.pdf", "sha256": "b" * 64,
                            "extension": ".pdf"}],
        }
        corpus = {"files": [
            {"source_name": "课件.pdf", "source_sha256": "b" * 64,
             "kind": "courseware",
             "sections": {"ASR 结果": ["不适用：该文件不含音轨。"],
                          "OCR 结果": ["第一行\n第二行"],
                          "GPT 融合校对结果": ["融合课件"]}},
        ]}
        with tempfile.TemporaryDirectory() as directory:
            target = Path(directory) / "result.docx"
            export(corpus, inventory, target)
            document = Document(target)
            ocr = next(p for p in document.paragraphs if p.text == "第一行\n第二行")
            self.assertEqual(ocr.alignment, WD_ALIGN_PARAGRAPH.LEFT)
            self.assertEqual(ocr.paragraph_format.first_line_indent.cm, 0)


if __name__ == "__main__":
    unittest.main()
