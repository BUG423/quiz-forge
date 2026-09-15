"""Offline report checks: exact deduplication, evidence, and valid Word content."""
from __future__ import annotations

import copy
import sys
import tempfile
import unittest
from pathlib import Path

from docx import Document
from docx.oxml.ns import qn

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
from media_pipeline.report import (
    _change_excerpt, _correction_rows, _extract_persistent_texts,
    _ocr_groups, _partition_persistent_corrections, _ranges, build_report,
)


def frame(index: int, text: str, timestamp: float | None = None) -> dict:
    timestamp = float(index) if timestamp is None else timestamp
    return {"index": index, "timestamp_seconds": timestamp,
            "actual_timestamp_seconds": timestamp + 0.02, "image_path": "cache/frame.png",
            "lines": [{"text": text, "confidence": 0.91, "box": []}] if text else []}


class ReportTests(unittest.TestCase):
    def setUp(self) -> None:
        self.manifest = {"source_path": "data/测试视频.mp4", "source_name": "测试视频.mp4",
                         "duration_seconds": 5.0, "sha256": "a" * 64}
        self.asr = {"model": "mimo-v2.5-asr", "segments": [
            {"index": 0, "start_seconds": 0, "end_seconds": 5, "text": "二号踏施工。"}]}
        self.ocr = {"engine": "RapidOCR", "sample_interval_seconds": 1.0,
                    "expected_frame_count": 5,
                    "frames": [frame(0, "二号塔"), frame(1, "二号塔"), frame(2, ""),
                               frame(3, "二号踏"), frame(4, "二号塔")]}
        self.review = {"reviewer": "测试复核", "reviewed_at": "2026-09-09",
                       "checked_frames": [0, 3], "notes": ["查看了两张画面。"], "unresolved": ["背景小字待核实。"],
                       "asr_corrections": [{"segment_index": 0, "old": "二号踏施工。", "new": "二号塔施工。",
                                            "reason": "同音误字", "evidence": "00:00:00.000 画面标题为二号塔"}],
                       "ocr_corrections": [{"frame_index": 3, "line_index": 0, "old": "二号踏", "new": "二号塔",
                                            "reason": "字形误识别", "evidence": "查看 frame index 3 的原始画面"}]}

    def test_only_adjacent_identical_full_text_is_grouped(self) -> None:
        groups = _ocr_groups(self.ocr)
        self.assertEqual([group["count"] for group in groups], [2, 1, 1])
        self.assertEqual((groups[0]["start"], groups[0]["end"]), (0, 1))
        # A blank frame or nonadjacent timestamp breaks repetition.
        gapped = {"frames": [frame(0, "相同"), frame(1, "相同", 3), frame(2, "相 同", 4)]}
        self.assertEqual(len(_ocr_groups(gapped)), 3)

    def test_compact_frame_ranges_and_long_asr_changes(self) -> None:
        self.assertEqual(_ranges(list(range(272))), "0—271")
        self.assertEqual(_ranges([5, 0, 1, 4, 4, 8]), "0—1、4—5、8")
        old = "塔身施工安全可靠。" * 50 + "二号踏" + "施工按照计划进行。" * 50
        new = old.replace("二号踏", "二号塔")
        old_excerpt, new_excerpt = _change_excerpt(old, new)
        self.assertIn("二号踏", old_excerpt)
        self.assertIn("二号塔", new_excerpt)
        self.assertLess(len(old_excerpt), 30)
        self.assertLess(len(new_excerpt), 30)

    def test_ocr_correction_rows_preserve_discrete_ranges_in_one_row(self) -> None:
        changes = []
        for index in [0, 1, 3, 4]:
            changes.append({"kind": "OCR", "frame_index": index, "line_index": 0,
                            "old": "误字", "new": "正字", "reason": "形近", "evidence": "同一标题画面",
                            "location": f"帧 {index}"})
        original = copy.deepcopy(changes)
        rows = _correction_rows(changes, self.ocr)
        self.assertEqual(len(rows), 1)
        self.assertIn("行 0：帧 0—1、3—4", rows[0][0])
        self.assertIn("00:00:00.000—00:00:01.000、00:00:03.000—00:00:04.000", rows[0][0])
        self.assertEqual(changes, original)
        changes[-1]["evidence"] = "另一依据"
        self.assertEqual(len(_correction_rows(changes, self.ocr)), 2)

    def test_explicit_persistent_text_never_deduplicates_ordinary_reappearance(self) -> None:
        ocr = {"frames": [frame(0, "固定水印\n正文甲"), frame(1, "固定水印\n正文甲"),
                          frame(2, "正文乙"), frame(3, "固定水印\n正文甲"),
                          frame(4, "固定水印的其他文字")]}
        original = copy.deepcopy(ocr)
        automatic, definitions = _extract_persistent_texts(ocr, {})
        self.assertEqual(automatic, original)
        self.assertEqual(definitions, [])
        reviewed, definitions = _extract_persistent_texts(ocr, {"persistent_texts": [
            {"text": "固定水印", "reason": "固定栏目", "evidence": "逐帧查看"}]})
        self.assertEqual(definitions[0]["frame_indices"], [0, 1, 3])
        groups = _ocr_groups(reviewed)
        self.assertEqual([group["text"] for group in groups], ["正文甲", "正文乙", "正文甲", "固定水印的其他文字"])
        self.assertEqual([group["count"] for group in groups], [2, 1, 1, 1])
        self.assertEqual(ocr, original)

    def test_multiline_persistent_text_requires_complete_exact_match(self) -> None:
        ocr = {"frames": [frame(0, "奋斗有我\n为梦前行\n正文"), frame(1, "奋斗有我\n为梦前型")]}
        result, definitions = _extract_persistent_texts(ocr, {"persistent_texts": [
            {"text": "奋斗有我\n为梦前行", "reason": "水印", "evidence": "画面"}]})
        self.assertEqual(definitions[0]["frame_indices"], [0])
        self.assertEqual(_ocr_groups(result)[0]["text"], "正文")
        self.assertEqual(_ocr_groups(result)[1]["text"], "奋斗有我\n为梦前型")

    def test_only_explicit_duplicate_watermark_deletions_enter_summary(self) -> None:
        definitions = [{"text": "固定水印", "frame_indices": [0, 1]}]
        changes = [
            {"kind": "OCR", "frame_index": 0, "line_index": 0, "old": "误识别", "new": "固定水印"},
            {"kind": "OCR", "frame_index": 0, "line_index": 1, "old": "局部", "new": "", "persistent_text": "固定水印"},
            {"kind": "OCR", "frame_index": 0, "line_index": 2, "old": "姓名重复", "new": ""},
        ]
        ordinary, summarized = _partition_persistent_corrections(changes, definitions)
        self.assertEqual(ordinary, changes[2:])
        self.assertEqual(len(summarized[0]["changes"]), 2)
        changes[1]["frame_index"] = 5
        with self.assertRaisesRegex(ValueError, "没有匹配"):
            _partition_persistent_corrections(changes, definitions)

    def test_ai_supplements_and_fixed_legend_are_traceable(self) -> None:
        self.review["ocr_supplements"] = [{"start_seconds": 2, "end_seconds": 2.5, "text": "画面补录小字",
                                           "reason": "自动 OCR 漏识别", "evidence": "查看 2 秒原图"}]
        self.review["persistent_texts"] = [{"text": "二号塔", "reason": "测试固定文字", "evidence": "对应原始帧"}]
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "test.docx"
            build_report(self.manifest, self.asr, self.ocr, self.review, path)
            report = Document(path)
            text = "\n".join(paragraph.text for paragraph in report.paragraphs)
            self.assertIn("AI直接读图补录", text)
            self.assertIn("不是 RapidOCR 原始识别结果", text)
            self.assertIn("画面补录小字", text)
            self.assertIn("匹配帧数：4；帧 index：0—1、3—4", text)
            self.assertIn("固定标语/水印校对汇总", text)
            self.assertIn("完整逐条原文、改文、理由与画面依据仍保存在任务 review.json", text)
            self.assertIn("除上述集中列示的固定文字外，本次没有其他识别文字", text)

    def test_supplements_validate_text_evidence_and_source_time_bounds(self) -> None:
        good = {"start_seconds": 0, "end_seconds": 1, "text": "补录", "reason": "漏字", "evidence": "原图"}
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "invalid.docx"
            for key, value in [("start_seconds", -1), ("end_seconds", 6), ("start_seconds", 2),
                               ("end_seconds", float("nan")), ("start_seconds", True),
                               ("text", " "), ("reason", ""), ("evidence", None)]:
                with self.subTest(key=key, value=value):
                    self.review["ocr_supplements"] = [{**good, key: value}]
                    with self.assertRaises(ValueError):
                        build_report(self.manifest, self.asr, self.ocr, self.review, path)

    def test_report_corrections_counts_and_input_immutability(self) -> None:
        originals = copy.deepcopy((self.asr, self.ocr, self.review))
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "test.docx"
            self.assertEqual(build_report(self.manifest, self.asr, self.ocr, self.review, path), path)
            report = Document(path)
            text = "\n".join(paragraph.text for paragraph in report.paragraphs)
            tables = "\n".join(cell.text for table in report.tables for row in table.rows for cell in row.cells)
            self.assertEqual(report.paragraphs[0].text, "测试视频.mp4")
            self.assertIn("应采样 5 帧，实际记录 5 帧，其中无文字 1 帧；展示 2 组文字", text)
            self.assertIn("记录已查看画面 2 / 5 帧", text)
            self.assertIn("二号塔施工。", text)
            self.assertIn("附录：原始 ASR 结果", text)
            self.assertIn("二号踏施工。", text)
            self.assertIn("非逐字时间戳", text)
            self.assertIn("不能保证捕捉", text)
            self.assertIn("背景小字待核实", text)
            self.assertIn("字形误识别", tables)
            self.assertIn("00:00:00.000 画面标题为二号塔", tables)
            self.assertIn("ASR / OCR 纠错数\n1 / 1", tables)
            for table in report.tables:
                self.assertIsNotNone(table.rows[0]._tr.find(qn("w:trPr" ) + "/" + qn("w:tblHeader")))
            self.assertTrue(report.sections[0].footer._element.xpath('.//w:fldSimple'))
        self.assertEqual((self.asr, self.ocr, self.review), originals)

    def test_exact_old_text_required(self) -> None:
        self.review["asr_corrections"][0]["old"] = "二号踏"
        with tempfile.TemporaryDirectory() as directory, self.assertRaisesRegex(ValueError, "原文不匹配"):
            build_report(self.manifest, self.asr, self.ocr, self.review, Path(directory) / "bad.docx")

    def test_invalid_and_duplicate_corrections_are_rejected(self) -> None:
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "bad.docx"
            review = copy.deepcopy(self.review)
            review["ocr_corrections"][0]["line_index"] = -1
            with self.assertRaisesRegex(ValueError, "非负整数"):
                build_report(self.manifest, self.asr, self.ocr, review, path)
            review = copy.deepcopy(self.review)
            review["ocr_corrections"].append(copy.deepcopy(review["ocr_corrections"][0]))
            with self.assertRaisesRegex(ValueError, "重复纠错"):
                build_report(self.manifest, self.asr, self.ocr, review, path)
            review = copy.deepcopy(self.review)
            review["asr_corrections"][0]["evidence"] = ""
            with self.assertRaisesRegex(ValueError, "evidence"):
                build_report(self.manifest, self.asr, self.ocr, review, path)
            review = copy.deepcopy(self.review)
            review["checked_frames"] = [50]
            with self.assertRaisesRegex(ValueError, "checked_frames"):
                build_report(self.manifest, self.asr, self.ocr, review, path)

    def test_no_claim_of_full_review_or_false_silence(self) -> None:
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "test.docx"
            build_report(self.manifest, {"segments": []}, self.ocr, {}, path)
            text = "\n".join(paragraph.text for paragraph in Document(path).paragraphs)
            self.assertIn("记录已查看画面 0 / 5 帧", text)
            self.assertIn("不能据此判定视频无语音", text)
            self.assertNotIn("附录：原始 ASR 结果", text)


if __name__ == "__main__":
    unittest.main()
