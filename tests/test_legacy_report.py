from __future__ import annotations

import tempfile
from pathlib import Path
import unittest
from unittest.mock import patch

from docx import Document

from media_pipeline.legacy_report import extract


class LegacyReportTests(unittest.TestCase):
    def test_extracts_sections_without_frame_metadata(self):
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / "old.docx"
            document = Document()
            for index in range(1, 59):
                document.add_heading(f"{index:02d}. 文件{index}.mp4", 1)
                document.add_heading("OCR 结果", 2)
                document.add_heading("画面 00:00:01", 3)
                document.add_paragraph(f"画面文字{index}")
                document.add_heading("ASR 结果", 2)
                document.add_paragraph(f"语音文字{index}")
            document.save(path)
            result = extract(path)
            self.assertEqual(result["file_count"], 58)
            self.assertEqual(result["files"][0]["source_name"], "文件1.mp4")
            self.assertEqual(result["files"][0]["ocr_blocks"], [
                {"label": "画面 00:00:01", "text": "画面文字1"}])
            self.assertEqual(result["files"][0]["asr_paragraphs"], ["语音文字1"])


if __name__ == "__main__":
    unittest.main()
