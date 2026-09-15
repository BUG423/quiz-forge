from __future__ import annotations

from copy import deepcopy
import json
from pathlib import Path
import tempfile
import unittest

from docx import Document
from docx.shared import RGBColor

from scripts import exam_points_pipeline as ep


class ExamPointsPipelineTests(unittest.TestCase):
    def setUp(self) -> None:
        self.temporary = tempfile.TemporaryDirectory()
        self.root = Path(self.temporary.name)
        self.manifest_path = self.root / "asset_manifest.json"
        self.evidence_path = self.root / "evidence_corpus.json"
        self.task_dir = self.root / "tasks"
        self.result_dir = self.root / "results"
        self.final_path = self.root / "exam_points_final.json"
        self.word_path = self.root / "exam_points.docx"
        self.count = 4
        self.manifest = self._manifest()
        self.evidence = self._evidence(self.manifest)
        self._write(self.manifest_path, self.manifest)
        self._write(self.evidence_path, self.evidence)

    def tearDown(self) -> None:
        self.temporary.cleanup()

    @staticmethod
    def _write(path: Path, value: dict) -> None:
        path.parent.mkdir(parents=True, exist_ok=True)
        path.write_text(json.dumps(value, ensure_ascii=False), encoding="utf-8")

    def _manifest(self) -> dict:
        assets = []
        for order in range(1, self.count + 1):
            digest = f"{order:064x}"
            assets.append({
                "asset_id": f"2026/video/{digest}",
                "source_sha256": digest,
                "source_name": f"课程{order}.mp4",
                "source_path": f"data/2026/课程{order}.mp4",
                "year": 2026,
                "kind": "video",
                "catalog_position": order,
                "catalog_suborder": 1,
                "catalog_anchor_position": order,
                "coverage_positions": [order],
            })
        return {
            "asset_count": self.count,
            "catalog": {
                "entry_count": self.count,
                "entries": [f"目录原文{order}" for order in range(1, self.count + 1)],
            },
            "assets": assets,
        }

    @staticmethod
    def _evidence(manifest: dict) -> dict:
        files = []
        for order, asset in enumerate(manifest["assets"], 1):
            raw = {
                "asset_id": asset["asset_id"],
                "source_sha256": asset["source_sha256"],
                "source_path": asset["source_path"],
                "source_name": asset["source_name"],
                "year": asset["year"],
                "kind": asset["kind"],
                "catalog_position": asset["catalog_position"],
                "catalog_suborder": asset["catalog_suborder"],
                "asset_manifest_record": asset,
                "asr_result": {
                    "status": "complete",
                    "provider": "Xiaomi MiMo",
                    "model": "mimo-v2.5-asr",
                    "result": {"segments": [{"index": 0, "text": f"语音证据{order}"}]},
                },
                "ocr_result": {
                    "status": "complete",
                    "frames": [{"frame_index": 0, "lines": [{"text": f"画面证据{order}"}]}],
                },
            }
            files.append({
                "asset_id": asset["asset_id"],
                "source_sha256": asset["source_sha256"],
                "evidence_sha256": ep.canonical_sha256(raw),
                "sections": {
                    "ASR 结果": [f"语音证据{order}"],
                    "OCR 结果": [f"画面证据{order}"],
                },
                "raw_evidence": raw,
            })
        return {
            "schema_version": "first-principles-evidence/v1",
            "status": "complete",
            "artifact_role": "authoritative_evidence",
            "asset_count": len(files),
            "file_count": len(files),
            "blockers": [],
            "content_policy": {
                "generated_prose_fields": False,
                "asr_source": "verified Xiaomi MiMo only",
                "local_asr_substitution": False,
            },
            "files": files,
        }

    def _prepare(self) -> dict:
        return ep.prepare_tasks(
            manifest_path=self.manifest_path,
            evidence_path=self.evidence_path,
            output_dir=self.task_dir,
            expected_assets=self.count,
            expected_catalog=self.count,
            batch_count=2,
        )["index"]

    def _result(self, task: dict, *, suffix: str = "") -> dict:
        order = task["order"]
        return {
            "order": order,
            "batch_id": task["batch_id"],
            "catalog_position": task["catalog_position"],
            "catalog_suborder": task["catalog_suborder"],
            "catalog_entry": task["catalog_entry"],
            "coverage_positions": task["coverage_positions"],
            "asset_id": task["asset_id"],
            "source_name": task["source_name"],
            "source_sha256": task["source_sha256"],
            "evidence_sha256": task["evidence_sha256"],
            "task_sha256": task["task_sha256"],
            "model": "gpt-test-model",
            "source_limitations": [],
            "dimension_review": {
                key: "reviewed" for key in ep.DIMENSION_KEYS
            },
            "points": [{
                "point_id": f"A{order:03d}-KP001",
                "dimension": ["定义概念", "数字日期比例单位"],
                "statement": f"课程{order}由GPT直接理解后撰写的考点{suffix}。",
                "question_angles": ["判断题", "单选题"],
                "answer_elements": [f"语音证据{order}", f"画面证据{order}"],
                "traps": ["不得调换主体"],
                "evidence_refs": [
                    {
                        "modality": "ASR",
                        "locator": "/raw_evidence/asr_result/result/segments/0",
                        "support": "direct",
                    },
                    {
                        "modality": "OCR",
                        "locator": "/raw_evidence/ocr_result/frames/0/lines/0",
                        "support": "corroborating",
                    },
                ],
                "evidence_status": "ASR+OCR互证",
            }],
            "asset_review": {
                "full_evidence_read": True,
                "file_level_final_review": True,
                "numbers_rechecked": True,
                "names_rechecked": True,
                "conditions_negations_rechecked": True,
                "process_order_rechecked": True,
                "ui_fields_rechecked": True,
                "unresolved_conflicts": [],
            },
        }

    def _write_results(self, index: dict, count: int | None = None) -> None:
        self.result_dir.mkdir(parents=True, exist_ok=True)
        for task in index["tasks"][:count]:
            self._write(
                self.result_dir / f"result-{task['order']:04d}.json",
                self._result(task),
            )

    def test_prepare_binds_verbatim_evidence_and_batches_without_writing_points(self) -> None:
        index = self._prepare()
        self.assertEqual(self.count, index["file_count"])
        self.assertEqual(2, index["batch_count"])
        self.assertEqual(self.count, sum(item["task_count"] for item in index["batches"]))
        for task_index, evidence_item in zip(index["tasks"], self.evidence["files"]):
            task = json.loads((self.task_dir / task_index["task_file"]).read_text())
            self.assertEqual(set(ep.TASK_FILE_KEYS), set(task))
            self.assertEqual(evidence_item["raw_evidence"]["asr_result"], task["ASR 结果"])
            self.assertEqual(evidence_item["raw_evidence"]["ocr_result"], task["OCR 结果"])
            self.assertFalse(task["task_constraints"]["automatic_point_generation_allowed"])
            self.assertFalse(task["task_constraints"]["old_gpt_fusion_allowed"])
            self.assertNotIn("points", task)
            self.assertEqual(task_index["task_sha256"], ep.canonical_sha256(task))

    def test_incomplete_debug_or_generated_evidence_is_refused_before_task_write(self) -> None:
        incomplete_path = self.root / "evidence_corpus.incomplete.json"
        self._write(incomplete_path, self.evidence)
        with self.assertRaisesRegex(ep.ExamPointsError, "incomplete"):
            ep.prepare_tasks(
                manifest_path=self.manifest_path, evidence_path=incomplete_path,
                output_dir=self.root / "refused-1", expected_assets=self.count,
                expected_catalog=self.count,
            )
        self.assertFalse((self.root / "refused-1").exists())

        blocked = deepcopy(self.evidence)
        blocked["status"] = "incomplete"
        blocked["artifact_role"] = "debug_only_not_for_final_fusion"
        blocked["blockers"] = [{"code": "video_scene_ocr_incomplete"}]
        blocked_path = self.root / "blocked.json"
        self._write(blocked_path, blocked)
        with self.assertRaises(ep.ExamPointsError):
            ep.prepare_tasks(
                manifest_path=self.manifest_path, evidence_path=blocked_path,
                output_dir=self.root / "refused-2", expected_assets=self.count,
                expected_catalog=self.count,
            )

        generated = deepcopy(self.evidence)
        generated["files"][0]["raw_evidence"]["gpt_final"] = "旧稿"
        generated_path = self.root / "generated.json"
        self._write(generated_path, generated)
        with self.assertRaisesRegex(ep.ExamPointsError, "旧GPT"):
            ep.prepare_tasks(
                manifest_path=self.manifest_path, evidence_path=generated_path,
                output_dir=self.root / "refused-3", expected_assets=self.count,
                expected_catalog=self.count,
            )

    def test_merge_preserves_model_points_and_progress_counts_only_valid_results(self) -> None:
        index = self._prepare()
        self._write_results(index, count=2)
        progress = ep.build_progress(
            task_dir=self.task_dir, result_dir=self.result_dir,
            evidence_path=self.evidence_path, output_json=self.root / "progress.json",
            output_markdown=self.root / "progress.md", expected_assets=self.count,
        )
        self.assertEqual(2, progress["qa_passed_count"])
        self.assertEqual(2, progress["point_count"])
        self.assertEqual("[##########----------] 2/4", progress["progress_bar"])
        self.assertEqual(2, sum(item["qa_passed_count"] for item in progress["batches"]))
        self.assertEqual(["QA通过", "QA通过", "已绑定", "已绑定"],
                         [item["status"] for item in progress["assets"]])

        self._write_results(index)
        nested = self.result_dir / index["tasks"][0]["batch_id"]
        nested.mkdir()
        first_path = nested / "result-0001.json"
        (self.result_dir / "result-0001.json").rename(first_path)
        first = json.loads(first_path.read_text())
        first["points"][0]["statement"] = "这是逐字保留、不能由程序改写的考点。"
        self._write(first_path, first)
        final = ep.merge_results(
            task_dir=self.task_dir, result_dir=self.result_dir,
            evidence_path=self.evidence_path, output_path=self.final_path,
            expected_assets=self.count,
        )
        self.assertEqual("这是逐字保留、不能由程序改写的考点。",
                         final["assets"][0]["points"][0]["statement"])
        self.assertFalse(final["basis"]["automatic_point_generation_used"])
        self.assertFalse(final["basis"]["old_gpt_fusion_used"])
        self.assertEqual(self.count, final["point_count"])

    def test_task_cannot_be_rehashed_after_substituting_authoritative_asr_or_ocr(self) -> None:
        index = self._prepare()
        self._write_results(index)
        first_index = index["tasks"][0]
        task_path = self.task_dir / first_index["task_file"]
        task = json.loads(task_path.read_text())
        task["ASR 结果"]["result"]["segments"][0]["text"] = "伪造语音证据"
        self._write(task_path, task)
        index["tasks"][0]["task_sha256"] = ep.canonical_sha256(task)
        self._write(self.task_dir / "index.json", index)
        result = json.loads((self.result_dir / "result-0001.json").read_text())
        result["task_sha256"] = index["tasks"][0]["task_sha256"]
        self._write(self.result_dir / "result-0001.json", result)
        with self.assertRaisesRegex(ep.ExamPointsError, "ASR与正式证据不逐字一致"):
            ep.merge_results(
                task_dir=self.task_dir, result_dir=self.result_dir,
                evidence_path=self.evidence_path, output_path=None,
                expected_assets=self.count,
            )

    def test_strict_result_schema_identity_pointer_and_file_review_are_enforced(self) -> None:
        index = self._prepare()
        self._write_results(index)
        original = json.loads((self.result_dir / "result-0001.json").read_text())
        variants = []
        wrong_sha = deepcopy(original)
        wrong_sha["source_sha256"] = "f" * 64
        variants.append(wrong_sha)
        wrong_task = deepcopy(original)
        wrong_task["task_sha256"] = "f" * 64
        variants.append(wrong_task)
        wrong_pointer = deepcopy(original)
        wrong_pointer["points"][0]["evidence_refs"][0]["locator"] = (
            "/raw_evidence/asr_result/result/segments/999"
        )
        variants.append(wrong_pointer)
        broad_pointer = deepcopy(original)
        broad_pointer["points"][0]["evidence_refs"][0]["locator"] = (
            "/raw_evidence/asr_result"
        )
        variants.append(broad_pointer)
        wrong_modality = deepcopy(original)
        wrong_modality["points"][0]["evidence_status"] = "ASR单证"
        variants.append(wrong_modality)
        no_final_review = deepcopy(original)
        no_final_review["asset_review"]["file_level_final_review"] = False
        variants.append(no_final_review)
        extra_field = deepcopy(original)
        extra_field["summary"] = "代码生成的摘要"
        variants.append(extra_field)
        for number, variant in enumerate(variants):
            self._write(self.result_dir / "result-0001.json", variant)
            with self.subTest(number=number), self.assertRaises(ep.ExamPointsError):
                ep.merge_results(
                    task_dir=self.task_dir, result_dir=self.result_dir,
                    evidence_path=self.evidence_path, output_path=None,
                    expected_assets=self.count,
                )
        self._write(self.result_dir / "result-0001.json", original)

        index_document = json.loads((self.task_dir / "index.json").read_text())
        index_document["batches"][0]["orders"] = []
        self._write(self.task_dir / "index.json", index_document)
        with self.assertRaisesRegex(ep.ExamPointsError, "批次汇总"):
            ep.load_task_index(self.task_dir, expected_assets=self.count)

    def test_independent_word_and_read_only_qa_are_verbatim(self) -> None:
        index = self._prepare()
        self._write_results(index)
        ep.merge_results(
            task_dir=self.task_dir, result_dir=self.result_dir,
            evidence_path=self.evidence_path, output_path=self.final_path,
            expected_assets=self.count,
        )
        with self.assertRaisesRegex(ep.ExamPointsError, "拒绝写入主"):
            ep.export_word(
                manifest_path=self.manifest_path, task_dir=self.task_dir,
                evidence_path=self.evidence_path, final_path=self.final_path,
                output_path=self.root / "2026年YW第一课全量_ASR-OCR-GPT融合校对.docx",
                expected_assets=self.count, expected_catalog=self.count,
            )
        ep.export_word(
            manifest_path=self.manifest_path, task_dir=self.task_dir,
            evidence_path=self.evidence_path, final_path=self.final_path,
            output_path=self.word_path, expected_assets=self.count,
            expected_catalog=self.count,
        )
        qa = ep.qa_delivery(
            manifest_path=self.manifest_path, task_dir=self.task_dir,
            evidence_path=self.evidence_path, final_path=self.final_path,
            word_path=self.word_path, report_path=self.root / "qa.json",
            expected_assets=self.count, expected_catalog=self.count,
        )
        self.assertTrue(qa["passed"])
        self.assertEqual(self.count, qa["point_count"])
        document = Document(self.word_path)
        self.assertIn(
            "考点陈述：课程1由GPT直接理解后撰写的考点。",
            [paragraph.text for paragraph in document.paragraphs],
        )
        self.assertFalse(any(
            run.text and run.font.color.rgb is not None
            and str(run.font.color.rgb) == ep.RED
            for paragraph in document.paragraphs for run in paragraph.runs
        ))

        target = next(
            paragraph for paragraph in document.paragraphs
            if paragraph.text == "考点陈述：课程1由GPT直接理解后撰写的考点。"
        )
        target.text += "篡改"
        for run in target.runs:
            run.font.color.rgb = RGBColor.from_string(ep.DARK)
        document.save(self.word_path)
        with self.assertRaisesRegex(ep.ExamPointsError, "不逐字一致"):
            ep.qa_delivery(
                manifest_path=self.manifest_path, task_dir=self.task_dir,
                evidence_path=self.evidence_path, final_path=self.final_path,
                word_path=self.word_path, report_path=None,
                expected_assets=self.count, expected_catalog=self.count,
            )

    def test_word_only_display_name_is_identity_bound_and_qa_checked(self) -> None:
        index = self._prepare()
        self._write_results(index)
        ep.merge_results(
            task_dir=self.task_dir, result_dir=self.result_dir,
            evidence_path=self.evidence_path, output_path=self.final_path,
            expected_assets=self.count,
        )
        first = self.manifest["assets"][0]
        display_path = self.root / "word_display_names.json"
        self._write(display_path, {
            "schema_version": ep.DISPLAY_NAMES_SCHEMA_VERSION,
            "entries": [{
                "asset_id": first["asset_id"],
                "source_sha256": first["source_sha256"],
                "source_name": first["source_name"],
                "catalog_position": first["catalog_position"],
                "catalog_suborder": first["catalog_suborder"],
                "display_title": "国家安全正式课程名称",
                "basis": "课程画面标题与语音开场共同确认。",
            }],
        })
        ep.export_word(
            manifest_path=self.manifest_path, task_dir=self.task_dir,
            evidence_path=self.evidence_path, final_path=self.final_path,
            output_path=self.word_path, display_names_path=display_path,
            expected_assets=self.count, expected_catalog=self.count,
        )
        qa = ep.qa_delivery(
            manifest_path=self.manifest_path, task_dir=self.task_dir,
            evidence_path=self.evidence_path, final_path=self.final_path,
            word_path=self.word_path, report_path=None,
            display_names_path=display_path,
            expected_assets=self.count, expected_catalog=self.count,
        )
        self.assertTrue(qa["passed"])
        document = Document(self.word_path)
        texts = [paragraph.text for paragraph in document.paragraphs]
        self.assertIn("A001 [国家安全正式课程名称]", texts)
        self.assertIn("源文件名：课程1.mp4（仅保留用于证据追溯）", texts)
        self.assertNotIn("A001 [课程1.mp4]", texts)


if __name__ == "__main__":
    unittest.main()
