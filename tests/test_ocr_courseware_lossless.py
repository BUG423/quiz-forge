"""Lossless courseware OCR contracts; all inference is fake and offline."""

from __future__ import annotations

import io
import sys
import tempfile
import unittest
from pathlib import Path
from types import SimpleNamespace
from unittest.mock import patch
from zipfile import ZIP_DEFLATED, ZipFile

import numpy as np
import pymupdf
from openpyxl import Workbook
from pptx import Presentation
from pptx.chart.data import ChartData
from pptx.enum.chart import XL_CHART_TYPE
from pptx.util import Inches

from scripts import ocr_courseware_lossless as lossless


class FakeEngine:
    def __call__(self, image, *, text_score):
        assert text_score == 0.0
        return SimpleNamespace(
            txts=("云南电纲", "A", "A"),
            scores=(0.99, 0.01, 0.01),
            boxes=np.array([
                [[1, 1], [20, 1], [20, 10], [1, 10]],
                [[2, 2], [8, 2], [8, 9], [2, 9]],
                [[2, 2], [8, 2], [8, 9], [2, 9]],
            ]),
        )


def make_pdf(path: Path, pages: int) -> None:
    document = pymupdf.open()
    for index in range(pages):
        page = document.new_page(width=320, height=180)
        page.insert_text((30, 80), f"page {index + 1}")
    document.save(path)
    document.close()


def add_external_relationship(path: Path, target: str) -> None:
    with ZipFile(path) as source:
        content = {name: source.read(name) for name in source.namelist()}
    rel = "ppt/slides/_rels/slide1.xml.rels"
    root = lossless.ET.fromstring(content[rel])
    lossless.ET.SubElement(root, f"{{{lossless.REL_NS}}}Relationship", {
        "Id": "rIdMissingFixture",
        "Type": f"{lossless.OFFICE_REL_NS}/hyperlink",
        "Target": target,
        "TargetMode": "External",
    })
    content[rel] = lossless.ET.tostring(root, encoding="utf-8", xml_declaration=True)
    with ZipFile(path, "w", ZIP_DEFLATED) as destination:
        for name, payload in content.items():
            destination.writestr(name, payload)


def make_pptx(path: Path) -> None:
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    box = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(4), Inches(1))
    box.text = "A"
    data = ChartData()
    data.categories = ["一", "二"]
    data.add_series("重复", (1, 2))
    slide.shapes.add_chart(
        XL_CHART_TYPE.COLUMN_CLUSTERED, Inches(1), Inches(2), Inches(5), Inches(3), data,
    )
    second = presentation.slides.add_slide(presentation.slide_layouts[6])
    second.shapes.add_textbox(Inches(1), Inches(1), Inches(4), Inches(1)).text = "第二页"
    presentation.save(path)


class LosslessCoursewareTests(unittest.TestCase):
    def test_raw_ocr_keeps_corrections_single_char_low_score_and_duplicates(self):
        lines = lossless.recognize_raw(FakeEngine(), np.zeros((20, 30, 3), np.uint8))
        self.assertEqual([item["text"] for item in lines], ["云南电纲", "A", "A"])
        self.assertEqual([item["confidence"] for item in lines], [0.99, 0.01, 0.01])
        self.assertEqual([item["index"] for item in lines], [0, 1, 2])

    def test_native_text_and_embedded_workbook_are_mapped_to_source_slide(self):
        with tempfile.TemporaryDirectory() as temporary:
            source = Path(temporary) / "fixture.pptx"
            make_pptx(source)
            inventory = lossless.pptx_inventory(source)
            self.assertEqual(len(inventory["slide_parts"]), 2)
            self.assertIn("A", [item["text"] for item in inventory["native_by_page"][1]])
            workbooks = [item for item in inventory["objects"] if item["extension"] == ".xlsx"]
            self.assertEqual(len(workbooks), 1)
            self.assertEqual(workbooks[0]["source_pages"], [1])

    def test_xlsx_native_success_and_render_failure_is_an_audited_risk(self):
        with tempfile.TemporaryDirectory() as temporary:
            root = Path(temporary)
            source = root / "fixture.pptx"
            make_pptx(source)
            inventory = lossless.pptx_inventory(source)
            with patch.object(
                lossless, "run_soffice",
                side_effect=lossless.LosslessOCRError("LibreOffice fixture failure"),
            ):
                objects, errors, risks = lossless.process_embedded_objects(
                    source, inventory, FakeEngine(), root / "work",
                    Path("/not/used/soffice"), 1.0,
                )
            workbook = next(item for item in objects if item["extension"] == ".xlsx")
            self.assertEqual(errors, [])
            self.assertEqual(workbook["status"], lossless.XLSX_RENDER_DEGRADED_STATUS)
            self.assertGreater(sum(len(sheet["cells"]) for sheet in workbook["native_content"]), 0)
            self.assertEqual(workbook["source_pages"], [1])
            self.assertEqual(workbook["raw_ocr_frames"], [])
            matching = [risk for risk in risks
                        if risk["code"] == lossless.XLSX_RENDER_DEGRADED_RISK]
            self.assertEqual(len(matching), 1)
            self.assertEqual(matching[0]["coverage"], "source_full_page_ocr")
            self.assertIn("LibreOffice fixture failure", matching[0]["exception_message"])

    def test_xlsx_native_extraction_failure_remains_a_hard_error(self):
        with tempfile.TemporaryDirectory() as temporary:
            root = Path(temporary)
            source = root / "fixture.pptx"
            make_pptx(source)
            inventory = lossless.pptx_inventory(source)
            with patch.object(lossless, "_xlsx_native", side_effect=ValueError("bad workbook")):
                objects, errors, risks = lossless.process_embedded_objects(
                    source, inventory, FakeEngine(), root / "work",
                    Path("/not/used/soffice"), 1.0,
                )
            workbook = next(item for item in objects if item["extension"] == ".xlsx")
            self.assertEqual(workbook["status"], "error")
            self.assertEqual(workbook["native_content"], [])
            self.assertEqual(len(errors), 1)
            self.assertEqual(errors[0]["code"], "embedded_object_processing_failed")
            self.assertEqual(risks, [])

    def test_stale_cache_adds_only_missing_embedded_member_and_binds_inventory(self):
        with tempfile.TemporaryDirectory() as temporary:
            root = Path(temporary)
            source = root / "fixture.pptx"
            make_pptx(source)
            stale = {"embedded_objects": [], "risks": [{
                "code": "embedded_object_page_unmapped", "member": "legacy",
            }]}
            with patch.object(
                lossless, "run_soffice",
                side_effect=lossless.LosslessOCRError("fixture render failure"),
            ):
                refreshed, added = lossless.refresh_cached_embedded_inventory(
                    stale, source, FakeEngine(), root / "work",
                    Path("/not/used/soffice"), 1.0,
                )
            self.assertEqual(added, 1)
            self.assertEqual(len(refreshed["embedded_objects"]), 1)
            self.assertEqual(
                refreshed["embedded_inventory_version"],
                lossless.EMBEDDED_INVENTORY_VERSION,
            )
            self.assertEqual(len(refreshed["embedded_inventory_sha256"]), 64)
            self.assertNotIn(
                "embedded_object_page_unmapped",
                [risk["code"] for risk in refreshed["risks"]],
            )

    def test_zero_byte_source_member_is_preserved_as_explicit_corruption_risk(self):
        with tempfile.TemporaryDirectory() as temporary:
            root = Path(temporary)
            source = root / "empty.pptx"
            with ZipFile(source, "w") as archive:
                archive.writestr("ppt/media/image1.jpeg", b"")
            inventory = {"objects": [{
                "member": "ppt/media/image1.jpeg",
                "member_sha256": lossless.sha256_bytes(b""), "size_bytes": 0,
                "extension": ".jpeg", "source_pages": [1],
                "relationship_types": ["image"],
            }]}
            objects, errors, risks = lossless.process_embedded_objects(
                source, inventory, FakeEngine(), root / "work",
                Path("/not/used/soffice"), 1.0,
            )
            self.assertEqual(errors, [])
            self.assertEqual(objects[0]["status"], "source_member_empty_unreadable")
            self.assertEqual(objects[0]["raw_ocr_frames"], [])
            self.assertEqual([risk["code"] for risk in risks], ["empty_embedded_member"])

    def test_wdp_uses_jpegxr_decoder_and_returns_bgr_pixels(self):
        rgb = np.zeros((3, 4, 3), dtype=np.uint8)
        rgb[:, :, 0] = 255
        fake_codec = SimpleNamespace(jpegxr_decode=lambda payload: rgb)
        with patch.dict(sys.modules, {"imagecodecs": fake_codec}):
            frames = lossless._decode_member_frames(b"fixture", ".wdp")
        self.assertEqual(len(frames), 1)
        self.assertEqual(frames[0].shape, (3, 4, 3))
        self.assertEqual(frames[0][0, 0].tolist(), [0, 0, 255])

    def test_missing_file_relationship_becomes_risk_with_slide(self):
        with tempfile.TemporaryDirectory() as temporary:
            source = Path(temporary) / "fixture.pptx"
            make_pptx(source)
            add_external_relationship(source, "附件/缺失材料.docx")
            inventory = lossless.pptx_inventory(source)
            risks = lossless.external_risks(source, inventory["external_relationships"])
            match = [item for item in risks if item["target"] == "附件/缺失材料.docx"]
            self.assertEqual(len(match), 1)
            self.assertEqual(match[0]["source_page"], 1)

    def test_only_known_final_slide_gap_gets_explicit_fallback(self):
        name = "南网在线“速捷电”品牌系列功能建设介绍.pptx"
        plan = lossless.validate_ppt_render_plan(name, 48, 47)
        self.assertEqual(len(plan), 48)
        self.assertEqual(plan[-1], "native_text_fallback")
        with self.assertRaisesRegex(lossless.LosslessOCRError, "页数"):
            lossless.validate_ppt_render_plan("其它.pptx", 48, 47)

    def test_fixture_ppt_has_equal_source_and_page_records_with_image_hashes(self):
        with tempfile.TemporaryDirectory() as temporary:
            root = Path(temporary)
            source = root / "fixture.pptx"
            make_pptx(source)
            render_root = root / "renders"
            render_dir = render_root / "01"
            render_dir.mkdir(parents=True)
            digest = lossless.sha256_file(source)[:16]
            render = render_dir / f"input-{digest}.pdf"
            make_pdf(render, 2)
            (render_dir / f"input-{digest}.pptx").write_bytes(source.read_bytes())

            # Avoid rendering the embedded chart workbook in this integration
            # check; its page mapping is covered independently above.
            original = lossless.process_embedded_objects
            lossless.process_embedded_objects = lambda *args, **kwargs: ([], [], [])
            try:
                result = lossless.process_file(
                    source, FakeEngine(), root / "work", [render_root],
                    Path("/not/used/soffice"), 1.0, progress=lambda _: None,
                )
            finally:
                lossless.process_embedded_objects = original
            self.assertEqual(result["render"]["source_slide_count"], 2)
            self.assertEqual(result["page_record_count"], 2)
            self.assertTrue(result["processing_complete"])
            for page in result["pages"]:
                evidence = Path(page["evidence_path"])
                self.assertTrue(evidence.is_file())
                self.assertEqual(lossless.sha256_file(evidence), page["image_sha256"])
                self.assertIn("native_lines", page)
                self.assertEqual([line["text"] for line in page["raw_ocr_lines"]],
                                 ["云南电纲", "A", "A"])


if __name__ == "__main__":
    unittest.main()
