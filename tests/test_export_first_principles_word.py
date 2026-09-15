from __future__ import annotations

from copy import deepcopy
import json
from pathlib import Path
import tempfile
import unittest

from docx import Document
from docx.shared import RGBColor

from scripts.export_first_principles_word import (
    ASSET_COUNT,
    BLACK,
    BLUE,
    CATALOG_ENTRY_COUNT,
    ExportValidationError,
    GPT_SECTION,
    OCR_PARAGRAPH_CHARACTERS,
    OCR_SECTION,
    RED,
    export_word,
    load_bundle,
    mechanical_ocr_paragraphs,
    qa_word,
    validate_asset_manifest,
    validate_display_names,
)


class FirstPrinciplesWordTest(unittest.TestCase):
    def setUp(self) -> None:
        self.temporary = tempfile.TemporaryDirectory()
        self.root = Path(self.temporary.name)
        self.manifest_path = self.root / "asset_manifest.json"
        self.evidence_path = self.root / "evidence_corpus.json"
        self.gpt_path = self.root / "gpt_final.json"
        self.word_path = self.root / "delivery.docx"

        assets = []
        evidence_files = []
        gpt_files = []
        for index in range(1, ASSET_COUNT + 1):
            digest = f"{index:064x}"
            asset_id = f"asset-{index:03d}"
            # Duplicate visible names prove that identity is asset_id + SHA, not
            # the human-facing filename.
            source_name = "同名课程.mp4" if index in {1, 2} else f"课程-{index:03d}.mp4"
            catalog_position = min(index, CATALOG_ENTRY_COUNT)
            assets.append({
                "asset_id": asset_id,
                "source_sha256": digest,
                "source_name": source_name,
                "year": 2026 if index % 2 else 2025,
                "kind": "video" if index % 3 else "courseware",
                "catalog_position": catalog_position,
                "catalog_suborder": (
                    index - CATALOG_ENTRY_COUNT + 1
                    if catalog_position == CATALOG_ENTRY_COUNT else 1
                ),
                "coverage_positions": [catalog_position],
            })
            evidence_files.append({
                "asset_id": asset_id,
                "source_sha256": digest,
                "evidence_sha256": f"{index + 1000:064x}",
                "sections": {
                    "ASR 结果": [f"第{index}项语音原文\n第二句。"],
                    "OCR 结果": [f"第{index}页 OCR碎片A", "OCR碎片B\nOCR碎片C"],
                },
            })
            gpt_files.append({
                "asset_id": asset_id,
                "source_sha256": digest,
                "evidence_sha256": f"{index + 1000:064x}",
                "task_sha256": f"{index + 2000:064x}",
                "paragraphs": [f"第{index}项由大模型直接撰写的完整融合正文。"],
            })
        self.manifest = {
            "asset_count": ASSET_COUNT,
            "catalog": {
                "entry_count": CATALOG_ENTRY_COUNT,
                "entries": [
                    f"目录原文-{index:02d}"
                    for index in range(1, CATALOG_ENTRY_COUNT + 1)
                ],
            },
            "assets": assets,
        }
        self.evidence = {"file_count": ASSET_COUNT, "files": evidence_files}
        self.gpt = {"model": "gpt-test-large", "file_count": ASSET_COUNT, "files": gpt_files}
        self._write_inputs()

    def tearDown(self) -> None:
        self.temporary.cleanup()

    def _write_inputs(self) -> None:
        for path, value in (
            (self.manifest_path, self.manifest),
            (self.evidence_path, self.evidence),
            (self.gpt_path, self.gpt),
        ):
            path.write_text(json.dumps(value, ensure_ascii=False), encoding="utf-8")

    def _bundle(self):
        return load_bundle(self.manifest_path, self.evidence_path, self.gpt_path)

    def test_export_has_81_row_directory_91_assets_and_exact_color_boundary(self) -> None:
        bundle = self._bundle()
        self.assertEqual(
            list(range(1, ASSET_COUNT + 1)),
            [asset["order"] for asset in bundle["assets"]],
        )
        export_word(bundle, self.word_path)
        result = qa_word(bundle, self.word_path)
        self.assertTrue(result["passed"])
        self.assertEqual(CATALOG_ENTRY_COUNT, result["catalog_row_count"])
        self.assertEqual(ASSET_COUNT, result["asset_count"])
        self.assertTrue(result["gpt_verbatim"])
        self.assertTrue(result["asr_verbatim"])
        self.assertTrue(result["ocr_verbatim"])
        self.assertEqual(result["ocr_source_fragment_count"], ASSET_COUNT * 2)
        self.assertEqual(result["ocr_word_paragraph_count"], ASSET_COUNT)
        self.assertEqual(result["ocr_fragment_separator"], "single_space")
        self.assertTrue(result["red_not_leaked"])

        document = Document(self.word_path)
        self.assertEqual(CATALOG_ENTRY_COUNT + 1, len(document.tables[0].rows))
        headings = [p for p in document.paragraphs if p.style.name == "Heading 2"]
        self.assertEqual(ASSET_COUNT * 3, len(headings))
        self.assertTrue(all(
            all(str(run.font.color.rgb) == BLUE for run in paragraph.runs if run.text)
            for paragraph in headings
        ))
        gpt_text = self.gpt["files"][0]["paragraphs"][0]
        paragraph = next(p for p in document.paragraphs if p.text == gpt_text)
        self.assertTrue(all(str(run.font.color.rgb) == RED for run in paragraph.runs if run.text))

        expected_ocr = "第1页 OCR碎片A OCR碎片B OCR碎片C"
        self.assertIn(expected_ocr, [p.text for p in document.paragraphs])

    def test_binding_is_strict_and_old_fusion_cannot_be_a_fallback(self) -> None:
        cases = []
        missing = deepcopy(self.gpt)
        missing["files"].pop()
        missing["file_count"] -= 1
        cases.append(missing)

        wrong_sha = deepcopy(self.gpt)
        wrong_sha["files"][0]["source_sha256"] = "f" * 64
        cases.append(wrong_sha)

        wrong_evidence_sha = deepcopy(self.gpt)
        wrong_evidence_sha["files"][0]["evidence_sha256"] = "f" * 64
        cases.append(wrong_evidence_sha)

        old_shape = deepcopy(self.gpt)
        item = old_shape["files"][0]
        item.pop("paragraphs")
        item["sections"] = {GPT_SECTION: ["旧融合稿"]}
        cases.append(old_shape)

        forbidden = deepcopy(self.gpt)
        forbidden["files"][0]["paragraphs"] = ["【ASR底稿】不得进入最终正文。"]
        cases.append(forbidden)

        for paragraph in (
            "以下是完整的语音识别底稿。",
            "以下是完整的OCR识别底稿。",
            "ASR/OCR原文附后。",
            "根据识别结果整理如下。",
        ):
            forbidden_variant = deepcopy(self.gpt)
            forbidden_variant["files"][0]["paragraphs"] = [paragraph]
            cases.append(forbidden_variant)

        for value in cases:
            with self.subTest(case=cases.index(value)):
                self.gpt_path.write_text(json.dumps(value, ensure_ascii=False), encoding="utf-8")
                with self.assertRaises(ExportValidationError):
                    self._bundle()

    def test_read_only_qa_detects_text_tampering_and_red_leakage(self) -> None:
        bundle = self._bundle()
        export_word(bundle, self.word_path)

        document = Document(self.word_path)
        first_gpt = self.gpt["files"][0]["paragraphs"][0]
        paragraph = next(p for p in document.paragraphs if p.text == first_gpt)
        paragraph.text = first_gpt + "被篡改"
        for run in paragraph.runs:
            run.font.color.rgb = RGBColor.from_string(RED)
        document.save(self.word_path)
        with self.assertRaisesRegex(ExportValidationError, "逐字一致"):
            qa_word(bundle, self.word_path)

        export_word(bundle, self.word_path)
        document = Document(self.word_path)
        asr_text = "第1项语音原文 第二句。"
        paragraph = next(p for p in document.paragraphs if p.text == asr_text)
        for run in paragraph.runs:
            run.font.color.rgb = RGBColor.from_string(RED)
        document.save(self.word_path)
        with self.assertRaisesRegex(ExportValidationError, "红色泄漏"):
            qa_word(bundle, self.word_path)

    def test_ocr_partition_is_mechanical_and_lossless(self) -> None:
        source = ["  甲\n乙  ", "重复", "重复", "丙 " * 80, " 丁"]
        paragraphs = mechanical_ocr_paragraphs(source, maximum=100)
        normalized = "甲 乙 重复 重复 " + ("丙 " * 80).strip() + " 丁"
        self.assertEqual(normalized, " ".join(paragraphs))
        self.assertGreater(len(paragraphs), 1)
        self.assertEqual(" ".join(paragraphs).count("重复"), 2)

        # The production default packs at roughly ten thousand characters;
        # it never turns each incoming OCR box into its own Word paragraph.
        self.assertEqual(OCR_PARAGRAPH_CHARACTERS, 10_000)
        packed = mechanical_ocr_paragraphs(["甲" * 4_000, "乙" * 4_000, "丙" * 3_000])
        self.assertEqual([8_001, 3_000], [len(item) for item in packed])
        self.assertEqual("甲" * 4_000 + " " + "乙" * 4_000 + " " + "丙" * 3_000,
                         " ".join(packed))

    def test_read_only_qa_detects_ocr_fragment_stream_tampering(self) -> None:
        bundle = self._bundle()
        export_word(bundle, self.word_path)
        document = Document(self.word_path)
        expected = "第1页 OCR碎片A OCR碎片B OCR碎片C"
        paragraph = next(p for p in document.paragraphs if p.text == expected)
        paragraph.text = paragraph.text.replace("OCR碎片B", "OCR碎片甲", 1)
        for run in paragraph.runs:
            run.font.color.rgb = RGBColor.from_string(BLACK)
        document.save(self.word_path)
        with self.assertRaisesRegex(ExportValidationError, r"OCR\s*结果.*逐字一致"):
            qa_word(bundle, self.word_path)

    def test_manifest_uses_phase_a_schema_and_rejects_array_disorder(self) -> None:
        validated = validate_asset_manifest(self.manifest)
        self.assertEqual(1, validated["assets"][0]["order"])
        self.assertEqual(ASSET_COUNT, validated["assets"][-1]["order"])
        self.assertNotIn("catalog_positions", validated["assets"][0])
        self.assertEqual([1], validated["assets"][0]["coverage_positions"])

        invalid_position = deepcopy(self.manifest)
        invalid_position["assets"][1]["catalog_position"] = 0.5
        with self.assertRaisesRegex(ExportValidationError, "catalog_position"):
            validate_asset_manifest(invalid_position)

        disordered = deepcopy(self.manifest)
        disordered["assets"][0]["catalog_suborder"] = 2
        disordered["assets"][1]["catalog_position"] = 1
        disordered["assets"][1]["catalog_suborder"] = 1
        with self.assertRaisesRegex(ExportValidationError, "数组未按"):
            validate_asset_manifest(disordered)

    def test_word_display_names_are_identity_bound_and_presentation_only(self) -> None:
        # Two assets deliberately share a source filename.  Only asset-001 is
        # retitled, proving this is asset-id/SHA binding rather than a global
        # filename substitution.
        display_path = self.root / "word_display_names.json"
        display = {
            "schema_version": "word-display-names-v1",
            "scope": "presentation only",
            "entries": [{
                "asset_id": "asset-001",
                "source_sha256": f"{1:064x}",
                "source_name": "同名课程.mp4",
                "catalog_position": 1,
                "catalog_suborder": 1,
                "display_title": "国家安全正式课名",
                "basis": "开场自报课名。",
            }],
        }
        display_path.write_text(json.dumps(display, ensure_ascii=False), encoding="utf-8")
        bundle = load_bundle(
            self.manifest_path, self.evidence_path, self.gpt_path, display_path)
        self.assertEqual("国家安全正式课名", bundle["assets"][0]["display_title"])
        self.assertEqual("同名课程.mp4", bundle["assets"][1]["display_title"])
        self.assertEqual(
            [asset["asset_id"] for asset in validate_asset_manifest(self.manifest)["assets"]],
            [asset["asset_id"] for asset in bundle["assets"]],
        )
        self.assertEqual(
            [asset["source_sha256"] for asset in validate_asset_manifest(self.manifest)["assets"]],
            [asset["source_sha256"] for asset in bundle["assets"]],
        )

        export_word(bundle, self.word_path)
        qa_word(bundle, self.word_path)
        document = Document(self.word_path)
        headings = [p.text for p in document.paragraphs if p.style.name == "Heading 1"]
        self.assertIn("1. [2026] 国家安全正式课名", headings)
        self.assertIn("2. [2025] 同名课程.mp4", headings)
        subtitles = [p.text for p in document.paragraphs if p.style.name == "Subtitle"]
        self.assertIn("对应课程目录第 1 条；源文件名：同名课程.mp4", subtitles)
        # The course catalog is copied verbatim even when an asset title changes.
        self.assertEqual("目录原文-01", document.tables[0].rows[1].cells[1].text)

        wrong_sha = deepcopy(display)
        wrong_sha["entries"][0]["source_sha256"] = "f" * 64
        with self.assertRaisesRegex(ExportValidationError, "SHA绑定失败"):
            validate_display_names(wrong_sha, validate_asset_manifest(self.manifest))

        duplicate_title = deepcopy(display)
        duplicate_title["entries"].append({
            **duplicate_title["entries"][0],
            "asset_id": "asset-002",
            "source_sha256": f"{2:064x}",
            "catalog_position": 2,
        })
        with self.assertRaisesRegex(ExportValidationError, "重复正式标题"):
            validate_display_names(duplicate_title, validate_asset_manifest(self.manifest))


if __name__ == "__main__":
    unittest.main()
