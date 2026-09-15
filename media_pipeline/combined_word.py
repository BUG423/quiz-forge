"""Export and validate the single user-facing 2025 Word deliverable."""
from __future__ import annotations

import argparse
from datetime import date
import json
from pathlib import Path
import re
import sys
from typing import Any
from zipfile import ZipFile

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

from .bulk_fusion import SECTION_TITLES, validate_corpus
from .runner import atomic_json, read_json


BLUE = "175A8A"
DARK = "203746"
ORANGE = "D97706"
LIGHT_BLUE = "EAF3F8"
LIGHT_ORANGE = "FFF7E8"
LIGHT_GRAY = "F3F4F6"
GRAY = "667085"
FONT_CN = "Microsoft YaHei"


def _font(style: Any, size: float, *, bold: bool = False, color: str | None = None) -> None:
    style.font.name = FONT_CN
    style.font.size = Pt(size)
    style.font.bold = bold
    if color:
        style.font.color.rgb = RGBColor.from_string(color)
    fonts = style.element.get_or_add_rPr().get_or_add_rFonts()
    fonts.set(qn("w:eastAsia"), FONT_CN)
    fonts.set(qn("w:ascii"), "Aptos")
    fonts.set(qn("w:hAnsi"), "Aptos")


def _shade(cell: Any, color: str) -> None:
    properties = cell._tc.get_or_add_tcPr()
    node = properties.find(qn("w:shd"))
    if node is None:
        node = OxmlElement("w:shd")
        properties.append(node)
    node.set(qn("w:fill"), color)


def _no_split(row: Any) -> None:
    row._tr.get_or_add_trPr().append(OxmlElement("w:cantSplit"))


def _repeat_header(row: Any) -> None:
    node = OxmlElement("w:tblHeader")
    node.set(qn("w:val"), "true")
    row._tr.get_or_add_trPr().append(node)


def setup(document: Document) -> None:
    section = document.sections[0]
    section.page_width, section.page_height = Cm(21), Cm(29.7)
    section.top_margin, section.bottom_margin = Cm(1.75), Cm(1.65)
    section.left_margin, section.right_margin = Cm(2.05), Cm(1.95)
    section.header_distance, section.footer_distance = Cm(0.75), Cm(0.7)

    normal = document.styles["Normal"]
    _font(normal, 10.5)
    normal.paragraph_format.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE
    normal.paragraph_format.space_after = Pt(5)
    normal.paragraph_format.widow_control = True
    for name, size, color, before, after in (
        ("Title", 27, BLUE, 0, 18),
        ("Subtitle", 12, GRAY, 0, 8),
        ("Heading 1", 16, DARK, 0, 10),
        ("Heading 2", 12.5, ORANGE, 10, 6),
    ):
        style = document.styles[name]
        _font(style, size, bold=True if name != "Subtitle" else False, color=color)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True

    settings = document.settings._element
    update = OxmlElement("w:updateFields")
    update.set(qn("w:val"), "true")
    settings.append(update)

    header = section.header.paragraphs[0]
    header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = header.add_run("2025年全部文件 · ASR + OCR + GPT 融合校对")
    run.font.name, run.font.size, run.font.color.rgb = FONT_CN, Pt(8.5), RGBColor.from_string(GRAY)
    run._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:eastAsia"), FONT_CN)
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    footer.add_run("第 ")
    field = OxmlElement("w:fldSimple")
    field.set(qn("w:instr"), "PAGE")
    footer._p.append(field)
    footer.add_run(" 页")
    for footer_run in footer.runs:
        footer_run.font.name, footer_run.font.size = FONT_CN, Pt(8.5)
        footer_run._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:eastAsia"), FONT_CN)


def add_cover(document: Document, corpus: dict) -> None:
    for _ in range(4):
        document.add_paragraph()
    title = document.add_heading("2025年全部文件", 0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle = document.add_paragraph("ASR + OCR + GPT 融合校对结果", style="Subtitle")
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    video_count = sum(item["kind"] == "video" for item in corpus["files"])
    courseware_count = len(corpus["files"]) - video_count
    summary = document.add_paragraph(
        f"共 {len(corpus['files'])} 个文件：{video_count} 个视频，{courseware_count} 个 PDF / 图片 / PPT / Word 课件",
        style="Subtitle",
    )
    summary.alignment = WD_ALIGN_PARAGRAPH.CENTER
    stamp = document.add_paragraph(f"整理日期：{date.today().isoformat()}", style="Subtitle")
    stamp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    document.add_paragraph().add_run().add_break(WD_BREAK.PAGE)


def add_directory(document: Document, corpus: dict) -> None:
    document.add_heading("文件目录", 1)
    table = document.add_table(rows=1, cols=3)
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    widths = [Cm(1.25), Cm(2.1), Cm(13.8)]
    for cell, value, width in zip(table.rows[0].cells, ("序号", "类型", "文件名"), widths):
        cell.text, cell.width = value, width
        _shade(cell, BLUE)
        for run in cell.paragraphs[0].runs:
            run.bold = True
            run.font.color.rgb = RGBColor(255, 255, 255)
    _repeat_header(table.rows[0])
    for index, item in enumerate(corpus["files"], 1):
        extension = Path(item["source_name"]).suffix.upper().lstrip(".")
        cells = table.add_row().cells
        for cell, value, width in zip(cells, (str(index), extension, item["source_name"]), widths):
            cell.text, cell.width = value, width
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            for paragraph in cell.paragraphs:
                paragraph.paragraph_format.space_after = Pt(2)
                paragraph.paragraph_format.space_before = Pt(2)
                for run in paragraph.runs:
                    run.font.size = Pt(9)
                    run._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:eastAsia"), FONT_CN)
        if index % 2 == 0:
            for cell in cells:
                _shade(cell, LIGHT_GRAY)
        _no_split(table.rows[-1])
    document.add_paragraph().add_run().add_break(WD_BREAK.PAGE)


def add_file(document: Document, index: int, item: dict) -> None:
    heading = document.add_heading(item["source_name"], 1)
    heading.paragraph_format.keep_with_next = True
    for section_index, title in enumerate(SECTION_TITLES):
        section = document.add_heading(title, 2)
        # Subtle semantic color bar without adding process metadata.
        properties = section._p.get_or_add_pPr()
        borders = OxmlElement("w:pBdr")
        border = OxmlElement("w:left")
        border.set(qn("w:val"), "single")
        border.set(qn("w:sz"), "18")
        border.set(qn("w:space"), "8")
        border.set(qn("w:color"), BLUE if section_index < 2 else ORANGE)
        borders.append(border)
        properties.append(borders)
        for text in item["sections"][title]:
            paragraph = document.add_paragraph(str(text))
            # Courseware OCR keeps visual reading order with manual line
            # breaks.  Justifying those lines makes Word/LibreOffice stretch
            # individual Chinese glyphs across the page, so keep multiline
            # OCR left-aligned and reserve justified prose for ASR/fusion.
            multiline_ocr = title == "OCR 结果" and "\n" in str(text)
            paragraph.alignment = (WD_ALIGN_PARAGRAPH.LEFT if multiline_ocr
                                   else WD_ALIGN_PARAGRAPH.JUSTIFY)
            paragraph.paragraph_format.first_line_indent = (Cm(0) if multiline_ocr
                                                             else Cm(0.74))
            paragraph.paragraph_format.keep_together = False
    document.add_paragraph().add_run().add_break(WD_BREAK.PAGE)


def export(corpus: dict, inventory: dict, output: Path) -> Path:
    if output.suffix.lower() != ".docx":
        raise ValueError("最终输出必须是 .docx Word 文件")
    validate_corpus(corpus, inventory)
    document = Document()
    setup(document)
    document.core_properties.title = "2025年全部文件 ASR+OCR+GPT融合校对结果"
    document.core_properties.subject = "ASR、OCR与融合校对结果"
    add_cover(document, corpus)
    add_directory(document, corpus)
    for index, item in enumerate(corpus["files"], 1):
        add_file(document, index, item)
    # Remove the otherwise blank trailing page break.
    last = document.paragraphs[-1]
    last._element.getparent().remove(last._element)
    output.parent.mkdir(parents=True, exist_ok=True)
    temporary = output.with_suffix(".tmp.docx")
    document.save(temporary)
    temporary.replace(output)
    return output


def qa_word(output: Path, corpus: dict) -> dict:
    document = Document(output)
    file_headings = [paragraph.text for paragraph in document.paragraphs
                     if paragraph.style.name == "Heading 1" and paragraph.text != "文件目录"]
    section_headings = [paragraph.text for paragraph in document.paragraphs
                        if paragraph.style.name == "Heading 2"]
    expected_file_headings = [item["source_name"] for item in corpus["files"]]
    if file_headings != expected_file_headings:
        raise ValueError("Word 文件标题顺序或数量不正确")
    if section_headings != list(SECTION_TITLES) * len(corpus["files"]):
        raise ValueError("Word 三段式标题顺序或数量不正确")
    text = "\n".join(paragraph.text for paragraph in document.paragraphs)
    forbidden = ("frame_index", "timestamp_seconds", "置信度", "画面 00:", "第 0 帧")
    if any(token in text for token in forbidden):
        raise ValueError("Word 中出现了不应展示的逐帧/技术字段")
    with ZipFile(output) as archive:
        xml = b"\n".join(archive.read(name) for name in archive.namelist() if name.endswith(".xml"))
    if re.search(rb"(?:sk|tp)-[A-Za-z0-9_-]{12,}", xml):
        raise ValueError("Word 中疑似包含API认证信息")
    return {
        "output": str(output.resolve()),
        "size_bytes": output.stat().st_size,
        "file_heading_count": len(file_headings),
        "section_heading_count": len(section_headings),
        "paragraph_count": len(document.paragraphs),
        "table_count": len(document.tables),
        "character_count": len(text),
        "forbidden_fields_absent": True,
        "credential_pattern_absent": True,
    }


def main() -> int:
    parser = argparse.ArgumentParser(description="导出2025全部文件融合校对总Word")
    parser.add_argument("--content", type=Path, default=Path(".work/batch-2025/fused_content.json"))
    parser.add_argument("--inventory", type=Path, default=Path(".work/batch-2025/inventory.json"))
    parser.add_argument("--output", type=Path, required=True)
    parser.add_argument("--qa", type=Path, default=Path(".work/batch-2025/delivery_qa.json"))
    args = parser.parse_args()
    try:
        corpus, inventory = read_json(args.content.resolve()), read_json(args.inventory.resolve())
        output = export(corpus, inventory, args.output.resolve())
        qa = qa_word(output, corpus)
        atomic_json(args.qa.resolve(), qa)
        print(json.dumps(qa, ensure_ascii=False, indent=2), flush=True)
        return 0
    except (OSError, ValueError, KeyError, TypeError) as exc:
        print(f"ERROR {type(exc).__name__}: {exc}", file=sys.stderr)
        return 1


if __name__ == "__main__":
    raise SystemExit(main())
