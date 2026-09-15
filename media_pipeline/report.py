"""Traceable, single-video Word reports; original recognition dictionaries stay intact."""
from __future__ import annotations

import copy
import difflib
import math
from pathlib import Path
from typing import Any

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor


def _time(value: float) -> str:
    milliseconds = round(float(value) * 1000)
    if milliseconds < 0:
        raise ValueError("时间不能为负数")
    seconds, milliseconds = divmod(milliseconds, 1000)
    hours, seconds = divmod(seconds, 3600)
    minutes, seconds = divmod(seconds, 60)
    return f"{hours:02d}:{minutes:02d}:{seconds:02d}.{milliseconds:03d}"


def _index(items: list[dict], label: str) -> dict[int, dict]:
    result = {}
    for item in items:
        index = item["index"]
        if not isinstance(index, int) or isinstance(index, bool) or index in result:
            raise ValueError(f"{label} 的 index 必须是唯一整数：{index!r}")
        result[index] = item
    return result


def _apply_review(asr: dict, ocr: dict, review: dict) -> tuple[dict, dict, list[dict]]:
    """Apply whole-segment/whole-line replacements only, without mutating inputs."""
    asr_result, ocr_result = copy.deepcopy(asr), copy.deepcopy(ocr)
    segments = _index(asr_result.get("segments", []), "ASR 段落")
    frames = _index(ocr_result.get("frames", []), "OCR 帧")
    changes = []
    seen = set()
    for kind, corrections in (
        ("ASR", review.get("asr_corrections", [])),
        ("OCR", review.get("ocr_corrections", [])),
    ):
        for correction in corrections:
            old, new = correction.get("old"), correction.get("new")
            if not isinstance(old, str) or not isinstance(new, str) or old == new:
                raise ValueError("纠错 old/new 必须为不同的字符串")
            if not all(isinstance(correction.get(key), str) and correction[key].strip()
                       for key in ("reason", "evidence")):
                raise ValueError("每条纠错都必须包含非空 reason 和 evidence")
            try:
                if kind == "ASR":
                    index = correction["segment_index"]
                    target = segments[index]
                    key = (kind, index)
                    location = f"段 {index}：{_time(target['start_seconds'])}—{_time(target['end_seconds'])}"
                else:
                    index, line_index = correction["frame_index"], correction["line_index"]
                    frame = frames[index]
                    if not isinstance(line_index, int) or isinstance(line_index, bool) or line_index < 0:
                        raise ValueError("OCR line_index 必须是非负整数（从 0 开始）")
                    target = frame["lines"][line_index]
                    key = (kind, index, line_index)
                    location = f"帧 {index} / 行 {line_index}：{_time(frame['timestamp_seconds'])}"
            except (KeyError, IndexError, TypeError) as exc:
                raise ValueError(f"纠错目标不存在：{correction!r}") from exc
            if key in seen:
                raise ValueError(f"同一目标存在重复纠错：{key}")
            seen.add(key)
            if target["text"] != old:
                raise ValueError(f"{kind} 原文不匹配，拒绝纠错：{location}")
            target["text"] = new
            changes.append({"kind": kind, "location": location, **correction})
    checked = review.get("checked_frames", [])
    if any(not isinstance(index, int) or isinstance(index, bool) or index not in frames for index in checked):
        raise ValueError("checked_frames 包含不存在的帧 index")
    return asr_result, ocr_result, changes


def _ocr_groups(ocr: dict) -> list[dict]:
    """Group exact full-text repetitions only across adjacent sample times."""
    frames = sorted(ocr.get("frames", []), key=lambda frame: (frame["timestamp_seconds"], frame["index"]))
    interval = float(ocr.get("sample_interval_seconds", 1.0))
    groups = []
    previous = None
    for frame in frames:
        text = "\n".join(line["text"] for line in frame.get("lines", []))
        if not text.strip():
            previous = None
            continue
        timestamp = float(frame["timestamp_seconds"])
        continuous = previous is not None and math.isclose(
            timestamp - previous["end"], interval, abs_tol=1e-6, rel_tol=0.0)
        actual = float(frame.get("actual_timestamp_seconds", timestamp))
        if continuous and previous["text"] == text:
            previous["end"], previous["actual_end"] = timestamp, actual
            previous["count"] += 1
        else:
            previous = {"start": timestamp, "end": timestamp,
                        "actual_start": actual, "actual_end": actual,
                        "count": 1, "text": text}
            groups.append(previous)
    return groups


def _ranges(indices: list[int]) -> str:
    ranges = []
    start = end = None
    for index in sorted(set(indices)):
        if start is None:
            start = end = index
        elif index == end + 1:
            end = index
        else:
            ranges.append(str(start) if start == end else f"{start}—{end}")
            start = end = index
    if start is not None:
        ranges.append(str(start) if start == end else f"{start}—{end}")
    return "、".join(ranges)


def _frame_intervals(ocr: dict, indices: list[int]) -> list[dict]:
    frames = _index(ocr.get("frames", []), "OCR 帧")
    interval = float(ocr.get("sample_interval_seconds", 1.0))
    groups = []
    previous = None
    for index in sorted(set(indices), key=lambda value: (frames[value]["timestamp_seconds"], value)):
        frame = frames[index]
        timestamp = float(frame["timestamp_seconds"])
        actual = float(frame.get("actual_timestamp_seconds", timestamp))
        if (previous is not None and index == previous["end_index"] + 1
                and math.isclose(timestamp - previous["end"], interval, abs_tol=1e-6, rel_tol=0.0)):
            previous["end_index"], previous["end"], previous["actual_end"] = index, timestamp, actual
            previous["count"] += 1
        else:
            previous = {"start_index": index, "end_index": index, "start": timestamp, "end": timestamp,
                        "actual_start": actual, "actual_end": actual, "count": 1}
            groups.append(previous)
    return groups


def _time_intervals(ocr: dict, indices: list[int]) -> str:
    return "、".join(_time(group["start"]) if group["count"] == 1 else
                    f"{_time(group['start'])}—{_time(group['end'])}"
                    for group in _frame_intervals(ocr, indices))


def _nonempty_fields(item: dict, fields: tuple[str, ...], label: str) -> None:
    if not isinstance(item, dict) or any(not isinstance(item.get(key), str) or not item[key].strip()
                                         for key in fields):
        raise ValueError(f"{label} 必须包含非空字符串：{', '.join(fields)}")


def _supplements(review: dict, duration: float) -> list[dict]:
    supplements = copy.deepcopy(review.get("ocr_supplements", []))
    if not isinstance(supplements, list):
        raise ValueError("ocr_supplements 必须为列表")
    for item in supplements:
        _nonempty_fields(item, ("text", "reason", "evidence"), "OCR AI 直接读图补录")
        start, end = item.get("start_seconds"), item.get("end_seconds")
        if (any(isinstance(value, bool) or not isinstance(value, (int, float)) or not math.isfinite(value)
                for value in (start, end)) or not 0 <= start <= end <= duration):
            raise ValueError("OCR AI 直接读图补录时间必须有限且满足 0 ≤ start_seconds ≤ end_seconds ≤ 视频时长")
    return sorted(supplements, key=lambda item: (item["start_seconds"], item["end_seconds"]))


def _extract_persistent_texts(ocr: dict, review: dict) -> tuple[dict, list[dict]]:
    """Only explicitly reviewed, exact full-line signatures may leave timed body text."""
    definitions = review.get("persistent_texts", [])
    if not isinstance(definitions, list):
        raise ValueError("persistent_texts 必须为列表")
    extracted = []
    for definition in definitions:
        _nonempty_fields(definition, ("text", "reason", "evidence"), "persistent_texts")
        tokens = definition["text"].split("\n")
        if any(not token.strip() for token in tokens):
            raise ValueError("persistent_texts 的每一行都必须非空")
        for existing in extracted:
            other = existing["tokens"]
            if any(short == long[index:index + len(short)]
                   for short, long in ((tokens, other), (other, tokens))
                   for index in range(len(long) - len(short) + 1)):
                raise ValueError("persistent_texts 不能重复或包含另一项，避免重叠匹配")
        extracted.append({**copy.deepcopy(definition), "tokens": tokens, "frame_indices": []})
    result = copy.deepcopy(ocr)
    if not extracted:
        return result, []
    for frame in result.get("frames", []):
        # Empty lines arise from explicit duplicate-box corrections and have no visible text.
        tokens = [token for line in frame.get("lines", []) for token in line["text"].split("\n") if token]
        removed = False
        for definition in extracted:
            signature = definition["tokens"]
            kept, index, found = [], 0, False
            while index < len(tokens):
                if tokens[index:index + len(signature)] == signature:
                    found, removed = True, True
                    index += len(signature)
                else:
                    kept.append(tokens[index])
                    index += 1
            tokens = kept
            if found:
                definition["frame_indices"].append(frame["index"])
        if removed:
            # For display only; original boxes/confidences and recognition remain in cache.
            frame["lines"] = [{"text": token} for token in tokens]
    for definition in extracted:
        definition["frame_indices"] = sorted(set(definition["frame_indices"]))
    return result, extracted


def _change_excerpt(old: str, new: str) -> tuple[str, str]:
    """Show change spans with ten nearby characters; full text remains elsewhere."""
    old_parts, new_parts = [], []
    matcher = difflib.SequenceMatcher(a=old, b=new, autojunk=False)
    for group in matcher.get_grouped_opcodes(n=10):
        old_start, old_end = group[0][1], group[-1][2]
        new_start, new_end = group[0][3], group[-1][4]
        old_parts.append(("…" if old_start else "") + old[old_start:old_end] + ("…" if old_end < len(old) else ""))
        new_parts.append(("…" if new_start else "") + new[new_start:new_end] + ("…" if new_end < len(new) else ""))
    def compact(parts: list[str]) -> str:
        value = "\n".join(parts)
        return value if len(value) <= 200 else value[:90] + "\n…（节略，完整文字见正文或附录）…\n" + value[-90:]
    return compact(old_parts), compact(new_parts)


def _correction_positions(changes: list[dict]) -> str:
    by_line = {}
    for change in changes:
        by_line.setdefault(change["line_index"], []).append(change["frame_index"])
    return "；".join(f"行 {line}：帧 {_ranges(indices)}" for line, indices in sorted(by_line.items()))


def _partition_persistent_corrections(changes: list[dict], persistent: list[dict]) -> tuple[list[dict], list[dict]]:
    definitions = {definition["text"]: definition for definition in persistent}
    grouped = {}
    ordinary = []
    for change in changes:
        label = None
        if change["kind"] == "OCR":
            explicit = change.get("persistent_text")
            if explicit is not None:
                if explicit not in definitions or change["new"] not in ("", explicit):
                    raise ValueError("纠错 persistent_text 必须指向已定义的固定全文，改文只能是该全文或置空")
                if change["frame_index"] not in definitions[explicit]["frame_indices"]:
                    raise ValueError("固定水印重复框置空所在帧没有匹配的完整固定文字")
                label = explicit
            elif change["new"] in definitions:
                label = change["new"]
        if label is None:
            ordinary.append(change)
        else:
            grouped.setdefault(label, []).append(change)
    return ordinary, [{"definition": definitions[label], "changes": entries} for label, entries in grouped.items()]


def _correction_rows(changes: list[dict], ocr: dict) -> list[list[str]]:
    rows = []
    for change in changes:
        if change["kind"] == "ASR":
            old, new = _change_excerpt(change["old"], change["new"])
            rows.append([f"ASR\n{change['location']}", old, new,
                         f"理由：{change['reason']}\n依据：{change['evidence']}"])
    by_identity = {}
    for change in changes:
        if change["kind"] == "OCR":
            identity = tuple(change[key] for key in ("old", "new", "reason", "evidence"))
            by_identity.setdefault(identity, []).append(change)
    for repeated in by_identity.values():
        change = repeated[0]
        if len(repeated) == 1:
            location = change["location"]
        else:
            by_line = {}
            for entry in repeated:
                by_line.setdefault(entry["line_index"], []).append(entry["frame_index"])
            positions = [f"行 {line}：帧 {_ranges(indices)}" for line, indices in sorted(by_line.items())]
            indices = [entry["frame_index"] for entry in repeated]
            location = (f"{len(repeated)} 条修改 / {len(set(indices))} 帧\n" + "\n".join(positions)
                        + "\n采样时间：" + _time_intervals(ocr, indices))
        rows.append([f"OCR\n{location}", change["old"], change["new"] or "（置空；具体理由见右栏）",
                     f"理由：{change['reason']}\n依据：{change['evidence']}"])
    return rows


def _font(style: Any, name: str, size: float, *, bold: bool = False) -> None:
    style.font.name = name
    style.font.size = Pt(size)
    style.font.bold = bold
    fonts = style.element.get_or_add_rPr().get_or_add_rFonts()
    for script in ("ascii", "hAnsi", "eastAsia", "cs"):
        fonts.set(qn(f"w:{script}"), name)


def _setup_document() -> Any:
    document = Document()
    section = document.sections[0]
    section.page_width, section.page_height = Cm(21), Cm(29.7)
    section.top_margin, section.bottom_margin = Cm(1.8), Cm(1.8)
    section.left_margin, section.right_margin = Cm(2.0), Cm(2.0)
    section.header_distance, section.footer_distance = Cm(0.8), Cm(0.8)
    normal = document.styles["Normal"]
    _font(normal, "宋体", 10.5)
    normal.paragraph_format.line_spacing = 1.15
    normal.paragraph_format.space_after = Pt(4)
    normal.paragraph_format.widow_control = True
    for name, size in (("Title", 19), ("Heading 1", 15), ("Heading 2", 12)):
        style = document.styles[name]
        _font(style, "微软雅黑", size, bold=True)
        style.font.color.rgb = RGBColor.from_string("203746")
        style.paragraph_format.keep_with_next = True
        style.paragraph_format.page_break_before = False
        style.paragraph_format.space_before = Pt(12 if name != "Title" else 0)
        style.paragraph_format.space_after = Pt(6)
    _font(document.styles["Caption"], "宋体", 9)
    caption = document.styles["Caption"].paragraph_format
    caption.space_before, caption.space_after = Pt(6), Pt(2)
    caption.keep_with_next = True
    footer = section.footer.paragraphs[0]
    footer.alignment = 1
    footer.add_run("第 ")
    field = OxmlElement("w:fldSimple")
    field.set(qn("w:instr"), "PAGE")
    footer._p.append(field)
    footer.add_run(" 页")
    for run in footer.runs:
        run.font.size = Pt(9)
    return document


def _table(document: Any, headers: list[str], rows: list[list[str]], widths: list[float]) -> Any:
    table = document.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    table.autofit = False
    for column, width in zip(table.columns, widths):
        column.width = Cm(width)
    for cell, header, width in zip(table.rows[0].cells, headers, widths):
        cell.text, cell.width = header, Cm(width)
        for run in cell.paragraphs[0].runs:
            run.bold = True
    repeat = OxmlElement("w:tblHeader")
    repeat.set(qn("w:val"), "true")
    table.rows[0]._tr.get_or_add_trPr().append(repeat)
    for values in rows:
        cells = table.add_row().cells
        for cell, value, width in zip(cells, values, widths):
            cell.text, cell.width = str(value), Cm(width)
    for row in table.rows:
        no_split = OxmlElement("w:cantSplit")
        row._tr.get_or_add_trPr().append(no_split)
        for cell in row.cells:
            for paragraph in cell.paragraphs:
                paragraph.paragraph_format.space_after = Pt(3)
                paragraph.paragraph_format.keep_with_next = False
                for run in paragraph.runs:
                    run.font.size = Pt(9)
    return table


def build_report(manifest: dict, asr: dict, ocr: dict, review: dict, output_path: Path) -> Path:
    """Write one Word report with evidence-linked corrections and honest review scope."""
    output_path = Path(output_path)
    if output_path.suffix.lower() != ".docx":
        raise ValueError("报告文件必须使用 .docx 扩展名")
    corrected_asr, corrected_ocr, changes = _apply_review(asr, ocr, review)
    frames = corrected_ocr.get("frames", [])
    display_ocr, persistent = _extract_persistent_texts(corrected_ocr, review)
    supplements = _supplements(review, float(manifest["duration_seconds"]))
    ordinary_changes, persistent_changes = _partition_persistent_corrections(changes, persistent)
    groups = _ocr_groups(display_ocr)
    segments = sorted(corrected_asr.get("segments", []), key=lambda segment: (segment["start_seconds"], segment["index"]))
    blank = sum(not any(line["text"].strip() for line in frame.get("lines", [])) for frame in frames)
    expected = ocr.get("expected_frame_count")
    checked_frames = sorted(set(review.get("checked_frames", [])))
    document = _setup_document()
    document.core_properties.title = str(manifest["source_name"])
    document.core_properties.subject = "单视频 OCR / ASR 识别及校验记录"
    document.core_properties.author = str(review.get("reviewer") or "")
    document.add_heading(str(manifest["source_name"]), 0)
    document.add_paragraph(f"源文件：{manifest['source_path']}")
    document.add_paragraph(f"视频时长：{_time(manifest['duration_seconds'])}；一源文件对应一份报告。")
    document.add_paragraph(f"ASR：{asr.get('model', '未记录')}；OCR初稿：{ocr.get('engine', '未记录')}。")
    document.add_paragraph(f"读图校对与补录：{review.get('reviewer') or '未记录'}。")

    document.add_heading("一、OCR 结果", 1)
    document.add_paragraph(
        f"采样间隔：{float(ocr.get('sample_interval_seconds', 1.0)):g} 秒；"
        f"应采样 {expected if expected is not None else '未记录'} 帧，实际记录 {len(frames)} 帧，"
        f"其中无文字 {blank} 帧；展示 {len(groups)} 组文字。")
    document.add_paragraph(
        "仅合并连续采样且全文完全相同的结果，保留采样起止时间和帧数；"
        "不同时间重新出现的文字分别列出。每秒采样不能保证捕捉到帧间短暂闪现的文字。")
    if persistent:
        document.add_heading("固定画面文字（集中列示）", 2)
        document.add_paragraph(
            "以下仅为复核记录明确指定的固定页眉或水印，按校对后完整文字逐行精确匹配，"
            "从后续逐时正文移至此处集中列示。保留全部实际匹配帧的离散区间；未匹配的相似文字及其他字幕仍在正文。")
        for definition in persistent:
            document.add_paragraph(definition["text"])
            indices = definition["frame_indices"]
            document.add_paragraph(f"匹配帧数：{len(indices)}；帧 index：{_ranges(indices) or '无精确匹配'}。")
            intervals = _frame_intervals(corrected_ocr, indices)
            seconds = "、".join(f"{item['start']:g}" if item["count"] == 1 else
                                f"{item['start']:g}—{item['end']:g}" for item in intervals)
            document.add_paragraph(f"出现采样秒（离散范围）：{seconds or '无'}。")
            if all(item["start"] == item["actual_start"] and item["end"] == item["actual_end"] for item in intervals):
                document.add_paragraph("实际取帧时间与上述采样秒完全一致。")
            else:
                actual_seconds = "、".join(f"{item['actual_start']:.6f}—{item['actual_end']:.6f}" for item in intervals)
                document.add_paragraph(f"对应实际取帧秒（逐范围对应）：{actual_seconds}。")
            document.add_paragraph(f"集中列示理由：{definition['reason']}\n读图依据：{definition['evidence']}")
        document.add_heading("按时间排列的其他画面文字", 2)
    if not groups:
        document.add_paragraph("除上述集中列示的固定文字外，本次没有其他识别文字。" if persistent
                               else "本次采样帧未识别出文字。")
    for group in groups:
        sampling = _time(group["start"])
        actual = _time(group["actual_start"])
        if group["count"] > 1:
            sampling += "—" + _time(group["end"])
            actual += "—" + _time(group["actual_end"])
        document.add_paragraph(f"采样时间 {sampling}｜{group['count']} 帧｜实际帧时间 {actual}", style="Caption")
        document.add_paragraph(group["text"])
    if supplements:
        document.add_heading("AI直接读图补录", 2)
        document.add_paragraph(
            "以下为 AI 直接查看视频画面后的补漏，独立列示补录文字、时间及依据，"
            "不是 RapidOCR 原始识别结果；原始帧识别记录保持不变。")
        for supplement in supplements:
            document.add_paragraph(
                f"画面时间 {_time(supplement['start_seconds'])}—{_time(supplement['end_seconds'])}", style="Caption")
            document.add_paragraph(supplement["text"])
            document.add_paragraph(f"补录理由：{supplement['reason']}\n读图依据：{supplement['evidence']}")

    document.add_heading("二、ASR 结果", 1)
    document.add_paragraph(
        f"共 {len(segments)} 段。以下时间为源视频中的音频分段范围，非逐字时间戳；"
        "校对只应用于校验记录中明确列出的修改。")
    if not segments:
        document.add_paragraph("未提供 ASR 段落；不能据此判定视频无语音。")
    for segment in segments:
        document.add_paragraph(
            f"段 {segment['index']}｜{_time(segment['start_seconds'])}—{_time(segment['end_seconds'])}",
            style="Caption")
        document.add_paragraph(segment["text"] or "（本段识别结果为空）")

    document.add_heading("三、校验记录", 1)
    document.add_paragraph(
        f"复核人：{review.get('reviewer') or '未记录'}；复核时间：{review.get('reviewed_at') or '未记录'}。")
    document.add_paragraph(
        f"记录已查看画面 {len(checked_frames)} / {len(frames)} 帧。"
        "复核范围以本节记录为准，不代表所有文字均已确认无误。")
    if checked_frames:
        document.add_paragraph("已查看帧 index（与原始记录一致，连续编号以范围表示）：" + _ranges(checked_frames))
    coverage = "未记录应采样帧数，无法判定覆盖"
    if expected is not None:
        coverage = "记录帧数与应采样帧数一致" if len(frames) == expected else "记录帧数与应采样帧数不一致，需核实"
    _table(document, ["检查项", "记录"], [
        ["OCR 帧数核对", coverage],
        ["OCR 有文字 / 无文字", f"{len(frames) - blank} / {blank}"],
        ["OCR 展示组数", str(len(groups))],
        ["OCR 固定文字 / AI 读图补录", f"{len(persistent)} 项 / {len(supplements)} 条"],
        ["ASR 段落数", str(len(segments))],
        ["ASR / OCR 纠错数", f"{sum(change['kind'] == 'ASR' for change in changes)} / {sum(change['kind'] == 'OCR' for change in changes)}"],
    ], [5, 12])
    document.add_heading("纠错明细", 2)
    if changes:
        document.add_paragraph(
            "ASR 表内列出变化片段及邻近上下文，完整校对文字见 ASR 正文，完整原文见附录。"
            "OCR 原文、改文、理由和依据均相同的修改合并列示，保留帧号与行号的全部离散范围；"
            "统计仍按实际修改条数计算。正文中的普通字幕不会因此合并。")
        if ordinary_changes:
            _table(document, ["类型 / 位置", "原文 / 变化片段", "校对文字 / 变化片段", "理由及依据"],
                   _correction_rows(ordinary_changes, ocr), [3.4, 4.0, 4.0, 5.6])
        if persistent_changes:
            document.add_heading("固定标语/水印校对汇总", 2)
            document.add_paragraph(
                "以下合并展示已明确指定的固定文字校正及显式标记的同帧重复水印框置空。"
                "列出全部原识别变体及准确帧号/行号范围；完整逐条原文、改文、理由与画面依据仍保存在任务 review.json。"
                "“↵”表示原识别文字中的换行。")
            for group in persistent_changes:
                definition, entries = group["definition"], group["changes"]
                full = [entry for entry in entries if entry["new"]]
                empty = [entry for entry in entries if not entry["new"]]
                document.add_paragraph("校对后固定全文：" + definition["text"])
                document.add_paragraph(
                    f"本项共 {len(entries)} 条修改：完整校正 {len(full)} 条，重复水印局部框置空 {len(empty)} 条。")
                variants = sorted(set(entry["old"] for entry in entries))
                document.add_paragraph("原识别变体集合：" + "、".join("“" + value.replace("\n", "↵") + "”" for value in variants))
                if full:
                    document.add_paragraph("完整校正位置：" + _correction_positions(full))
                if empty:
                    document.add_paragraph("重复框置空位置：" + _correction_positions(empty))
                document.add_paragraph(f"汇总理由：{definition['reason']}\n读图依据：{definition['evidence']}")
    else:
        document.add_paragraph("未记录文字修改；这不等于识别文字全部正确。")
    document.add_heading("复核说明", 2)
    for note in review.get("notes", []) or ["未提供额外复核说明。"]:
        document.add_paragraph(str(note))
    document.add_heading("待核实事项", 2)
    for note in review.get("unresolved", []) or ["未记录待核实事项；不代表已排除所有识别错误。"]:
        document.add_paragraph(str(note))
    document.add_paragraph("源文件 SHA-256：" + str(manifest.get("sha256", "未记录")))
    document.add_paragraph("原始识别结果保留在任务缓存中，未被校对覆盖。OCR 行号 line_index 从 0 开始。")
    if any(change["kind"] == "ASR" for change in changes):
        document.add_heading("附录：原始 ASR 结果", 1)
        for segment in sorted(asr.get("segments", []), key=lambda value: (value["start_seconds"], value["index"])):
            document.add_paragraph(
                f"段 {segment['index']}｜{_time(segment['start_seconds'])}—{_time(segment['end_seconds'])}",
                style="Caption")
            document.add_paragraph(segment["text"] or "（本段识别结果为空）")
    output_path.parent.mkdir(parents=True, exist_ok=True)
    document.save(output_path)
    return output_path
