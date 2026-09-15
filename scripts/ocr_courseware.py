#!/usr/bin/env python3
"""Render courseware from an archive and OCR every visual page.

The archive is treated as read-only. Results are written as JSON so the Word
builder can merge them with the reviewed video transcript corpus.
"""

from __future__ import annotations

import argparse
import hashlib
import io
import json
import os
import re
import subprocess
import sys
import time
import xml.etree.ElementTree as ET
from difflib import SequenceMatcher
from pathlib import Path
from typing import Any, Iterable
from zipfile import ZipFile

import cv2
import numpy as np
import pymupdf
from docx import Document
from PIL import Image, ImageSequence
from pptx import Presentation


TYPE_ORDER = {".pdf": 0, ".png": 1, ".jpg": 2, ".jpeg": 2, ".pptx": 3, ".docx": 4}
TYPE_LABEL = {
    ".pdf": "PDF",
    ".png": "PNG",
    ".jpg": "JPG",
    ".jpeg": "JPEG",
    ".pptx": "PPTX",
    ".docx": "DOCX",
}
PROCESSOR_VERSION = "courseware-ocr-lossless-v5"

# Only high-confidence normalizations are applied. Native PPTX/DOCX text is
# used as a proofreading reference; it is never presented as raw OCR output.
TEXT_CORRECTIONS = (
    ("云南电纲", "云南电网"),
    ("云南电间", "云南电网"),
    ("中国共声党", "中国共产党"),
    ("共青四", "共青团"),
    ("青马微课常", "青马微课堂"),
    ("职场沟迦", "职场沟通"),
)


def clean_text(text: str) -> str:
    text = str(text or "").replace("\u3000", " ").replace("\ufeff", "")
    text = re.sub(r"[ \t]+", " ", text)
    text = re.sub(r"\s*\n\s*", "\n", text).strip()
    for source, replacement in TEXT_CORRECTIONS:
        text = text.replace(source, replacement)
    return text


def key_text(text: str) -> str:
    return re.sub(r"[^0-9a-z\u3400-\u9fff]+", "", clean_text(text).casefold())


def split_lines(values: Iterable[str]) -> list[str]:
    output: list[str] = []
    seen: set[str] = set()
    for value in values:
        for raw in re.split(r"[\r\n]+", str(value or "")):
            line = clean_text(raw)
            key = key_text(line)
            if len(key) < 2 or key in seen:
                continue
            seen.add(key)
            output.append(line)
    return output


def safe_extract(archive: Path, output_dir: Path) -> list[Path]:
    output_dir.mkdir(parents=True, exist_ok=True)
    extracted: list[Path] = []
    root = output_dir.resolve()
    with ZipFile(archive) as handle:
        for item in handle.infolist():
            if item.is_dir():
                continue
            target = (output_dir / item.filename).resolve()
            if root not in target.parents:
                raise RuntimeError(f"压缩包包含不安全路径：{item.filename}")
            target.parent.mkdir(parents=True, exist_ok=True)
            with handle.open(item) as source, target.open("wb") as destination:
                while chunk := source.read(1024 * 1024):
                    destination.write(chunk)
            extracted.append(target)
    return extracted


def nested_shape_text(shape: Any) -> list[str]:
    values: list[str] = []
    if getattr(shape, "has_text_frame", False):
        values.extend(paragraph.text for paragraph in shape.text_frame.paragraphs)
    if getattr(shape, "has_table", False):
        for row in shape.table.rows:
            values.extend(cell.text for cell in row.cells)
    for child in getattr(shape, "shapes", ()):
        values.extend(nested_shape_text(child))
    return values


def pptx_native_pages(path: Path) -> list[list[str]]:
    presentation = Presentation(path)
    pages: list[list[str]] = []
    for slide in presentation.slides:
        values = [text for shape in slide.shapes for text in nested_shape_text(shape)]
        values.extend(
            node.text or ""
            for node in slide.element.iter()
            if str(node.tag).endswith("}t")
        )
        values.extend(
            value
            for node in slide.element.iter()
            for key, value in node.attrib.items()
            if str(key).endswith(("}descr", "}title")) and value
        )
        if slide.has_notes_slide:
            values.extend(
                paragraph.text
                for paragraph in slide.notes_slide.notes_text_frame.paragraphs
            )
        for shape in slide.shapes:
            if not getattr(shape, "has_chart", False):
                continue
            chart = shape.chart
            if chart.has_title:
                values.append(chart.chart_title.text_frame.text)
            for series in chart.series:
                values.append(str(series.name or ""))
                values.extend(str(value) for value in series.values if value is not None)
        pages.append(split_lines(values))
    with ZipFile(path) as archive:
        for name in sorted(archive.namelist()):
            if not name.startswith("ppt/embeddings/") or not name.lower().endswith(".xlsx"):
                continue
            try:
                from openpyxl import load_workbook
                workbook = load_workbook(io.BytesIO(archive.read(name)), read_only=True, data_only=False)
                embedded = [f"嵌入工作簿：{Path(name).name}"]
                for sheet in workbook.worksheets:
                    embedded.append(f"工作表：{sheet.title}")
                    for row in sheet.iter_rows():
                        values = [str(cell.value) for cell in row if cell.value is not None]
                        if values:
                            embedded.append(" | ".join(values))
                if pages:
                    pages[0] = split_lines([*pages[0], *embedded])
            except Exception:
                # The rendered chart still remains in the whole-slide OCR.
                continue
    return pages


def nested_picture_images(shape: Any) -> Iterable[np.ndarray]:
    """Yield raster pictures embedded in a PPTX shape tree.

    PPTX text is read directly from OOXML because that is more accurate than
    OCR.  Pictures still need OCR so that screenshots and text baked into
    diagrams are not silently missed.  Unsupported vector image formats are
    skipped here; their captions and nearby labels remain available from the
    native slide text.
    """
    for child in getattr(shape, "shapes", ()):
        yield from nested_picture_images(child)
    try:
        image = getattr(shape, "image", None)
    except ValueError:
        # Linked-picture placeholders can advertise a picture shape without
        # carrying an embedded image part. The rendered slide OCR below still
        # captures their visible appearance when LibreOffice can resolve it.
        return
    if image is None:
        return
    try:
        blob = image.blob
    except ValueError:
        return
    if not blob:
        return
    try:
        animated = Image.open(io.BytesIO(blob))
        if getattr(animated, "n_frames", 1) > 1:
            seen_frames: set[str] = set()
            for frame in ImageSequence.Iterator(animated):
                rgb = np.array(frame.convert("RGB"))
                digest = hashlib.sha256(rgb.tobytes()).hexdigest()
                if digest in seen_frames:
                    continue
                seen_frames.add(digest)
                yield cv2.cvtColor(rgb, cv2.COLOR_RGB2BGR)
            return
    except Exception:
        pass
    decoded = cv2.imdecode(np.frombuffer(blob, dtype=np.uint8), cv2.IMREAD_COLOR)
    if decoded is not None:
        yield decoded


def pptx_pages(
    path: Path,
    engine: Any,
    minimum_confidence: float,
) -> Iterable[tuple[int, list[dict[str, Any]], list[str]]]:
    """Extract exact slide text and OCR every raster picture on each slide."""
    presentation = Presentation(path)
    for page_number, slide in enumerate(presentation.slides, 1):
        native = split_lines(
            text
            for shape in slide.shapes
            for text in nested_shape_text(shape)
        )
        picture_lines: list[dict[str, Any]] = []
        seen: set[str] = set()
        for shape in slide.shapes:
            for picture in nested_picture_images(shape):
                for item in ocr_image(engine, picture, minimum_confidence):
                    key = key_text(item["text"])
                    if key in seen:
                        continue
                    seen.add(key)
                    picture_lines.append(item)
        yield page_number, picture_lines, native


def docx_native_text(path: Path) -> list[str]:
    document = Document(path)
    values = [paragraph.text for paragraph in document.paragraphs]
    values.extend(cell.text for table in document.tables for row in table.rows for cell in row.cells)
    with ZipFile(path) as archive:
        for name in sorted(archive.namelist()):
            if not name.startswith("word/") or not name.endswith(".xml"):
                continue
            try:
                root = ET.fromstring(archive.read(name))
            except ET.ParseError:
                continue
            values.extend(
                node.text or ""
                for node in root.iter()
                if str(node.tag).endswith("}t")
            )
            values.extend(
                value
                for node in root.iter()
                for key, value in node.attrib.items()
                if str(key).endswith(("}descr", "}title")) and value
            )
    return split_lines(values)


def docx_media_ocr(path: Path, engine: Any, minimum_confidence: float) -> list[dict[str, Any]]:
    output: list[dict[str, Any]] = []
    seen: set[str] = set()
    with ZipFile(path) as archive:
        for name in sorted(archive.namelist()):
            if not name.startswith("word/media/") or name.endswith("/"):
                continue
            payload = archive.read(name)
            if not payload:
                continue
            decoded = cv2.imdecode(
                np.frombuffer(payload, dtype=np.uint8), cv2.IMREAD_COLOR)
            if decoded is None:
                continue
            for item in ocr_image(engine, decoded, minimum_confidence):
                key = key_text(item["text"])
                if not key or key in seen:
                    continue
                seen.add(key)
                output.append(item)
    return output


def pdf_native_pages(path: Path) -> list[list[str]]:
    with pymupdf.open(path) as document:
        return [split_lines(page.get_text("text").splitlines()) for page in document]


def run_soffice(source: Path, output_dir: Path, soffice: Path, profile_dir: Path) -> Path:
    output_dir.mkdir(parents=True, exist_ok=True)
    profile_dir.mkdir(parents=True, exist_ok=True)
    # LibreOffice 24.2 in the portable runtime silently rejects a few supplied
    # long/non-ASCII PPTX basenames.  A source-bound ASCII hardlink avoids that
    # importer bug without modifying or duplicating the original document.
    alias = output_dir / f"input-{sha256_file(source)[:16]}{source.suffix.lower()}"
    if not alias.exists():
        os.link(source, alias)
    destination = alias.with_suffix(".pdf")
    if destination.is_file():
        return destination
    profile_uri = profile_dir.resolve().as_uri()
    command = [
        str(soffice),
        "--headless",
        "--nologo",
        "--nodefault",
        "--nofirststartwizard",
        f"-env:UserInstallation={profile_uri}",
        "--convert-to",
        "pdf",
        "--outdir",
        str(output_dir),
        str(alias),
    ]
    environment = os.environ.copy()
    environment["SOFFICE_TASK_TMPDIR"] = str((profile_dir / "tmp").resolve())
    completed = subprocess.run(
        command,
        text=True,
        stdout=subprocess.PIPE,
        stderr=subprocess.STDOUT,
        timeout=900,
        env=environment,
    )
    if completed.returncode != 0 or not destination.is_file():
        raise RuntimeError(f"LibreOffice 渲染失败：{completed.stdout.strip()}")
    return destination


def slice_array(image: np.ndarray) -> list[np.ndarray]:
    height, width = image.shape[:2]
    slice_height = max(1600, round(width * 1.45))
    if height <= slice_height:
        return [image]
    overlap = min(160, slice_height // 12)
    slices: list[np.ndarray] = []
    start = 0
    while start < height:
        end = min(height, start + slice_height)
        slices.append(image[start:end])
        if end == height:
            break
        start = end - overlap
    return slices


def pdf_images(path: Path, zoom: float) -> Iterable[tuple[int, int, int, np.ndarray]]:
    with pymupdf.open(path) as document:
        matrix = pymupdf.Matrix(zoom, zoom)
        for index, page in enumerate(document, 1):
            embedded = page.get_images(full=True)
            if embedded:
                largest = max(embedded, key=lambda item: int(item[2]) * int(item[3]))
                embedded_width, embedded_height = int(largest[2]), int(largest[3])
                if embedded_height / max(1, embedded_width) > 2.5:
                    payload = document.extract_image(largest[0])["image"]
                    decoded = cv2.imdecode(np.frombuffer(payload, dtype=np.uint8), cv2.IMREAD_COLOR)
                    if decoded is not None:
                        slices = slice_array(decoded)
                        for part, image in enumerate(slices, 1):
                            yield index, part, len(slices), image
                        continue
            # Some supplied PDFs are a single very tall poster page. Rendering
            # those at a fixed zoom makes Chinese glyphs too small for OCR.
            page_zoom = max(zoom, min(14.0, 2400.0 / max(1.0, page.rect.width)))
            page_matrix = pymupdf.Matrix(page_zoom, page_zoom) if page.rect.height / page.rect.width > 2.5 else matrix
            pixmap = page.get_pixmap(matrix=page_matrix, alpha=False, colorspace=pymupdf.csRGB)
            rgb = np.frombuffer(pixmap.samples, dtype=np.uint8).reshape(pixmap.height, pixmap.width, 3)
            slices = slice_array(cv2.cvtColor(rgb, cv2.COLOR_RGB2BGR))
            for part, image in enumerate(slices, 1):
                yield index, part, len(slices), image


def image_slices(path: Path) -> Iterable[tuple[int, int, int, np.ndarray]]:
    raw = np.fromfile(path, dtype=np.uint8)
    image = cv2.imdecode(raw, cv2.IMREAD_COLOR)
    if image is None:
        raise RuntimeError(f"无法读取图片：{path.name}")
    slices = slice_array(image)
    for part, chunk in enumerate(slices, 1):
        yield 1, part, len(slices), chunk


def ocr_image(engine: Any, image: np.ndarray, minimum_confidence: float) -> list[dict[str, Any]]:
    height, width = image.shape[:2]
    maximum_dimension = 3200
    scale = min(1.0, maximum_dimension / max(height, width))
    if scale < 1.0:
        image = cv2.resize(image, None, fx=scale, fy=scale, interpolation=cv2.INTER_AREA)
    result = engine(image, text_score=0.0)
    lines: list[dict[str, Any]] = []
    seen: set[str] = set()
    for text, score in zip(result.txts or (), result.scores or ()):
        text = clean_text(text)
        key = key_text(text)
        score = float(score)
        if score < minimum_confidence or len(key) < 2 or key in seen:
            continue
        seen.add(key)
        lines.append({"text": text, "confidence": round(score, 4)})
    return lines


def is_represented(line: str, references: list[str]) -> bool:
    candidate = key_text(line)
    if not candidate:
        return True
    for reference in references:
        ref = key_text(reference)
        if candidate in ref or ref in candidate:
            return True
        if SequenceMatcher(None, candidate, ref, autojunk=False).ratio() >= 0.62:
            return True
    return False


def merge_with_native(ocr_lines: list[dict[str, Any]], native_lines: list[str]) -> tuple[list[str], int]:
    """Prefer exact native text and retain OCR-only text from images/charts."""
    merged = list(native_lines)
    supplements = 0
    for item in ocr_lines:
        line = item["text"]
        if is_represented(line, merged):
            continue
        if float(item["confidence"]) < 0.58:
            continue
        merged.append(line)
        supplements += 1
    return merged, supplements


def atomic_json(path: Path, value: Any) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    temporary = path.with_suffix(path.suffix + ".tmp")
    temporary.write_text(json.dumps(value, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    temporary.replace(path)


def sha256_file(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for block in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(block)
    return digest.hexdigest()


def process_file(
    source: Path,
    index: int,
    engine: Any,
    args: argparse.Namespace,
) -> dict[str, Any]:
    extension = source.suffix.lower()
    rendered_dir = args.work_dir / "rendered" / f"{index:02d}"
    native_pages: list[list[str]] = []
    native_document: list[str] = []
    picture_ocr_by_page: dict[int, list[dict[str, Any]]] = {}
    document_media_ocr: list[dict[str, Any]] = []
    if extension == ".pptx":
        native_pages = pptx_native_pages(source)
        for page_number, picture_lines, _ in pptx_pages(
                source, engine, args.minimum_confidence):
            picture_ocr_by_page[page_number] = picture_lines
        visual = run_soffice(source, rendered_dir, args.soffice, args.work_dir / "lo-profile")
        pages = pdf_images(visual, args.zoom)
    elif extension == ".docx":
        native_document = docx_native_text(source)
        document_media_ocr = docx_media_ocr(source, engine, args.minimum_confidence)
        visual = run_soffice(source, rendered_dir, args.soffice, args.work_dir / "lo-profile")
        pages = pdf_images(visual, args.zoom)
    elif extension == ".pdf":
        native_pages = pdf_native_pages(source)
        pages = pdf_images(source, args.zoom)
    elif extension in {".png", ".jpg", ".jpeg"}:
        pages = image_slices(source)
    else:
        raise RuntimeError(f"不支持的课件格式：{extension}")

    started = time.perf_counter()
    page_records: list[dict[str, Any]] = []
    for page_number, part_number, part_count, image in pages:
        ocr_lines = ocr_image(engine, image, args.minimum_confidence)
        if extension == ".pptx":
            seen = {key_text(item["text"]) for item in ocr_lines}
            for item in picture_ocr_by_page.get(page_number, []):
                key = key_text(item["text"])
                if key and key not in seen:
                    seen.add(key)
                    ocr_lines.append(item)
        native = native_pages[page_number - 1] if page_number <= len(native_pages) else []
        merged, supplements = merge_with_native(ocr_lines, native) if native else (
            [item["text"] for item in ocr_lines],
            0,
        )
        page_records.append(
            {
                "page": page_number,
                "part": part_number,
                "part_count": part_count,
                "label": (
                    f"第 {page_number} 页 · 分片 {part_number}/{part_count}"
                    if part_count > 1
                    else f"第 {page_number} 页"
                ),
                "ocr_lines": ocr_lines,
                "native_reference_lines": native,
                "reviewed_lines": merged,
                "ocr_supplement_count": supplements,
                "render_mode": (
                    "LibreOffice整页渲染OCR + OOXML原生文字/备注/图表"
                    if extension == ".pptx"
                    else "整页OCR + 原生文字交叉校正"
                ),
            }
        )
        print(
            f"    第 {page_number} 页：{len(native)} 行原生文字，{len(ocr_lines)} 行整页OCR",
            flush=True,
        )

    if native_document:
        all_ocr = [item for page in page_records for item in page["ocr_lines"]]
        all_ocr.extend(document_media_ocr)
        reviewed, supplements = merge_with_native(all_ocr, native_document)
        # DOCX pagination cannot be mapped safely from OOXML, so retain the
        # reviewed whole-document text and the raw per-page OCR evidence.
        reviewed_document = reviewed
        document_supplements = supplements
    else:
        reviewed_document = []
        document_supplements = 0

    return {
        "source": source.name,
        "source_sha256": sha256_file(source),
        "archive": args.source_label,
        "file_type": TYPE_LABEL[extension],
        "extension": extension,
        "page_count": len(page_records),
        "pages": page_records,
        "reviewed_document_lines": reviewed_document,
        "document_ocr_supplement_count": document_supplements,
        "processor_version": PROCESSOR_VERSION,
        "minimum_confidence": args.minimum_confidence,
        "render_zoom": args.zoom,
        "processing_seconds": round(time.perf_counter() - started, 3),
    }


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description="Render and OCR archived courseware")
    source = parser.add_mutually_exclusive_group(required=True)
    source.add_argument("--archive", type=Path)
    source.add_argument(
        "--source-dir",
        type=Path,
        action="append",
        help="直接读取一个或多个课件目录（可重复指定），不复制源文件",
    )
    parser.add_argument("--work-dir", type=Path, required=True)
    parser.add_argument("--soffice", type=Path, required=True)
    parser.add_argument("--output", type=Path, required=True)
    parser.add_argument("--zoom", type=float, default=2.0)
    parser.add_argument("--minimum-confidence", type=float, default=0.48)
    return parser.parse_args()


def main() -> int:
    args = parse_args()
    args.work_dir = args.work_dir.resolve()
    args.soffice = args.soffice.resolve()
    args.output = args.output.resolve()
    if args.archive is not None:
        args.archive = args.archive.resolve()
        extracted = safe_extract(args.archive, args.work_dir / "courseware")
        args.source_label = args.archive.name
    else:
        roots = [path.resolve(strict=True) for path in args.source_dir]
        extracted = [path for root in roots for path in root.rglob("*") if path.is_file()]
        args.source_label = " + ".join(str(path) for path in roots)
    sources = sorted(
        (path for path in extracted if path.suffix.lower() in TYPE_ORDER),
        key=lambda path: (TYPE_ORDER[path.suffix.lower()], path.name.casefold()),
    )
    duplicate_names = sorted(
        name for name in {path.name for path in sources}
        if sum(path.name == name for path in sources) > 1
    )
    if duplicate_names:
        raise RuntimeError(f"课件目录存在同名文件：{duplicate_names[0]}")
    from rapidocr import RapidOCR

    engine = RapidOCR()
    records: list[dict[str, Any]] = []
    cache_dir = args.work_dir / "ocr-cache"
    cache_dir.mkdir(parents=True, exist_ok=True)
    for index, source in enumerate(sources, 1):
        cache = cache_dir / f"{index:02d}.json"
        print(f"[{index}/{len(sources)}] {source.name}", flush=True)
        if cache.is_file():
            cached = json.loads(cache.read_text(encoding="utf-8"))
            if (cached.get("source") == source.name
                    and cached.get("source_sha256") == sha256_file(source)
                    and cached.get("processor_version") == PROCESSOR_VERSION
                    and cached.get("minimum_confidence") == args.minimum_confidence
                    and cached.get("render_zoom") == args.zoom):
                records.append(cached)
                print("    使用已完成缓存", flush=True)
                continue
        record = process_file(source, index, engine, args)
        atomic_json(cache, record)
        records.append(record)
    result = {
        "schema_version": "1.0",
        "archive": args.source_label,
        "archive_sha256": sha256_file(args.archive) if args.archive is not None else None,
        "engine": "RapidOCR 3.9.2 / PP-OCRv6 small",
        "processor_version": PROCESSOR_VERSION,
        "file_count": len(records),
        "page_count": sum(item["page_count"] for item in records),
        "files": records,
    }
    atomic_json(args.output, result)
    print(f"完成：{len(records)} 个课件，{result['page_count']} 个页面/图片分片", flush=True)
    return 0


if __name__ == "__main__":
    sys.exit(main())
