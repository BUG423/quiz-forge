#!/usr/bin/env python3
"""Bind, validate and export independently authored exam points.

This module never extracts, summarizes, rewrites or supplements an exam-point
body.  It only packages authoritative ASR/OCR evidence for GPT authors,
validates identity-bound result JSON, reports progress, lays validated text out
in a separate Word file, and performs read-only QA.
"""
from __future__ import annotations

import argparse
from collections import defaultdict
import hashlib
import json
import os
import re
import shutil
import sys
import tempfile
from pathlib import Path
from typing import Any, Iterable
from zipfile import ZipFile

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor


ROOT = Path(__file__).resolve().parent.parent
WORK = ROOT / ".work" / "first-principles-2026" / "exam-points"
DEFAULT_MANIFEST = ROOT / ".work" / "first-principles-2026" / "asset_manifest.json"
DEFAULT_EVIDENCE = ROOT / ".work" / "first-principles-2026" / "evidence_corpus.json"
DEFAULT_TASK_DIR = WORK / "tasks"
DEFAULT_RESULT_DIR = WORK / "results"
DEFAULT_FINAL = WORK / "exam_points_final.json"
DEFAULT_PROGRESS_JSON = WORK / "progress.json"
DEFAULT_PROGRESS_MD = WORK / "progress.md"
DEFAULT_WORD = ROOT / "outputs" / "最终交付" / "2026年YW第一课全量考点.docx"
DEFAULT_QA = WORK / "exam_points_qa.json"
DEFAULT_DISPLAY_NAMES = ROOT / "config" / "word_display_names.json"
DEFAULT_SPEC = ROOT / "docs" / "exam-points-delivery-spec.md"
DEFAULT_TEMPLATE = ROOT / "docs" / "exam-points-structure-template.md"

ASSET_COUNT = 91
CATALOG_ENTRY_COUNT = 81
KNOWN_UNCOVERED_CATALOG_POSITIONS = frozenset({6, 12, 13, 38, 45})
TASK_SCHEMA = "exam-points-task-index/v1"
FINAL_SCHEMA = "exam-points/v1"
PROGRESS_SCHEMA = "exam-points-progress/v1"
SHA_RE = re.compile(r"[0-9a-f]{64}\Z")
CONTROL_RE = re.compile(r"[\r\n\t\x00-\x08\x0b\x0c\x0e-\x1f\x7f]")
SECRET_PATTERNS = (
    re.compile(rb"(?:sk|tp)-[A-Za-z0-9_-]{12,}"),
    re.compile(rb"api[_ -]?key\s*[:=]\s*[A-Za-z0-9._-]{12,}", re.I),
    re.compile(rb"bearer\s+[A-Za-z0-9._-]{16,}", re.I),
)
FORBIDDEN_EVIDENCE_KEYS = frozenset({
    "paragraphs", "points", "exam_points", "gpt", "gpt_result", "gpt_final",
    "gpt_fusion", "gpt_fused", "fusion_text", "fused_content", "generated_prose",
    "GPT 融合校对结果", "GPT 融合结果",
})
FORBIDDEN_EVIDENCE_KEYS_CASEFOLD = frozenset(
    item.casefold() for item in FORBIDDEN_EVIDENCE_KEYS
)

DIMENSION_KEYS = (
    "concept_definition",
    "number_date_unit",
    "proper_noun",
    "subject_responsibility",
    "rule_scope",
    "process_sequence",
    "risk_prohibition",
    "case_fact_conclusion",
    "system_ui_operation",
    "comparison_confusion",
    "visual_only_information",
    "emphasis_and_negation",
)
DIMENSION_LABELS = frozenset({
    "定义概念", "数字日期比例单位", "专名", "主体职责", "制度条款",
    "范围条件例外", "流程顺序", "风险禁令", "案例事实与结论",
    "系统菜单字段按钮", "对比辨析与易错点", "画面或附件独有信息",
    "否定词与强调词",
})
MUST_CHECK = (
    "定义、概念、术语、原理、观点、目标、意义、特征和分类",
    "数字、日期、年份、次数、比例、金额、时限、阈值、数量、单位、层级、型号和序号",
    "人名、地名、组织名、制度名、文件名、工程名、系统名、品牌名和产品名",
    "主体、对象、职责、权限、分工、责任边界、审批关系和问责关系",
    "制度条款、适用范围、触发条件、前置条件、例外、豁免、禁令和后果",
    "流程、步骤、先后顺序、输入输出、办理路径、检查点、结束条件和回退路径",
    "风险、隐患、红线、错误做法、事故原因、防控措施和应急要求",
    "案例人物、时间、地点、背景、行为、处置、结果、责任、原因和启示",
    "系统入口、菜单、页签、字段、按钮、选项、状态、提示、材料和操作结果",
    "相近概念、易混选项、上下位关系、并列项、相反项和适用场景差异",
    "图表、流程图、时间线、组织图、界面截图以及画面或附件独有信息",
    "不、不得、禁止、严禁、除外、仅、至少、首先、之前、之后和同时等边界词",
)
EVIDENCE_STATUSES = frozenset({
    "ASR+OCR互证", "ASR单证", "OCR单证", "证据冲突待复核",
})
REVIEW_VALUES = frozenset({"reviewed", "not_applicable"})
SUPPORT_VALUES = frozenset({"direct", "corroborating", "conflict"})
RESULT_KEYS = frozenset({
    "order", "batch_id", "catalog_position", "catalog_suborder", "catalog_entry",
    "coverage_positions", "asset_id", "source_name", "source_sha256",
    "evidence_sha256", "task_sha256", "model", "source_limitations",
    "dimension_review", "points", "asset_review",
})
BASIS_KEYS = frozenset({
    "asset_manifest_sha256", "evidence_corpus_sha256", "delivery_spec_sha256",
    "structure_template_sha256", "old_gpt_fusion_used",
    "automatic_point_generation_used", "authoring_method",
})
POINT_KEYS = frozenset({
    "point_id", "dimension", "statement", "question_angles", "answer_elements",
    "traps", "evidence_refs", "evidence_status",
})
REFERENCE_KEYS = frozenset({"modality", "locator", "support"})
ASSET_REVIEW_KEYS = frozenset({
    "full_evidence_read", "file_level_final_review", "numbers_rechecked",
    "names_rechecked", "conditions_negations_rechecked", "process_order_rechecked",
    "ui_fields_rechecked", "unresolved_conflicts",
})
TASK_FILE_KEYS = (
    "identity", "ASR 结果", "OCR 结果", "task_constraints",
)
TASK_CONSTRAINT_KEYS = frozenset({
    "objective", "authoring_method", "automatic_point_generation_allowed",
    "old_gpt_fusion_allowed", "file_level_final_review_required", "must_check",
    "evidence_locator_format", "delivery_spec_sha256", "structure_template_sha256",
    "required_result_keys",
})
INDEX_TASK_KEYS = frozenset({
    "order", "batch_id", "catalog_position", "catalog_suborder", "catalog_entry",
    "coverage_positions", "asset_id", "source_name", "source_sha256",
    "evidence_sha256", "task_file", "task_sha256", "evidence_character_count",
})
INDEX_BATCH_KEYS = frozenset({
    "batch_id", "task_count", "evidence_character_count", "orders",
})
INDEX_ROOT_KEYS = frozenset({
    "schema_version", "status", "file_count", "catalog_entry_count", "batch_count",
    "asset_manifest_sha256", "evidence_corpus_sha256", "delivery_spec_sha256",
    "structure_template_sha256", "tasks", "batches",
})

FONT_CN = "Microsoft YaHei"
BLUE = "175A8A"
DARK = "203746"
GRAY = "667085"
LIGHT_GRAY = "F3F4F6"
RED = "C00000"
DISPLAY_NAMES_SCHEMA_VERSION = "word-display-names-v1"


class ExamPointsError(RuntimeError):
    """A deterministic refusal caused by invalid or incomplete input."""


def read_json(path: Path) -> dict[str, Any]:
    try:
        value = json.loads(Path(path).read_text(encoding="utf-8"))
    except (OSError, UnicodeError, json.JSONDecodeError) as exc:
        raise ExamPointsError(f"无法读取有效JSON：{path}") from exc
    if not isinstance(value, dict):
        raise ExamPointsError(f"JSON顶层必须是对象：{path}")
    return value


def canonical_bytes(value: Any) -> bytes:
    return json.dumps(
        value, ensure_ascii=False, sort_keys=True, separators=(",", ":"), allow_nan=False,
    ).encode("utf-8")


def canonical_sha256(value: Any) -> str:
    return hashlib.sha256(canonical_bytes(value)).hexdigest()


def atomic_json(path: Path, value: dict[str, Any]) -> None:
    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    descriptor, temporary = tempfile.mkstemp(prefix=f".{path.name}.", dir=path.parent)
    try:
        with os.fdopen(descriptor, "wb") as handle:
            handle.write(canonical_bytes(value) + b"\n")
            handle.flush()
            os.fsync(handle.fileno())
        os.replace(temporary, path)
    finally:
        if os.path.exists(temporary):
            os.unlink(temporary)


def atomic_text(path: Path, value: str) -> None:
    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    descriptor, temporary = tempfile.mkstemp(prefix=f".{path.name}.", dir=path.parent)
    try:
        with os.fdopen(descriptor, "w", encoding="utf-8") as handle:
            handle.write(value)
            handle.flush()
            os.fsync(handle.fileno())
        os.replace(temporary, path)
    finally:
        if os.path.exists(temporary):
            os.unlink(temporary)


def valid_sha(value: Any, field: str) -> str:
    if not isinstance(value, str) or SHA_RE.fullmatch(value.lower()) is None:
        raise ExamPointsError(f"{field}必须是64位SHA-256")
    return value.lower()


def nonempty(value: Any, field: str) -> str:
    if (not isinstance(value, str) or not value or value != value.strip()
            or CONTROL_RE.search(value)):
        raise ExamPointsError(f"{field}必须是无控制字符的非空纯字符串")
    return value


def string_list(
    value: Any, field: str, *, allow_empty: bool = False, unique: bool = False,
) -> list[str]:
    if not isinstance(value, list) or (not value and not allow_empty):
        raise ExamPointsError(f"{field}必须是{'可空' if allow_empty else '非空'}字符串数组")
    output = [nonempty(item, f"{field}[{index}]") for index, item in enumerate(value)]
    if unique and len(output) != len(set(output)):
        raise ExamPointsError(f"{field}不得含重复项")
    return output


def evidence_string_list(value: Any, field: str) -> list[str]:
    """Validate immutable OCR/ASR display fragments without rewriting them.

    OCR engines can legitimately retain leading or trailing spaces in a fragment.
    ASR segments can likewise retain embedded line breaks.  Whitespace is normalized
    only by the Word rendering layer; task preparation must keep the authoritative
    evidence byte-for-byte unchanged.
    """
    if not isinstance(value, list) or not value:
        raise ExamPointsError(f"{field}必须是非空字符串数组")
    for index, item in enumerate(value):
        if not isinstance(item, str):
            raise ExamPointsError(f"{field}[{index}]必须是字符串")
    return value


def assert_no_generated_evidence(value: Any, location: str = "evidence") -> None:
    if isinstance(value, dict):
        for key, child in value.items():
            normalized = str(key).strip().casefold()
            if (str(key) in FORBIDDEN_EVIDENCE_KEYS
                    or normalized in FORBIDDEN_EVIDENCE_KEYS_CASEFOLD
                    or "gpt" in normalized):
                raise ExamPointsError(f"正式证据含旧GPT/考点正文：{location}.{key}")
            assert_no_generated_evidence(child, f"{location}.{key}")
    elif isinstance(value, list):
        for index, child in enumerate(value):
            assert_no_generated_evidence(child, f"{location}[{index}]")


def assert_authoritative_evidence_header(document: dict[str, Any], expected_assets: int) -> None:
    """Reject a debug, blocked, locally substituted or generated evidence corpus."""
    assert_no_generated_evidence(document)
    if (document.get("schema_version") != "first-principles-evidence/v1"
            or document.get("status") != "complete"
            or document.get("artifact_role") != "authoritative_evidence"
            or document.get("blockers") != []):
        raise ExamPointsError("只接受无blocker的authoritative complete evidence_corpus")
    policy = document.get("content_policy")
    if (not isinstance(policy, dict)
            or policy.get("asr_source") != "verified Xiaomi MiMo only"
            or policy.get("local_asr_substitution") is not False
            or policy.get("generated_prose_fields") is not False):
        raise ExamPointsError("正式证据必须声明仅使用经验证的小米MiMo ASR且无本地替代/生成正文")
    files = document.get("files")
    if (not isinstance(files, list) or len(files) != expected_assets
            or document.get("file_count") != expected_assets
            or document.get("asset_count") != expected_assets):
        raise ExamPointsError(f"正式证据必须恰好包含{expected_assets}个资产")


def refuse_incomplete_evidence_path(path: Path) -> None:
    if "incomplete" in Path(path).name.casefold():
        raise ExamPointsError("拒绝使用名称含incomplete的调试证据生成或验收正式考点")


def validate_manifest(
    document: dict[str, Any], *, expected_assets: int = ASSET_COUNT,
    expected_catalog: int = CATALOG_ENTRY_COUNT,
) -> dict[str, Any]:
    catalog = document.get("catalog")
    if not isinstance(catalog, dict):
        raise ExamPointsError("asset_manifest.catalog必须是对象")
    entries = catalog.get("entries")
    if (not isinstance(entries, list) or len(entries) != expected_catalog
            or catalog.get("entry_count") != expected_catalog
            or not all(isinstance(item, str) and item.strip() for item in entries)):
        raise ExamPointsError(f"课程目录必须恰好包含{expected_catalog}条非空原文")
    assets = document.get("assets")
    if (not isinstance(assets, list) or len(assets) != expected_assets
            or document.get("asset_count") != expected_assets):
        raise ExamPointsError(f"资产清单必须恰好包含{expected_assets}项")

    normalized: list[dict[str, Any]] = []
    seen_ids: set[str] = set()
    seen_shas: set[str] = set()
    previous: tuple[float, int] | None = None
    for order, item in enumerate(assets, 1):
        if not isinstance(item, dict):
            raise ExamPointsError(f"assets[{order}]必须是对象")
        asset_id = nonempty(item.get("asset_id"), f"assets[{order}].asset_id")
        source_sha = valid_sha(item.get("source_sha256"), f"{asset_id}.source_sha256")
        if asset_id in seen_ids or source_sha in seen_shas:
            raise ExamPointsError("asset_id或source_sha256重复")
        source_name = nonempty(item.get("source_name"), f"{asset_id}.source_name")
        year = item.get("year")
        kind = item.get("kind")
        position = item.get("catalog_position")
        suborder = item.get("catalog_suborder")
        anchor = item.get("catalog_anchor_position")
        coverage = item.get("coverage_positions")
        if not isinstance(year, int) or isinstance(year, bool) or not 2000 <= year <= 2100:
            raise ExamPointsError(f"{asset_id}.year无效")
        if kind not in {"video", "courseware"}:
            raise ExamPointsError(f"{asset_id}.kind无效")
        if (not isinstance(position, (int, float)) or isinstance(position, bool)
                or not 1 <= float(position) <= expected_catalog
                or not (float(position) * 2).is_integer()):
            raise ExamPointsError(f"{asset_id}.catalog_position无效")
        if not isinstance(suborder, int) or isinstance(suborder, bool) or suborder < 1:
            raise ExamPointsError(f"{asset_id}.catalog_suborder无效")
        if not isinstance(anchor, int) or isinstance(anchor, bool) or not 1 <= anchor <= expected_catalog:
            raise ExamPointsError(f"{asset_id}.catalog_anchor_position无效")
        if (not isinstance(coverage, list) or not coverage
                or any(not isinstance(value, int) or isinstance(value, bool)
                       or not 1 <= value <= expected_catalog for value in coverage)
                or coverage != sorted(set(coverage))):
            raise ExamPointsError(f"{asset_id}.coverage_positions无效")
        sort_key = (float(position), suborder)
        if previous is not None and sort_key < previous:
            raise ExamPointsError("asset_manifest.assets未保持目录顺序")
        previous = sort_key
        seen_ids.add(asset_id)
        seen_shas.add(source_sha)
        normalized.append({
            "order": order,
            "catalog_position": position,
            "catalog_suborder": suborder,
            "catalog_anchor_position": anchor,
            "catalog_entry": entries[anchor - 1],
            "coverage_positions": list(coverage),
            "asset_id": asset_id,
            "source_name": source_name,
            "source_sha256": source_sha,
            "year": year,
            "kind": kind,
            "manifest_record": item,
        })
    return {
        "catalog_entries": list(entries),
        "assets": normalized,
        "manifest_sha256": canonical_sha256(document),
    }


def validate_display_names(
    document: dict[str, Any], manifest: dict[str, Any],
) -> dict[str, dict[str, str]]:
    """Validate identity-bound titles used only for Word presentation."""
    if document.get("schema_version") != DISPLAY_NAMES_SCHEMA_VERSION:
        raise ExamPointsError("Word显示名配置schema_version不正确")
    entries = document.get("entries")
    if not isinstance(entries, list):
        raise ExamPointsError("Word显示名配置entries必须是数组")
    manifest_by_id = {item["asset_id"]: item for item in manifest["assets"]}
    required = {
        "asset_id", "source_sha256", "source_name", "catalog_position",
        "catalog_suborder", "display_title", "basis",
    }
    output: dict[str, dict[str, str]] = {}
    titles: set[str] = set()
    for number, entry in enumerate(entries, 1):
        if not isinstance(entry, dict) or set(entry) != required:
            raise ExamPointsError(f"Word显示名配置第{number}项字段不正确")
        asset_id = nonempty(entry.get("asset_id"), f"display[{number}].asset_id")
        if asset_id in output:
            raise ExamPointsError("Word显示名配置含重复asset_id")
        asset = manifest_by_id.get(asset_id)
        if asset is None:
            raise ExamPointsError(f"Word显示名配置引用未知资产：{asset_id}")
        if valid_sha(entry.get("source_sha256"), f"{asset_id}.display.sha256") != asset["source_sha256"]:
            raise ExamPointsError(f"Word显示名SHA绑定失败：{asset_id}")
        source_name = nonempty(entry.get("source_name"), f"{asset_id}.display.source_name")
        if source_name != asset["source_name"]:
            raise ExamPointsError(f"Word显示名源文件名绑定失败：{asset_id}")
        if (entry.get("catalog_position") != asset["catalog_position"]
                or entry.get("catalog_suborder") != asset["catalog_suborder"]):
            raise ExamPointsError(f"Word显示名目录坐标绑定失败：{asset_id}")
        title = nonempty(entry.get("display_title"), f"{asset_id}.display_title").strip()
        basis = nonempty(entry.get("basis"), f"{asset_id}.basis").strip()
        if CONTROL_RE.search(title + basis):
            raise ExamPointsError(f"Word显示名或依据含控制字符：{asset_id}")
        if title == source_name or title in titles:
            raise ExamPointsError(f"Word显示名无效或重复：{title}")
        titles.add(title)
        output[asset_id] = {"display_title": title, "basis": basis}
    return output


def apply_display_names(
    assets: list[dict[str, Any]], display_names: dict[str, dict[str, str]],
) -> list[dict[str, Any]]:
    """Return presentation copies; never mutate final JSON or source identity."""
    presented: list[dict[str, Any]] = []
    for asset in assets:
        copy = dict(asset)
        mapping = display_names.get(asset["asset_id"])
        copy["display_title"] = mapping["display_title"] if mapping else asset["source_name"]
        copy["display_name_applied"] = mapping is not None
        presented.append(copy)
    return presented


def validate_authoritative_evidence(
    path: Path, manifest: dict[str, Any], *, expected_assets: int = ASSET_COUNT,
    document: dict[str, Any] | None = None, header_validated: bool = False,
) -> dict[str, dict[str, Any]]:
    path = Path(path)
    refuse_incomplete_evidence_path(path)
    document = read_json(path) if document is None else document
    if not header_validated:
        assert_authoritative_evidence_header(document, expected_assets)
    files = document.get("files")

    output: dict[str, dict[str, Any]] = {}
    for order, (item, asset) in enumerate(zip(files, manifest["assets"]), 1):
        if not isinstance(item, dict):
            raise ExamPointsError(f"evidence.files[{order}]必须是对象")
        asset_id = asset["asset_id"]
        if item.get("asset_id") != asset_id:
            raise ExamPointsError(f"证据顺序或asset_id不一致：{asset_id}")
        if valid_sha(item.get("source_sha256"), f"{asset_id}.source_sha256") != asset["source_sha256"]:
            raise ExamPointsError(f"证据source_sha256绑定失败：{asset_id}")
        raw = item.get("raw_evidence")
        if not isinstance(raw, dict):
            raise ExamPointsError(f"{asset_id}.raw_evidence缺失")
        evidence_sha = valid_sha(item.get("evidence_sha256"), f"{asset_id}.evidence_sha256")
        if canonical_sha256(raw) != evidence_sha:
            raise ExamPointsError(f"evidence_sha256重算不一致：{asset_id}")
        identity_values = {
            "asset_id": asset_id,
            "source_sha256": asset["source_sha256"],
            "source_name": asset["source_name"],
            "year": asset["year"],
            "kind": asset["kind"],
            "catalog_position": asset["catalog_position"],
            "catalog_suborder": asset["catalog_suborder"],
        }
        for field, expected in identity_values.items():
            if raw.get(field) != expected:
                raise ExamPointsError(f"raw_evidence.{field}绑定失败：{asset_id}")
        if raw.get("asset_manifest_record") != asset["manifest_record"]:
            raise ExamPointsError(f"raw_evidence.asset_manifest_record绑定失败：{asset_id}")
        if not isinstance(raw.get("asr_result"), dict) or not isinstance(raw.get("ocr_result"), dict):
            raise ExamPointsError(f"{asset_id}缺少完整ASR/OCR结构")
        sections = item.get("sections")
        if not isinstance(sections, dict) or set(sections) != {"ASR 结果", "OCR 结果"}:
            raise ExamPointsError(f"{asset_id}.sections只能含ASR结果和OCR结果")
        for title in ("ASR 结果", "OCR 结果"):
            evidence_string_list(sections.get(title), f"{asset_id}.{title}")
        output[asset_id] = {
            "record": item,
            "raw_evidence": raw,
            "evidence_sha256": evidence_sha,
        }
    return output


def _greedy_assignment(weights: list[int], batch_count: int) -> list[int]:
    if not isinstance(batch_count, int) or batch_count < 1:
        raise ExamPointsError("batch_count必须是正整数")
    totals = [0] * batch_count
    assignment = [0] * len(weights)
    for index in sorted(range(len(weights)), key=lambda i: (-weights[i], i)):
        target = min(range(batch_count), key=lambda batch: (totals[batch], batch))
        assignment[index] = target + 1
        totals[target] += weights[index]
    return assignment


def _write_task_directory(path: Path, tasks: list[tuple[str, dict[str, Any]]], index: dict[str, Any]) -> None:
    path = Path(path).resolve()
    if path.exists():
        raise ExamPointsError(f"任务目录已存在，拒绝覆盖：{path}")
    path.parent.mkdir(parents=True, exist_ok=True)
    staged = Path(tempfile.mkdtemp(prefix=f".{path.name}.", dir=path.parent))
    try:
        for filename, task in tasks:
            (staged / filename).write_bytes(canonical_bytes(task) + b"\n")
        (staged / "index.json").write_bytes(canonical_bytes(index) + b"\n")
        os.replace(staged, path)
    finally:
        if staged.exists():
            shutil.rmtree(staged)


def prepare_tasks(
    *, manifest_path: Path = DEFAULT_MANIFEST, evidence_path: Path = DEFAULT_EVIDENCE,
    output_dir: Path | None = DEFAULT_TASK_DIR, spec_path: Path = DEFAULT_SPEC,
    template_path: Path = DEFAULT_TEMPLATE, expected_assets: int = ASSET_COUNT,
    expected_catalog: int = CATALOG_ENTRY_COUNT, batch_count: int = 3,
) -> dict[str, Any]:
    refuse_incomplete_evidence_path(evidence_path)
    manifest_document = read_json(manifest_path)
    manifest = validate_manifest(
        manifest_document, expected_assets=expected_assets, expected_catalog=expected_catalog,
    )
    evidence_document = read_json(evidence_path)
    evidence = validate_authoritative_evidence(
        evidence_path, manifest, expected_assets=expected_assets, document=evidence_document,
    )
    try:
        spec_payload = Path(spec_path).read_bytes()
        template_payload = Path(template_path).read_bytes()
    except OSError as exc:
        raise ExamPointsError("无法读取考点规范或结构模板") from exc
    weights = []
    for asset in manifest["assets"]:
        raw = evidence[asset["asset_id"]]["raw_evidence"]
        weights.append(len(canonical_bytes(raw["asr_result"])) + len(canonical_bytes(raw["ocr_result"])))
    assignments = _greedy_assignment(weights, batch_count)

    tasks: list[tuple[str, dict[str, Any]]] = []
    index_tasks: list[dict[str, Any]] = []
    for asset, weight, batch in zip(manifest["assets"], weights, assignments):
        asset_id = asset["asset_id"]
        evidence_item = evidence[asset_id]
        batch_id = f"EP-BATCH-{batch:02d}"
        identity = {
            "order": asset["order"],
            "batch_id": batch_id,
            "catalog_position": asset["catalog_position"],
            "catalog_suborder": asset["catalog_suborder"],
            "catalog_entry": asset["catalog_entry"],
            "coverage_positions": asset["coverage_positions"],
            "asset_id": asset_id,
            "source_name": asset["source_name"],
            "source_sha256": asset["source_sha256"],
            "evidence_sha256": evidence_item["evidence_sha256"],
        }
        task = {
            "identity": identity,
            "ASR 结果": evidence_item["raw_evidence"]["asr_result"],
            "OCR 结果": evidence_item["raw_evidence"]["ocr_result"],
            "task_constraints": {
                "objective": "从考官视角穷举本资产全部可命题细节，不作摘要或重要性筛选",
                "authoring_method": "GPT直接完整阅读本资产ASR+OCR后撰写",
                "automatic_point_generation_allowed": False,
                "old_gpt_fusion_allowed": False,
                "file_level_final_review_required": True,
                "must_check": list(MUST_CHECK),
                "evidence_locator_format": "指向/raw_evidence/asr_result或/raw_evidence/ocr_result子项的可解析JSON Pointer",
                "delivery_spec_sha256": hashlib.sha256(spec_payload).hexdigest(),
                "structure_template_sha256": hashlib.sha256(template_payload).hexdigest(),
                "required_result_keys": sorted(RESULT_KEYS),
            },
        }
        if set(task) != set(TASK_FILE_KEYS):
            raise AssertionError("exam-point task schema changed")
        task_sha = canonical_sha256(task)
        filename = f"{asset['order']:04d}-{asset['source_sha256'][:16]}.json"
        tasks.append((filename, task))
        index_tasks.append({
            **identity,
            "task_file": filename,
            "task_sha256": task_sha,
            "evidence_character_count": weight,
        })
    batches = []
    for batch in range(1, batch_count + 1):
        members = [item for item in index_tasks if item["batch_id"] == f"EP-BATCH-{batch:02d}"]
        batches.append({
            "batch_id": f"EP-BATCH-{batch:02d}",
            "task_count": len(members),
            "evidence_character_count": sum(item["evidence_character_count"] for item in members),
            "orders": [item["order"] for item in members],
        })
    index = {
        "schema_version": TASK_SCHEMA,
        "status": "ready",
        "file_count": expected_assets,
        "catalog_entry_count": expected_catalog,
        "batch_count": batch_count,
        "asset_manifest_sha256": manifest["manifest_sha256"],
        "evidence_corpus_sha256": canonical_sha256(evidence_document),
        "delivery_spec_sha256": hashlib.sha256(spec_payload).hexdigest(),
        "structure_template_sha256": hashlib.sha256(template_payload).hexdigest(),
        "tasks": index_tasks,
        "batches": batches,
    }
    if output_dir is not None:
        _write_task_directory(output_dir, tasks, index)
    return {"index": index, "task_files": tasks}


def load_task_index(task_dir: Path, *, expected_assets: int = ASSET_COUNT) -> dict[str, Any]:
    task_dir = Path(task_dir)
    index = read_json(task_dir / "index.json")
    tasks = index.get("tasks")
    if (set(index) != INDEX_ROOT_KEYS
            or index.get("schema_version") != TASK_SCHEMA or index.get("status") != "ready"
            or index.get("file_count") != expected_assets
            or not isinstance(tasks, list) or len(tasks) != expected_assets):
        raise ExamPointsError(f"考点任务索引必须为ready且恰好{expected_assets}项")
    for field in (
        "asset_manifest_sha256", "evidence_corpus_sha256", "delivery_spec_sha256",
        "structure_template_sha256",
    ):
        valid_sha(index.get(field), f"task_index.{field}")
    batch_count = index.get("batch_count")
    batches = index.get("batches")
    if (not isinstance(batch_count, int) or isinstance(batch_count, bool) or batch_count < 1
            or not isinstance(batches, list) or len(batches) != batch_count):
        raise ExamPointsError("考点任务批次数量无效")
    valid_batch_ids = {f"EP-BATCH-{number:02d}" for number in range(1, batch_count + 1)}
    seen_ids: set[str] = set()
    seen_files: set[str] = set()
    previous_order = 0
    for item in tasks:
        if not isinstance(item, dict):
            raise ExamPointsError("考点任务索引项必须是对象")
        if set(item) != INDEX_TASK_KEYS:
            raise ExamPointsError("考点任务索引项字段不符合严格schema")
        order = item.get("order")
        asset_id = item.get("asset_id")
        filename = item.get("task_file")
        if order != previous_order + 1:
            raise ExamPointsError("考点任务order必须从1连续递增")
        previous_order = order
        if not isinstance(asset_id, str) or not asset_id or asset_id in seen_ids:
            raise ExamPointsError("考点任务asset_id缺失或重复")
        if item.get("batch_id") not in valid_batch_ids:
            raise ExamPointsError(f"考点任务batch_id无效：{asset_id}")
        if (not isinstance(item.get("evidence_character_count"), int)
                or item["evidence_character_count"] < 1):
            raise ExamPointsError(f"考点任务证据字符数无效：{asset_id}")
        if (not isinstance(filename, str) or Path(filename).name != filename
                or filename in seen_files or not filename.endswith(".json")):
            raise ExamPointsError(f"考点task_file无效：{asset_id}")
        task = read_json(task_dir / filename)
        if set(task) != set(TASK_FILE_KEYS):
            raise ExamPointsError(f"考点任务文件字段不符合契约：{asset_id}")
        task_sha = valid_sha(item.get("task_sha256"), f"{asset_id}.task_sha256")
        if canonical_sha256(task) != task_sha:
            raise ExamPointsError(f"考点任务文件SHA绑定失败：{asset_id}")
        if task.get("identity") != {key: item[key] for key in (
            "order", "batch_id", "catalog_position", "catalog_suborder", "catalog_entry",
            "coverage_positions", "asset_id", "source_name", "source_sha256", "evidence_sha256",
        )}:
            raise ExamPointsError(f"考点任务identity绑定失败：{asset_id}")
        valid_sha(item.get("source_sha256"), f"{asset_id}.source_sha256")
        valid_sha(item.get("evidence_sha256"), f"{asset_id}.evidence_sha256")
        seen_ids.add(asset_id)
        seen_files.add(filename)
    seen_batches: set[str] = set()
    for batch in batches:
        if not isinstance(batch, dict) or set(batch) != INDEX_BATCH_KEYS:
            raise ExamPointsError("考点任务batches字段不符合严格schema")
        batch_id = batch.get("batch_id")
        if batch_id not in valid_batch_ids or batch_id in seen_batches:
            raise ExamPointsError("考点任务batches含无效或重复batch_id")
        members = [item for item in tasks if item["batch_id"] == batch_id]
        if (batch.get("task_count") != len(members)
                or batch.get("orders") != [item["order"] for item in members]
                or batch.get("evidence_character_count")
                != sum(item["evidence_character_count"] for item in members)):
            raise ExamPointsError(f"{batch_id}批次汇总与逐任务身份不一致")
        seen_batches.add(batch_id)
    if seen_batches != valid_batch_ids:
        raise ExamPointsError("考点任务批次集合不完整")
    return index


def validate_task_evidence_binding(
    task_dir: Path, index: dict[str, Any], evidence_records: dict[str, dict[str, Any]],
) -> None:
    """Prove that every task still carries the exact authoritative ASR/OCR object."""
    for item in index["tasks"]:
        asset_id = item["asset_id"]
        evidence_item = evidence_records.get(asset_id)
        if evidence_item is None:
            raise ExamPointsError(f"考点任务找不到正式证据：{asset_id}")
        task = read_json(Path(task_dir) / item["task_file"])
        raw = evidence_item.get("raw_evidence")
        if not isinstance(raw, dict):
            raise ExamPointsError(f"考点任务对应的raw_evidence缺失：{asset_id}")
        if task.get("ASR 结果") != raw.get("asr_result"):
            raise ExamPointsError(f"考点任务ASR与正式证据不逐字一致：{asset_id}")
        if task.get("OCR 结果") != raw.get("ocr_result"):
            raise ExamPointsError(f"考点任务OCR与正式证据不逐字一致：{asset_id}")
        constraints = task.get("task_constraints")
        if (not isinstance(constraints, dict) or set(constraints) != TASK_CONSTRAINT_KEYS
                or constraints.get("objective") != "从考官视角穷举本资产全部可命题细节，不作摘要或重要性筛选"
                or constraints.get("automatic_point_generation_allowed") is not False
                or constraints.get("old_gpt_fusion_allowed") is not False
                or constraints.get("file_level_final_review_required") is not True
                or constraints.get("authoring_method") != "GPT直接完整阅读本资产ASR+OCR后撰写"
                or constraints.get("must_check") != list(MUST_CHECK)
                or constraints.get("evidence_locator_format")
                != "指向/raw_evidence/asr_result或/raw_evidence/ocr_result子项的可解析JSON Pointer"
                or constraints.get("delivery_spec_sha256") != index["delivery_spec_sha256"]
                or constraints.get("structure_template_sha256") != index["structure_template_sha256"]
                or constraints.get("required_result_keys") != sorted(RESULT_KEYS)):
            raise ExamPointsError(f"考点任务撰写约束无效：{asset_id}")


def _resolve_json_pointer(document: Any, pointer: str) -> Any:
    if not pointer.startswith("/"):
        raise ExamPointsError(f"证据定位必须是JSON Pointer：{pointer}")
    current = document
    for raw_token in pointer[1:].split("/"):
        token = raw_token.replace("~1", "/").replace("~0", "~")
        if isinstance(current, dict):
            if token not in current:
                raise ExamPointsError(f"证据定位不存在：{pointer}")
            current = current[token]
        elif isinstance(current, list):
            if not token.isdigit() or int(token) >= len(current):
                raise ExamPointsError(f"证据定位不存在：{pointer}")
            current = current[int(token)]
        else:
            raise ExamPointsError(f"证据定位穿过标量值：{pointer}")
    return current


def _validate_result(
    result: dict[str, Any], task: dict[str, Any], evidence_record: dict[str, Any],
) -> dict[str, Any]:
    asset_id = task["asset_id"]
    if set(result) != RESULT_KEYS:
        raise ExamPointsError(f"{asset_id}结果字段必须严格符合schema")
    for field in (
        "order", "batch_id", "catalog_position", "catalog_suborder", "catalog_entry",
        "coverage_positions", "asset_id", "source_name", "source_sha256", "evidence_sha256",
    ):
        if result.get(field) != task[field]:
            raise ExamPointsError(f"{asset_id}.{field}任务身份绑定失败")
    if valid_sha(result.get("task_sha256"), f"{asset_id}.task_sha256") != task["task_sha256"]:
        raise ExamPointsError(f"{asset_id}.task_sha256绑定失败")
    model = nonempty(result.get("model"), f"{asset_id}.model")
    limitations = string_list(
        result.get("source_limitations"), f"{asset_id}.source_limitations", allow_empty=True,
        unique=True,
    )
    raw_ocr = evidence_record.get("raw_evidence", {}).get("ocr_result", {})
    courseware_record = raw_ocr.get("courseware_record") if isinstance(raw_ocr, dict) else None
    source_risks = courseware_record.get("risks", []) if isinstance(courseware_record, dict) else []
    if (isinstance(source_risks, list)
            and any(isinstance(risk, dict) and risk.get("code") == "missing_external_asset"
                    for risk in source_risks)
            and not limitations):
        raise ExamPointsError(f"{asset_id}存在缺失外链附件，source_limitations不得为空")
    review = result.get("dimension_review")
    if not isinstance(review, dict) or set(review) != set(DIMENSION_KEYS):
        raise ExamPointsError(f"{asset_id}.dimension_review必须且只能包含固定12个维度")
    if any(value not in REVIEW_VALUES for value in review.values()):
        raise ExamPointsError(f"{asset_id}.dimension_review值无效")

    points = result.get("points")
    if not isinstance(points, list) or not points:
        raise ExamPointsError(f"{asset_id}.points不能为空，禁止以空结果兜底")
    validated_points: list[dict[str, Any]] = []
    conflict_points = 0
    for index, point in enumerate(points, 1):
        if not isinstance(point, dict) or set(point) != POINT_KEYS:
            raise ExamPointsError(f"{asset_id}.points[{index}]字段不符合schema")
        expected_id = f"A{task['order']:03d}-KP{index:03d}"
        if point.get("point_id") != expected_id:
            raise ExamPointsError(f"{asset_id}.point_id必须连续为{expected_id}")
        dimensions = string_list(point.get("dimension"), f"{expected_id}.dimension", unique=True)
        if any(value not in DIMENSION_LABELS for value in dimensions):
            raise ExamPointsError(f"{expected_id}.dimension含未知命题维度")
        statement = nonempty(point.get("statement"), f"{expected_id}.statement")
        question_angles = string_list(
            point.get("question_angles"), f"{expected_id}.question_angles", unique=True,
        )
        answer_elements = string_list(
            point.get("answer_elements"), f"{expected_id}.answer_elements", unique=True,
        )
        traps = string_list(
            point.get("traps"), f"{expected_id}.traps", allow_empty=True, unique=True,
        )
        status = point.get("evidence_status")
        if status not in EVIDENCE_STATUSES:
            raise ExamPointsError(f"{expected_id}.evidence_status无效")
        refs = point.get("evidence_refs")
        if not isinstance(refs, list) or not refs:
            raise ExamPointsError(f"{expected_id}.evidence_refs不能为空")
        validated_refs = []
        modalities: set[str] = set()
        for ref_index, ref in enumerate(refs, 1):
            if not isinstance(ref, dict) or set(ref) != REFERENCE_KEYS:
                raise ExamPointsError(f"{expected_id}.evidence_refs[{ref_index}]字段无效")
            modality = ref.get("modality")
            locator = nonempty(ref.get("locator"), f"{expected_id}.locator")
            support = ref.get("support")
            if modality not in {"ASR", "OCR"} or support not in SUPPORT_VALUES:
                raise ExamPointsError(f"{expected_id}.evidence_ref枚举值无效")
            prefix = "/raw_evidence/asr_result" if modality == "ASR" else "/raw_evidence/ocr_result"
            if not locator.startswith(prefix + "/"):
                raise ExamPointsError(f"{expected_id}.{modality}定位未指向对应证据分区")
            _resolve_json_pointer(evidence_record, locator)
            modalities.add(modality)
            validated_refs.append(dict(ref))
        if status in {"ASR+OCR互证", "证据冲突待复核"} and modalities != {"ASR", "OCR"}:
            raise ExamPointsError(f"{expected_id}.{status}必须同时引用ASR和OCR")
        if status == "ASR单证" and modalities != {"ASR"}:
            raise ExamPointsError(f"{expected_id}.ASR单证只能引用ASR")
        if status == "OCR单证" and modalities != {"OCR"}:
            raise ExamPointsError(f"{expected_id}.OCR单证只能引用OCR")
        if status == "证据冲突待复核":
            conflict_points += 1
        validated_points.append({
            "point_id": expected_id,
            "dimension": dimensions,
            "statement": statement,
            "question_angles": question_angles,
            "answer_elements": answer_elements,
            "traps": traps,
            "evidence_refs": validated_refs,
            "evidence_status": status,
        })

    asset_review = result.get("asset_review")
    if not isinstance(asset_review, dict) or set(asset_review) != ASSET_REVIEW_KEYS:
        raise ExamPointsError(f"{asset_id}.asset_review字段无效")
    for field in ASSET_REVIEW_KEYS - {"unresolved_conflicts"}:
        if asset_review.get(field) is not True:
            raise ExamPointsError(f"{asset_id}.asset_review.{field}必须为true")
    conflicts = string_list(
        asset_review.get("unresolved_conflicts"), f"{asset_id}.unresolved_conflicts",
        allow_empty=True, unique=True,
    )
    if bool(conflict_points) != bool(conflicts):
        raise ExamPointsError(f"{asset_id}冲突考点与unresolved_conflicts不一致")
    return {
        **{field: result[field] for field in RESULT_KEYS if field not in {
            "model", "source_limitations", "dimension_review", "points", "asset_review",
        }},
        "model": model,
        "source_limitations": limitations,
        "dimension_review": dict(review),
        "points": validated_points,
        "asset_review": {**asset_review, "unresolved_conflicts": conflicts},
    }


def _result_files(result_dir: Path) -> list[Path]:
    # Parallel model batches may write to EP-BATCH-XX subdirectories.
    return sorted(Path(result_dir).rglob("*.json")) if Path(result_dir).is_dir() else []


def merge_results(
    *, task_dir: Path = DEFAULT_TASK_DIR, result_dir: Path = DEFAULT_RESULT_DIR,
    evidence_path: Path = DEFAULT_EVIDENCE, output_path: Path | None = DEFAULT_FINAL,
    expected_assets: int = ASSET_COUNT,
) -> dict[str, Any]:
    index = load_task_index(task_dir, expected_assets=expected_assets)
    # Revalidate authoritative status and corpus digest before accepting any result.
    refuse_incomplete_evidence_path(evidence_path)
    evidence_document = read_json(evidence_path)
    assert_authoritative_evidence_header(evidence_document, expected_assets)
    if canonical_sha256(evidence_document) != index.get("evidence_corpus_sha256"):
        raise ExamPointsError("正式证据已变化，与考点任务索引不一致")
    manifest_like = {
        "assets": [
            {
                "asset_id": item["asset_id"], "source_sha256": item["source_sha256"],
                "source_name": item["source_name"], "year": (
                    evidence_document["files"][item["order"] - 1]["raw_evidence"]["year"]),
                "kind": evidence_document["files"][item["order"] - 1]["raw_evidence"]["kind"],
                "catalog_position": item["catalog_position"],
                "catalog_suborder": item["catalog_suborder"],
                "catalog_anchor_position": (
                    evidence_document["files"][item["order"] - 1]["raw_evidence"]
                    ["asset_manifest_record"]["catalog_anchor_position"]),
                "coverage_positions": item["coverage_positions"],
                "manifest_record": evidence_document["files"][item["order"] - 1]
                ["raw_evidence"]["asset_manifest_record"],
            }
            for item in index["tasks"]
        ]
    }
    evidence = validate_authoritative_evidence(
        evidence_path, manifest_like, expected_assets=expected_assets,
        document=evidence_document, header_validated=True,
    )
    validate_task_evidence_binding(task_dir, index, evidence)
    paths = _result_files(result_dir)
    if len(paths) != expected_assets:
        raise ExamPointsError(f"正式考点结果必须恰好{expected_assets}份，实际{len(paths)}")
    raw_results: dict[str, dict[str, Any]] = {}
    for path in paths:
        result = read_json(path)
        asset_id = result.get("asset_id")
        if not isinstance(asset_id, str) or not asset_id or asset_id in raw_results:
            raise ExamPointsError("考点结果asset_id缺失或重复")
        raw_results[asset_id] = result
    expected_ids = {item["asset_id"] for item in index["tasks"]}
    if set(raw_results) != expected_ids:
        raise ExamPointsError("考点结果资产集与任务索引不一致")

    assets: list[dict[str, Any]] = []
    models: set[str] = set()
    for task in index["tasks"]:
        validated = _validate_result(
            raw_results[task["asset_id"]], task, evidence[task["asset_id"]]["record"],
        )
        models.add(validated["model"])
        assets.append(validated)
    if len(models) != 1:
        raise ExamPointsError(f"91份考点结果的model必须一致，实际{sorted(models)}")
    final = {
        "schema_version": FINAL_SCHEMA,
        "status": "complete",
        "basis": {
            "asset_manifest_sha256": index["asset_manifest_sha256"],
            "evidence_corpus_sha256": index["evidence_corpus_sha256"],
            "delivery_spec_sha256": index["delivery_spec_sha256"],
            "structure_template_sha256": index["structure_template_sha256"],
            "old_gpt_fusion_used": False,
            "automatic_point_generation_used": False,
            "authoring_method": "GPT逐资产直接阅读完整ASR+OCR证据并经文件级复核",
        },
        "model": next(iter(models)),
        "catalog_entry_count": index["catalog_entry_count"],
        "asset_count": expected_assets,
        "completed_asset_count": expected_assets,
        "point_count": sum(len(item["points"]) for item in assets),
        "assets": assets,
    }
    if output_path is not None:
        atomic_json(output_path, final)
    return final


def build_progress(
    *, task_dir: Path = DEFAULT_TASK_DIR, result_dir: Path = DEFAULT_RESULT_DIR,
    evidence_path: Path = DEFAULT_EVIDENCE, output_json: Path | None = DEFAULT_PROGRESS_JSON,
    output_markdown: Path | None = DEFAULT_PROGRESS_MD, expected_assets: int = ASSET_COUNT,
) -> dict[str, Any]:
    index = load_task_index(task_dir, expected_assets=expected_assets)
    refuse_incomplete_evidence_path(evidence_path)
    evidence_document = read_json(evidence_path)
    assert_authoritative_evidence_header(evidence_document, expected_assets)
    if canonical_sha256(evidence_document) != index.get("evidence_corpus_sha256"):
        raise ExamPointsError("正式证据已变化，与考点任务索引不一致")
    evidence_records = {
        item["asset_id"]: item for item in evidence_document.get("files", [])
        if isinstance(item, dict) and isinstance(item.get("asset_id"), str)
    }
    validate_task_evidence_binding(
        task_dir, index,
        {asset_id: {"raw_evidence": item["raw_evidence"]}
         for asset_id, item in evidence_records.items()},
    )
    by_asset: dict[str, list[tuple[Path, dict[str, Any]]]] = defaultdict(list)
    unreadable: list[str] = []
    for path in _result_files(result_dir):
        try:
            raw = read_json(path)
        except ExamPointsError:
            unreadable.append(path.name)
            continue
        asset_id = raw.get("asset_id")
        if isinstance(asset_id, str):
            by_asset[asset_id].append((path, raw))
        else:
            unreadable.append(path.name)
    rows = []
    completed = 0
    point_count = 0
    for task in index["tasks"]:
        matches = by_asset.get(task["asset_id"], [])
        status = "已绑定"
        error = None
        points = 0
        if len(matches) > 1:
            status, error = "结果无效", "同一asset_id存在多个结果文件"
        elif len(matches) == 1:
            try:
                validated = _validate_result(
                    matches[0][1], task, evidence_records[task["asset_id"]],
                )
                status = "QA通过"
                points = len(validated["points"])
                completed += 1
                point_count += points
            except (ExamPointsError, KeyError) as exc:
                status, error = "结果无效", str(exc)
        rows.append({
            "order": task["order"], "batch_id": task["batch_id"],
            "catalog_position": task["catalog_position"], "asset_id": task["asset_id"],
            "source_name": task["source_name"], "source_sha256": task["source_sha256"],
            "evidence_sha256": task["evidence_sha256"], "status": status,
            "point_count": points, "error": error,
        })
    progress = {
        "schema_version": PROGRESS_SCHEMA,
        "status": "complete" if completed == expected_assets else "in_progress",
        "asset_count": expected_assets,
        "bound_count": expected_assets,
        "qa_passed_count": completed,
        "point_count": point_count,
        "progress_bar": (
            "[" + "#" * ((completed * 20) // expected_assets)
            + "-" * (20 - ((completed * 20) // expected_assets))
            + f"] {completed}/{expected_assets}"
        ),
        "batches": [
            {
                "batch_id": batch["batch_id"],
                "asset_count": batch["task_count"],
                "qa_passed_count": sum(
                    row["status"] == "QA通过" for row in rows
                    if row["batch_id"] == batch["batch_id"]
                ),
                "point_count": sum(
                    row["point_count"] for row in rows
                    if row["batch_id"] == batch["batch_id"]
                ),
            }
            for batch in index["batches"]
        ],
        "unreadable_or_unbound_result_files": unreadable,
        "assets": rows,
    }
    if output_json is not None:
        atomic_json(output_json, progress)
    if output_markdown is not None:
        lines = [
            "# 全量考点执行进度", "",
            f"- 已绑定正式证据：{expected_assets}/{expected_assets}",
            f"- QA通过：{completed}/{expected_assets}",
            f"- 进度条：{progress['progress_bar']}",
            f"- 累计考点：{point_count}", "",
            "| 顺序 | 批次 | 目录位 | 文件 | evidence_sha256 | 状态 | 考点数 | 错误 |",
            "|---:|---|---:|---|---|---|---:|---|",
        ]
        for row in rows:
            safe_name = row["source_name"].replace("|", "\\|")
            safe_error = (row["error"] or "—").replace("|", "\\|").replace("\n", " ")
            lines.append(
                f"| {row['order']:03d} | {row['batch_id']} | {row['catalog_position']} | "
                f"{safe_name} | `{row['evidence_sha256']}` | {row['status']} | "
                f"{row['point_count']} | {safe_error} |"
            )
        atomic_text(output_markdown, "\n".join(lines) + "\n")
    return progress


def validate_final(
    final: dict[str, Any], index: dict[str, Any], evidence_document: dict[str, Any],
    *, expected_assets: int = ASSET_COUNT,
) -> list[dict[str, Any]]:
    expected_root = {
        "schema_version", "status", "basis", "model", "catalog_entry_count",
        "asset_count", "completed_asset_count", "point_count", "assets",
    }
    if set(final) != expected_root or final.get("schema_version") != FINAL_SCHEMA or final.get("status") != "complete":
        raise ExamPointsError("正式考点母本顶层schema无效")
    if (final.get("asset_count") != expected_assets
            or final.get("completed_asset_count") != expected_assets
            or final.get("catalog_entry_count") != index.get("catalog_entry_count")):
        raise ExamPointsError("正式考点母本数量字段无效")
    basis = final.get("basis")
    if (not isinstance(basis, dict) or set(basis) != BASIS_KEYS
            or basis.get("asset_manifest_sha256") != index.get("asset_manifest_sha256")
            or basis.get("evidence_corpus_sha256") != index.get("evidence_corpus_sha256")
            or basis.get("delivery_spec_sha256") != index.get("delivery_spec_sha256")
            or basis.get("structure_template_sha256") != index.get("structure_template_sha256")
            or basis.get("old_gpt_fusion_used") is not False
            or basis.get("automatic_point_generation_used") is not False
            or basis.get("authoring_method") != "GPT逐资产直接阅读完整ASR+OCR证据并经文件级复核"):
        raise ExamPointsError("正式考点母本basis身份或人工模型撰写声明无效")
    assert_authoritative_evidence_header(evidence_document, expected_assets)
    if canonical_sha256(evidence_document) != index.get("evidence_corpus_sha256"):
        raise ExamPointsError("正式证据与任务索引SHA不一致")
    assets = final.get("assets")
    if not isinstance(assets, list) or len(assets) != expected_assets:
        raise ExamPointsError("正式考点母本assets数量无效")
    evidence_by_id = {item["asset_id"]: item for item in evidence_document["files"]}
    models: set[str] = set()
    validated = []
    for result, task in zip(assets, index["tasks"]):
        item = _validate_result(result, task, evidence_by_id[task["asset_id"]])
        models.add(item["model"])
        validated.append(item)
    if models != {final.get("model")}:
        raise ExamPointsError("正式考点母本model字段与逐资产结果不一致")
    if final.get("point_count") != sum(len(item["points"]) for item in validated):
        raise ExamPointsError("正式考点母本point_count与实际不一致")
    if any(pattern.search(canonical_bytes(final)) for pattern in SECRET_PATTERNS):
        raise ExamPointsError("正式考点母本含疑似API密钥或授权信息")
    return validated


def _font(style: Any, size: float, *, bold: bool = False, color: str | None = None) -> None:
    style.font.name = FONT_CN
    style.font.size = Pt(size)
    style.font.bold = bold
    if color:
        style.font.color.rgb = RGBColor.from_string(color)
    fonts = style.element.get_or_add_rPr().get_or_add_rFonts()
    fonts.set(qn("w:eastAsia"), FONT_CN)
    fonts.set(qn("w:ascii"), "Aptos")
    fonts.set(qn("w:hAnsi"), "Aptos")


def _setup_word(document: Document) -> None:
    section = document.sections[0]
    section.page_width, section.page_height = Cm(21), Cm(29.7)
    section.top_margin, section.bottom_margin = Cm(1.7), Cm(1.6)
    section.left_margin, section.right_margin = Cm(1.9), Cm(1.8)
    _font(document.styles["Normal"], 9.5, color=DARK)
    document.styles["Normal"].paragraph_format.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE
    for name, size, color in (
        ("Title", 24, BLUE), ("Heading 1", 15, DARK),
        ("Heading 2", 11.5, BLUE), ("Subtitle", 9.5, GRAY),
    ):
        _font(document.styles[name], size, bold=name != "Subtitle", color=color)
        document.styles[name].paragraph_format.keep_with_next = True
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    footer.add_run("第 ")
    field = OxmlElement("w:fldSimple")
    field.set(qn("w:instr"), "PAGE")
    footer._p.append(field)
    footer.add_run(" 页")


def _shade(cell: Any, color: str) -> None:
    node = OxmlElement("w:shd")
    node.set(qn("w:fill"), color)
    cell._tc.get_or_add_tcPr().append(node)


def _catalog_assets(final_assets: list[dict[str, Any]]) -> dict[int, list[str]]:
    output: dict[int, list[str]] = defaultdict(list)
    for asset in final_assets:
        label = f"A{asset['order']:03d} {asset.get('display_title', asset['source_name'])}"
        for position in asset["coverage_positions"]:
            output[position].append(label)
    return output


def _body_rows(assets: list[dict[str, Any]]) -> list[tuple[str, str]]:
    rows: list[tuple[str, str]] = []
    for asset in assets:
        rows.append(("Heading 1", f"A{asset['order']:03d} [{asset.get('display_title', asset['source_name'])}]"))
        if asset.get("display_name_applied"):
            rows.append(("Subtitle", f"源文件名：{asset['source_name']}（仅保留用于证据追溯）"))
        rows.append(("Subtitle", f"目录位置：{asset['catalog_position']}；覆盖目录：{'、'.join(map(str, asset['coverage_positions']))}"))
        rows.append(("Subtitle", f"asset_id：{asset['asset_id']}"))
        rows.append(("Subtitle", f"source_sha256：{asset['source_sha256']}"))
        rows.append(("Subtitle", f"evidence_sha256：{asset['evidence_sha256']}"))
        limitations = "；".join(asset["source_limitations"]) if asset["source_limitations"] else "无"
        rows.append(("Normal", f"来源限制：{limitations}"))
        for point in asset["points"]:
            rows.append(("Heading 2", f"{point['point_id']}｜{'、'.join(point['dimension'])}"))
            rows.append(("Normal", f"考点陈述：{point['statement']}"))
            rows.append(("Normal", f"可考方式：{'；'.join(point['question_angles'])}"))
            rows.append(("Normal", f"标准答案要素：{'；'.join(point['answer_elements'])}"))
            rows.append(("Normal", f"易错点：{'；'.join(point['traps']) if point['traps'] else '无'}"))
            rows.append(("Normal", f"证据状态：{point['evidence_status']}"))
            refs = "；".join(f"{ref['modality']} {ref['locator']} ({ref['support']})" for ref in point["evidence_refs"])
            rows.append(("Normal", f"证据定位：{refs}"))
    return rows


def export_word(
    *, manifest_path: Path = DEFAULT_MANIFEST, task_dir: Path = DEFAULT_TASK_DIR,
    evidence_path: Path = DEFAULT_EVIDENCE, final_path: Path = DEFAULT_FINAL,
    output_path: Path = DEFAULT_WORD, expected_assets: int = ASSET_COUNT,
    expected_catalog: int = CATALOG_ENTRY_COUNT,
    display_names_path: Path | None = None,
) -> Path:
    refuse_incomplete_evidence_path(evidence_path)
    manifest = validate_manifest(
        read_json(manifest_path), expected_assets=expected_assets, expected_catalog=expected_catalog,
    )
    index = load_task_index(task_dir, expected_assets=expected_assets)
    evidence_document = read_json(evidence_path)
    validate_task_evidence_binding(
        task_dir, index,
        {item["asset_id"]: {"raw_evidence": item["raw_evidence"]}
         for item in evidence_document.get("files", []) if isinstance(item, dict)
         and isinstance(item.get("asset_id"), str) and isinstance(item.get("raw_evidence"), dict)},
    )
    assets = validate_final(
        read_json(final_path), index, evidence_document, expected_assets=expected_assets,
    )
    display_names = (
        validate_display_names(read_json(display_names_path), manifest)
        if display_names_path is not None else {}
    )
    assets = apply_display_names(assets, display_names)
    if index.get("asset_manifest_sha256") != manifest["manifest_sha256"]:
        raise ExamPointsError("Word资产清单与考点任务索引SHA不一致")
    output_path = Path(output_path)
    if output_path.suffix.lower() != ".docx":
        raise ExamPointsError("独立考点输出必须是.docx")
    if "ASR-OCR-GPT" in output_path.name.upper():
        raise ExamPointsError("考点导出器拒绝写入主ASR/OCR/GPT Word文件")

    document = Document()
    _setup_word(document)
    document.core_properties.title = "2026年YW第一课全量考点"
    document.add_heading("2026年YW第一课全量考点", 0).alignment = WD_ALIGN_PARAGRAPH.CENTER
    document.add_paragraph(
        f"共 {expected_catalog} 条课程目录、{expected_assets} 个课程资产；考点正文由GPT逐资产直接阅读完整ASR与OCR后撰写。",
        style="Subtitle",
    ).alignment = WD_ALIGN_PARAGRAPH.CENTER
    document.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
    document.add_heading("简要目录", 1)
    table = document.add_table(rows=1, cols=3)
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    for cell, label in zip(table.rows[0].cells, ("序号", "课程目录原文", "对应考点资产")):
        cell.text = label
        _shade(cell, BLUE)
        for run in cell.paragraphs[0].runs:
            run.bold = True
            run.font.color.rgb = RGBColor(255, 255, 255)
    by_position = _catalog_assets(assets)
    for position, entry in enumerate(manifest["catalog_entries"], 1):
        cells = table.add_row().cells
        values = (str(position), entry, "；".join(by_position[position]) or "无对应源资产")
        for cell, value in zip(cells, values):
            cell.text = value
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        if position % 2 == 0:
            for cell in cells:
                _shade(cell, LIGHT_GRAY)
    for style, value in _body_rows(assets):
        paragraph = document.add_paragraph(value, style=style)
        if style == "Heading 1":
            paragraph.paragraph_format.page_break_before = True
        elif style == "Normal" and value.startswith("考点陈述："):
            paragraph.paragraph_format.first_line_indent = Cm(0.7)
            paragraph.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    output_path.parent.mkdir(parents=True, exist_ok=True)
    temporary = output_path.with_suffix(".tmp.docx")
    document.save(temporary)
    temporary.replace(output_path)
    return output_path


def qa_delivery(
    *, manifest_path: Path = DEFAULT_MANIFEST, task_dir: Path = DEFAULT_TASK_DIR,
    evidence_path: Path = DEFAULT_EVIDENCE, final_path: Path = DEFAULT_FINAL,
    word_path: Path = DEFAULT_WORD, report_path: Path | None = DEFAULT_QA,
    expected_assets: int = ASSET_COUNT, expected_catalog: int = CATALOG_ENTRY_COUNT,
    allowed_uncovered: frozenset[int] | None = None,
    display_names_path: Path | None = None,
) -> dict[str, Any]:
    refuse_incomplete_evidence_path(evidence_path)
    manifest = validate_manifest(
        read_json(manifest_path), expected_assets=expected_assets, expected_catalog=expected_catalog,
    )
    index = load_task_index(task_dir, expected_assets=expected_assets)
    evidence_document = read_json(evidence_path)
    validate_task_evidence_binding(
        task_dir, index,
        {item["asset_id"]: {"raw_evidence": item["raw_evidence"]}
         for item in evidence_document.get("files", []) if isinstance(item, dict)
         and isinstance(item.get("asset_id"), str) and isinstance(item.get("raw_evidence"), dict)},
    )
    final_document = read_json(final_path)
    assets = validate_final(final_document, index, evidence_document, expected_assets=expected_assets)
    display_names = (
        validate_display_names(read_json(display_names_path), manifest)
        if display_names_path is not None else {}
    )
    assets = apply_display_names(assets, display_names)
    if index.get("asset_manifest_sha256") != manifest["manifest_sha256"]:
        raise ExamPointsError("QA资产清单与考点任务索引SHA不一致")
    covered = {position for asset in assets for position in asset["coverage_positions"]}
    uncovered = set(range(1, expected_catalog + 1)) - covered
    if allowed_uncovered is None:
        allowed_uncovered = (KNOWN_UNCOVERED_CATALOG_POSITIONS
                             if expected_assets == ASSET_COUNT and expected_catalog == CATALOG_ENTRY_COUNT
                             else frozenset())
    if uncovered != set(allowed_uncovered):
        raise ExamPointsError(f"课程目录覆盖缺口异常：{sorted(uncovered)}")

    word_path = Path(word_path).resolve(strict=True)
    document = Document(word_path)
    if len(document.tables) != 1 or len(document.tables[0].rows) != expected_catalog + 1:
        raise ExamPointsError("独立考点Word必须含一张完整课程目录表")
    if tuple(cell.text for cell in document.tables[0].rows[0].cells) != (
        "序号", "课程目录原文", "对应考点资产",
    ):
        raise ExamPointsError("独立考点Word目录表头不一致")
    by_position = _catalog_assets(assets)
    for position, row in enumerate(document.tables[0].rows[1:], 1):
        expected = (
            str(position), manifest["catalog_entries"][position - 1],
            "；".join(by_position[position]) or "无对应源资产",
        )
        if tuple(cell.text for cell in row.cells) != expected:
            raise ExamPointsError(f"独立考点Word目录第{position}行不一致")
    expected_rows = _body_rows(assets)
    first_heading = f"A{assets[0]['order']:03d} [{assets[0].get('display_title', assets[0]['source_name'])}]"
    start = next((i for i, p in enumerate(document.paragraphs) if p.text == first_heading), None)
    if start is None:
        raise ExamPointsError("独立考点Word缺少首个资产标题")
    expected_prefix = [
        ("Title", "2026年YW第一课全量考点"),
        ("Subtitle", f"共 {expected_catalog} 条课程目录、{expected_assets} 个课程资产；考点正文由GPT逐资产直接阅读完整ASR与OCR后撰写。"),
        ("Heading 1", "简要目录"),
    ]
    actual_prefix = [(p.style.name, p.text) for p in document.paragraphs[:start] if p.text]
    if actual_prefix != expected_prefix:
        raise ExamPointsError("独立考点Word封面或目录标题存在增删改")
    actual_rows = [(p.style.name, p.text) for p in document.paragraphs[start:] if p.text]
    if actual_rows != expected_rows:
        raise ExamPointsError("独立考点Word正文与正式JSON不逐字一致或顺序错误")
    for paragraph in document.paragraphs:
        for run in paragraph.runs:
            if run.text and run.font.color.rgb is not None and str(run.font.color.rgb) == RED:
                raise ExamPointsError("独立考点Word不得混用主Word的红色GPT正文样式")
    for table in document.tables:
        for row in table.rows:
            for cell in row.cells:
                for paragraph in cell.paragraphs:
                    for run in paragraph.runs:
                        if run.text and run.font.color.rgb is not None and str(run.font.color.rgb) == RED:
                            raise ExamPointsError("独立考点Word目录不得出现红色")
    for section in document.sections:
        for container in (section.header, section.footer):
            for paragraph in container.paragraphs:
                for run in paragraph.runs:
                    if run.text and run.font.color.rgb is not None and str(run.font.color.rgb) == RED:
                        raise ExamPointsError("独立考点Word页眉页脚不得出现红色")
    with ZipFile(word_path) as archive:
        xml = b"\n".join(archive.read(name) for name in archive.namelist() if name.endswith(".xml"))
    if any(pattern.search(xml) for pattern in SECRET_PATTERNS):
        raise ExamPointsError("独立考点Word含疑似API密钥或授权信息")
    report = {
        "schema_version": "exam-points-qa/v1",
        "passed": True,
        "word": str(word_path),
        "asset_count": len(assets),
        "catalog_entry_count": expected_catalog,
        "point_count": sum(len(item["points"]) for item in assets),
        "evidence_reference_count": sum(
            len(point["evidence_refs"]) for asset in assets for point in asset["points"]
        ),
        "uncovered_catalog_positions": sorted(uncovered),
        "allowed_uncovered_catalog_positions": sorted(allowed_uncovered),
        "identity_bindings_verified": True,
        "evidence_pointers_resolved": True,
        "model_authorship_declarations_verified": True,
        "word_verbatim": True,
        "word_separate_from_main_delivery": "ASR-OCR-GPT" not in word_path.name.upper(),
        "red_style_absent": True,
        "credential_patterns_absent": True,
    }
    if report_path is not None:
        atomic_json(report_path, report)
    return report


def parse_args(argv: list[str] | None = None) -> argparse.Namespace:
    parser = argparse.ArgumentParser(description=__doc__)
    sub = parser.add_subparsers(dest="command", required=True)
    prepare = sub.add_parser("prepare", help="从正式完整证据建立身份绑定任务；不生成考点正文")
    prepare.add_argument("--manifest", type=Path, default=DEFAULT_MANIFEST)
    prepare.add_argument("--evidence", type=Path, default=DEFAULT_EVIDENCE)
    prepare.add_argument("--output-dir", type=Path, default=DEFAULT_TASK_DIR)
    prepare.add_argument("--batch-count", type=int, default=3)
    progress = sub.add_parser("progress", help="只读校验已有结果并更新进度元数据")
    progress.add_argument("--task-dir", type=Path, default=DEFAULT_TASK_DIR)
    progress.add_argument("--results", type=Path, default=DEFAULT_RESULT_DIR)
    progress.add_argument("--evidence", type=Path, default=DEFAULT_EVIDENCE)
    progress.add_argument("--json", type=Path, default=DEFAULT_PROGRESS_JSON)
    progress.add_argument("--markdown", type=Path, default=DEFAULT_PROGRESS_MD)
    merge = sub.add_parser("merge", help="严格校验并机械合并91份GPT考点结果")
    merge.add_argument("--task-dir", type=Path, default=DEFAULT_TASK_DIR)
    merge.add_argument("--results", type=Path, default=DEFAULT_RESULT_DIR)
    merge.add_argument("--evidence", type=Path, default=DEFAULT_EVIDENCE)
    merge.add_argument("--output", type=Path, default=DEFAULT_FINAL)
    export = sub.add_parser("export", help="导出独立考点Word；不修改正文")
    export.add_argument("--manifest", type=Path, default=DEFAULT_MANIFEST)
    export.add_argument("--task-dir", type=Path, default=DEFAULT_TASK_DIR)
    export.add_argument("--evidence", type=Path, default=DEFAULT_EVIDENCE)
    export.add_argument("--final", type=Path, default=DEFAULT_FINAL)
    export.add_argument("--display-names", type=Path, default=DEFAULT_DISPLAY_NAMES)
    export.add_argument("--output", type=Path, default=DEFAULT_WORD)
    qa = sub.add_parser("qa", help="对正式JSON和独立Word执行只读QA")
    qa.add_argument("--manifest", type=Path, default=DEFAULT_MANIFEST)
    qa.add_argument("--task-dir", type=Path, default=DEFAULT_TASK_DIR)
    qa.add_argument("--evidence", type=Path, default=DEFAULT_EVIDENCE)
    qa.add_argument("--final", type=Path, default=DEFAULT_FINAL)
    qa.add_argument("--display-names", type=Path, default=DEFAULT_DISPLAY_NAMES)
    qa.add_argument("--word", type=Path, default=DEFAULT_WORD)
    qa.add_argument("--report", type=Path, default=DEFAULT_QA)
    return parser.parse_args(argv)


def main(argv: list[str] | None = None) -> int:
    args = parse_args(argv)
    try:
        if args.command == "prepare":
            result = prepare_tasks(
                manifest_path=args.manifest, evidence_path=args.evidence,
                output_dir=args.output_dir, batch_count=args.batch_count,
            )["index"]
            payload = {"status": "ready", "file_count": result["file_count"],
                       "batch_count": result["batch_count"], "output": str(args.output_dir)}
        elif args.command == "progress":
            payload = build_progress(
                task_dir=args.task_dir, result_dir=args.results, evidence_path=args.evidence,
                output_json=args.json, output_markdown=args.markdown,
            )
        elif args.command == "merge":
            result = merge_results(
                task_dir=args.task_dir, result_dir=args.results,
                evidence_path=args.evidence, output_path=args.output,
            )
            payload = {"status": result["status"], "asset_count": result["asset_count"],
                       "point_count": result["point_count"], "output": str(args.output)}
        elif args.command == "export":
            output = export_word(
                manifest_path=args.manifest, task_dir=args.task_dir,
                evidence_path=args.evidence, final_path=args.final, output_path=args.output,
                display_names_path=args.display_names,
            )
            payload = {"status": "exported", "output": str(output)}
        else:
            payload = qa_delivery(
                manifest_path=args.manifest, task_dir=args.task_dir,
                evidence_path=args.evidence, final_path=args.final,
                word_path=args.word, report_path=args.report,
                display_names_path=args.display_names,
            )
        print(json.dumps(payload, ensure_ascii=False, indent=2))
        return 0
    except (ExamPointsError, OSError, KeyError) as exc:
        print(json.dumps({"status": "refused", "reason": str(exc)}, ensure_ascii=False), file=sys.stderr)
        return 2


if __name__ == "__main__":
    raise SystemExit(main())
