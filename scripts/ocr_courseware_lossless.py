#!/usr/bin/env python3
"""Evidence-first OCR for every courseware leaf in the 2025/2026 corpus.

Unlike the legacy courseware extractor, this module deliberately keeps OCR
responses byte-for-byte at the text-field level: no spelling correction, no
normalisation, no confidence threshold, no de-duplication and no short-line
filter.  Native OOXML/PDF text is retained separately and is never labelled as
OCR.  Every rendered image is persisted with a SHA-256 digest so a later Word
builder can prove which pixels produced each OCR record.

The default invocation discovers the 20 directly supplied courseware files and
the nine already CRC-validated PPTX members of ``项目管理1-9.rar``.  Existing
LibreOffice PDF renders are reused only after their page count has been checked.
"""

from __future__ import annotations

import argparse
import hashlib
import importlib.metadata
import io
import json
import math
import os
import posixpath
import shutil
import subprocess
import sys
import tempfile
import time
import urllib.parse
import xml.etree.ElementTree as ET
from datetime import datetime
from pathlib import Path, PurePosixPath
from typing import Any, Callable, Iterable
from zipfile import BadZipFile, ZipFile

import cv2
import numpy as np
import pymupdf
from PIL import Image, ImageDraw, ImageFont, ImageSequence


ROOT = Path(__file__).resolve().parent.parent
PROCESSOR_VERSION = "courseware-ocr-lossless-v2"
EMBEDDED_INVENTORY_VERSION = "package-members-v1"
SCHEMA_VERSION = "1.0"
XLSX_RENDER_DEGRADED_STATUS = "native_complete_render_failed_full_page_ocr_covered"
XLSX_RENDER_DEGRADED_RISK = "embedded_xlsx_render_failed_full_page_ocr_covered"
SUPPORTED = {".pdf", ".png", ".jpg", ".jpeg", ".pptx", ".docx"}
RASTER = {".png", ".jpg", ".jpeg", ".bmp", ".tif", ".tiff", ".webp", ".wdp"}
VECTOR = {".emf", ".wmf", ".svg"}
ANIMATED = {".gif"}
AUDIO_VIDEO = {
    ".mp4", ".mov", ".wmv", ".avi", ".mkv", ".webm", ".m4a", ".mp3",
    ".wav", ".aac", ".flac", ".wma", ".ogg",
}
REL_NS = "http://schemas.openxmlformats.org/package/2006/relationships"
OFFICE_REL_NS = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
KNOWN_TRAILING_RENDER_FALLBACKS = {
    "南网在线“速捷电”品牌系列功能建设介绍.pptx": (48, 47),
}


class LosslessOCRError(RuntimeError):
    """A condition that must never be converted into an empty success."""


def sha256_bytes(payload: bytes) -> str:
    return hashlib.sha256(payload).hexdigest()


def sha256_file(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for block in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(block)
    return digest.hexdigest()


def atomic_bytes(path: Path, payload: bytes) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    descriptor, temporary = tempfile.mkstemp(prefix=f".{path.name}.", dir=path.parent)
    try:
        with os.fdopen(descriptor, "wb") as handle:
            handle.write(payload)
            handle.flush()
            os.fsync(handle.fileno())
        os.replace(temporary, path)
    finally:
        if os.path.exists(temporary):
            os.unlink(temporary)


def atomic_json(path: Path, value: Any) -> None:
    atomic_bytes(
        path,
        (json.dumps(value, ensure_ascii=False, indent=2, allow_nan=False) + "\n").encode("utf-8"),
    )


def local_name(value: str) -> str:
    return value.rsplit("}", 1)[-1]


def relationship_part(part: str) -> str:
    parent, name = posixpath.split(part)
    return posixpath.join(parent, "_rels", name + ".rels")


def resolve_part(part: str, target: str) -> str:
    decoded = urllib.parse.unquote(target).replace("\\", "/")
    return posixpath.normpath(posixpath.join(posixpath.dirname(part), decoded)).lstrip("/")


def relationships(archive: ZipFile, part: str) -> list[dict[str, Any]]:
    rel_part = relationship_part(part)
    try:
        root = ET.fromstring(archive.read(rel_part))
    except KeyError:
        return []
    except ET.ParseError as exc:
        raise LosslessOCRError(f"关系文件损坏：{rel_part}: {exc}") from exc
    output = []
    for node in root:
        if local_name(node.tag) != "Relationship":
            continue
        target = node.attrib.get("Target", "")
        external = node.attrib.get("TargetMode") == "External"
        rel_type = node.attrib.get("Type", "").rsplit("/", 1)[-1]
        output.append({
            "id": node.attrib.get("Id"),
            "type": rel_type,
            "target": target,
            "external": external,
            "resolved_part": None if external else resolve_part(part, target),
            "relationship_part": rel_part,
        })
    return output


def xml_native_text(archive: ZipFile, part: str, *, kind: str) -> list[dict[str, Any]]:
    """Return raw text fragments and accessibility strings without filtering."""
    try:
        root = ET.fromstring(archive.read(part))
    except KeyError:
        return []
    except ET.ParseError as exc:
        raise LosslessOCRError(f"XML 部件损坏：{part}: {exc}") from exc
    output: list[dict[str, Any]] = []
    sequence = 0
    for node in root.iter():
        tag = local_name(node.tag)
        # Chart caches store categories and values in c:v rather than a:t.
        if (tag == "t" or (kind == "chart" and tag == "v")) and node.text is not None:
            output.append({
                "index": sequence, "text": node.text, "kind": kind,
                "source_part": part,
            })
            sequence += 1
        for attribute, value in node.attrib.items():
            if local_name(attribute) in {"descr", "title"}:
                output.append({
                    "index": sequence, "text": value, "kind": "accessibility",
                    "source_part": part, "attribute": local_name(attribute),
                })
                sequence += 1
    return output


def _normalise_box(box: Any) -> list[list[float]]:
    try:
        points = [[float(point[0]), float(point[1])] for point in box]
    except (TypeError, ValueError, IndexError) as exc:
        raise LosslessOCRError("OCR 返回了非法坐标框。") from exc
    if len(points) != 4 or any(
            len(point) != 2 or not all(math.isfinite(value) for value in point)
            for point in points):
        raise LosslessOCRError("OCR 返回了非法四点坐标框。")
    return points


def recognize_raw(engine: Any, image: np.ndarray) -> list[dict[str, Any]]:
    """Preserve every engine row, including duplicates, one-char and low score."""
    result = engine(image, text_score=0.0)
    if result is None or not all(hasattr(result, key) for key in ("txts", "scores", "boxes")):
        raise LosslessOCRError("RapidOCR 返回结构不完整。")
    texts, scores, boxes = result.txts, result.scores, result.boxes
    if texts is None and scores is None and boxes is None:
        return []
    if any(value is None for value in (texts, scores, boxes)):
        raise LosslessOCRError("RapidOCR 只返回了部分字段。")
    if not (len(texts) == len(scores) == len(boxes)):
        raise LosslessOCRError("RapidOCR 文本、分数和坐标数量不一致。")
    output = []
    for index, (text, score, box) in enumerate(zip(texts, scores, boxes)):
        if not isinstance(text, str):
            raise LosslessOCRError("RapidOCR 文本字段不是字符串。")
        confidence = float(score)
        if not math.isfinite(confidence) or confidence < 0 or confidence > 1:
            raise LosslessOCRError("RapidOCR 置信度非法。")
        output.append({
            "index": index,
            "text": text,
            "confidence": confidence,
            "box": _normalise_box(box),
        })
    return output


def encode_png(image: np.ndarray) -> bytes:
    success, encoded = cv2.imencode(".png", image)
    if not success:
        raise LosslessOCRError("无法保存 OCR 像素证据。")
    return encoded.tobytes()


def evidence_record(
    engine: Any,
    image: np.ndarray,
    path: Path,
    *,
    page: int | None,
    part: int,
    part_count: int,
    render_sha256: str | None,
    render_mode: str,
) -> dict[str, Any]:
    payload = encode_png(image)
    atomic_bytes(path, payload)
    return {
        "page": page,
        "part": part,
        "part_count": part_count,
        "render_mode": render_mode,
        "render_sha256": render_sha256,
        "evidence_path": str(path.resolve()),
        "image_sha256": sha256_bytes(payload),
        "width": int(image.shape[1]),
        "height": int(image.shape[0]),
        "raw_ocr_lines": recognize_raw(engine, image),
        "native_lines": [],
    }


def slice_array(image: np.ndarray) -> list[np.ndarray]:
    height, width = image.shape[:2]
    slice_height = max(1600, round(width * 1.45))
    if height <= slice_height:
        return [image]
    overlap = min(160, slice_height // 12)
    output = []
    start = 0
    while start < height:
        end = min(height, start + slice_height)
        output.append(image[start:end])
        if end == height:
            break
        start = end - overlap
    return output


def pdf_page_images(path: Path, zoom: float) -> Iterable[tuple[int, int, int, np.ndarray]]:
    with pymupdf.open(path) as document:
        for page_index, page in enumerate(document, 1):
            embedded = page.get_images(full=True)
            if embedded:
                largest = max(embedded, key=lambda item: int(item[2]) * int(item[3]))
                if int(largest[3]) / max(1, int(largest[2])) > 2.5:
                    payload = document.extract_image(largest[0])["image"]
                    decoded = cv2.imdecode(np.frombuffer(payload, np.uint8), cv2.IMREAD_COLOR)
                    if decoded is not None:
                        chunks = slice_array(decoded)
                        for part, chunk in enumerate(chunks, 1):
                            yield page_index, part, len(chunks), chunk
                        continue
            page_zoom = zoom
            if page.rect.height / max(1.0, page.rect.width) > 2.5:
                page_zoom = max(zoom, min(14.0, 2400.0 / max(1.0, page.rect.width)))
            pixmap = page.get_pixmap(
                matrix=pymupdf.Matrix(page_zoom, page_zoom), alpha=False,
                colorspace=pymupdf.csRGB,
            )
            rgb = np.frombuffer(pixmap.samples, np.uint8).reshape(pixmap.height, pixmap.width, 3)
            chunks = slice_array(cv2.cvtColor(rgb, cv2.COLOR_RGB2BGR))
            for part, chunk in enumerate(chunks, 1):
                yield page_index, part, len(chunks), chunk


def pdf_full_page_images(path: Path, zoom: float) -> Iterable[tuple[int, np.ndarray]]:
    """Render office-exported PDF pages without poster/image special casing."""
    with pymupdf.open(path) as document:
        matrix = pymupdf.Matrix(zoom, zoom)
        for page_index, page in enumerate(document, 1):
            pixmap = page.get_pixmap(matrix=matrix, alpha=False, colorspace=pymupdf.csRGB)
            rgb = np.frombuffer(pixmap.samples, np.uint8).reshape(pixmap.height, pixmap.width, 3)
            yield page_index, cv2.cvtColor(rgb, cv2.COLOR_RGB2BGR)


def pdf_native_pages(path: Path) -> dict[int, list[dict[str, Any]]]:
    output: dict[int, list[dict[str, Any]]] = {}
    with pymupdf.open(path) as document:
        for page_index, page in enumerate(document, 1):
            output[page_index] = [
                {"index": index, "text": line, "kind": "pdf_text", "source_part": f"page:{page_index}"}
                for index, line in enumerate(page.get_text("text").splitlines())
            ]
    return output


def image_slices(path: Path) -> Iterable[tuple[int, int, int, np.ndarray]]:
    decoded = cv2.imdecode(np.frombuffer(path.read_bytes(), np.uint8), cv2.IMREAD_COLOR)
    if decoded is None:
        raise LosslessOCRError(f"无法解码图片：{path}")
    chunks = slice_array(decoded)
    for part, chunk in enumerate(chunks, 1):
        yield 1, part, len(chunks), chunk


def presentation_slide_parts(archive: ZipFile) -> list[str]:
    root = ET.fromstring(archive.read("ppt/presentation.xml"))
    rels = {item["id"]: item for item in relationships(archive, "ppt/presentation.xml")}
    output = []
    for node in root.iter():
        if local_name(node.tag) != "sldId":
            continue
        rel_id = node.attrib.get(f"{{{OFFICE_REL_NS}}}id")
        relation = rels.get(rel_id)
        if relation is None or relation["external"] or relation["type"] != "slide":
            raise LosslessOCRError(f"PPT 幻灯片关系缺失：{rel_id}")
        output.append(relation["resolved_part"])
    if not output:
        raise LosslessOCRError("PPT 没有幻灯片。")
    return output


def _file_like_external(target: str) -> bool:
    if target == "NULL":
        return True
    parsed = urllib.parse.urlsplit(target)
    if parsed.scheme.lower() in {"http", "https", "mailto", "ftp"}:
        return False
    return bool(PurePosixPath(urllib.parse.unquote(target).replace("\\", "/")).suffix)


def pptx_inventory(path: Path) -> dict[str, Any]:
    """Map native text, media, attachments and external links to slide numbers."""
    with ZipFile(path) as archive:
        names = set(archive.namelist())
        slides = presentation_slide_parts(archive)
        native_by_page: dict[int, list[dict[str, Any]]] = {}
        objects: dict[str, dict[str, Any]] = {}
        external: list[dict[str, Any]] = []
        native_traversable = {
            "chart", "diagramData", "comments", "notesSlide", "slideLayout", "slideMaster",
        }
        relationship_traversable = set(native_traversable)

        for page_number, slide_part in enumerate(slides, 1):
            native = xml_native_text(archive, slide_part, kind="slide_text")
            queue = [slide_part]
            visited: set[str] = set()
            while queue:
                part = queue.pop(0)
                if part in visited:
                    continue
                visited.add(part)
                for relation in relationships(archive, part):
                    if relation["external"]:
                        item = {
                            "slide": page_number,
                            "source_part": part,
                            "relationship_id": relation["id"],
                            "relationship_type": relation["type"],
                            "target": relation["target"],
                            "file_like": _file_like_external(relation["target"]),
                        }
                        external.append(item)
                        continue
                    target = relation["resolved_part"]
                    if target not in names:
                        external.append({
                            "slide": page_number,
                            "source_part": part,
                            "relationship_id": relation["id"],
                            "relationship_type": relation["type"],
                            "target": relation["target"],
                            "resolved_part": target,
                            "file_like": True,
                            "broken_internal": True,
                        })
                        continue
                    if target.startswith("ppt/media/") or target.startswith("ppt/embeddings/"):
                        entry = objects.setdefault(target, {
                            "member": target,
                            "member_sha256": sha256_bytes(archive.read(target)),
                            "size_bytes": len(archive.read(target)),
                            "extension": Path(target).suffix.lower(),
                            "relationship_types": set(),
                            "source_pages": set(),
                        })
                        entry["source_pages"].add(page_number)
                        entry["relationship_types"].add(relation["type"])
                    if relation["type"] in relationship_traversable:
                        if relation["type"] in native_traversable:
                            native.extend(xml_native_text(
                                archive, target,
                                kind="notes" if relation["type"] == "notesSlide" else relation["type"],
                            ))
                        queue.append(target)
            native_by_page[page_number] = native

        # Preserve package members even when their relationship is orphaned.
        # Such a member has no defensible page locator, which is represented by
        # an empty source_pages list and later elevated to an explicit risk.
        for member in sorted(
                name for name in names
                if (name.startswith("ppt/media/") or name.startswith("ppt/embeddings/"))
                and not name.endswith("/")):
            if member in objects:
                continue
            payload = archive.read(member)
            objects[member] = {
                "member": member,
                "member_sha256": sha256_bytes(payload),
                "size_bytes": len(payload),
                "extension": Path(member).suffix.lower(),
                "relationship_types": set(),
                "source_pages": set(),
            }

        # Broken master/layout relationships do not belong to only one slide,
        # but are still evidence about a damaged source package.
        for rel_name in sorted(name for name in names if name.endswith(".rels")):
            try:
                root = ET.fromstring(archive.read(rel_name))
            except ET.ParseError as exc:
                raise LosslessOCRError(f"关系文件损坏：{rel_name}: {exc}") from exc
            for node in root:
                if node.attrib.get("TargetMode") != "External":
                    continue
                target = node.attrib.get("Target", "")
                if target != "NULL":
                    continue
                marker = (rel_name, node.attrib.get("Id"), target)
                if any((item.get("relationship_part"), item.get("relationship_id"), item["target"]) == marker
                       for item in external):
                    continue
                external.append({
                    "slide": None,
                    "relationship_part": rel_name,
                    "relationship_id": node.attrib.get("Id"),
                    "relationship_type": node.attrib.get("Type", "").rsplit("/", 1)[-1],
                    "target": target,
                    "file_like": True,
                })

        serialised_objects = []
        for value in objects.values():
            serialised_objects.append({
                **value,
                "source_pages": sorted(value["source_pages"]),
                "relationship_types": sorted(value["relationship_types"]),
            })
        return {
            "slide_parts": slides,
            "native_by_page": native_by_page,
            "objects": sorted(serialised_objects, key=lambda item: item["member"]),
            "external_relationships": external,
        }


def docx_inventory(path: Path) -> dict[str, Any]:
    with ZipFile(path) as archive:
        native = []
        objects = []
        external = []
        for name in sorted(archive.namelist()):
            if name.startswith("word/") and name.endswith(".xml"):
                native.extend(xml_native_text(archive, name, kind=PurePosixPath(name).stem))
            if name.startswith("word/media/") and not name.endswith("/"):
                payload = archive.read(name)
                objects.append({
                    "member": name, "member_sha256": sha256_bytes(payload),
                    "size_bytes": len(payload), "extension": Path(name).suffix.lower(),
                    "source_pages": [], "relationship_types": ["image"],
                })
            if name.endswith(".rels"):
                try:
                    root = ET.fromstring(archive.read(name))
                except ET.ParseError as exc:
                    raise LosslessOCRError(f"关系文件损坏：{name}: {exc}") from exc
                for node in root:
                    if node.attrib.get("TargetMode") == "External":
                        target = node.attrib.get("Target", "")
                        external.append({
                            "slide": None, "relationship_part": name,
                            "relationship_id": node.attrib.get("Id"),
                            "relationship_type": node.attrib.get("Type", "").rsplit("/", 1)[-1],
                            "target": target, "file_like": _file_like_external(target),
                        })
        return {"native_document": native, "objects": objects,
                "external_relationships": external}


def find_cached_render(source: Path, roots: Iterable[Path]) -> Path | None:
    prefix = sha256_file(source)[:16]
    expected = f"input-{prefix}.pdf"
    for root in roots:
        if not root.is_dir():
            continue
        candidates = sorted(root.rglob(expected))
        for candidate in candidates:
            sibling = candidate.with_suffix(source.suffix.lower())
            if sibling.is_file() and sha256_file(sibling) != sha256_file(source):
                continue
            return candidate.resolve()
    return None


def run_soffice(source: Path, directory: Path, soffice: Path) -> Path:
    directory.mkdir(parents=True, exist_ok=True)
    digest = sha256_file(source)[:16]
    alias = directory / f"input-{digest}{source.suffix.lower()}"
    if not alias.exists():
        try:
            os.link(source, alias)
        except OSError:
            shutil.copyfile(source, alias)
    destination = alias.with_suffix(".pdf")
    if destination.is_file():
        return destination
    profile = directory / "lo-profile"
    profile.mkdir(exist_ok=True)
    environment = os.environ.copy()
    environment["SOFFICE_TASK_TMPDIR"] = str((directory / "tmp").resolve())
    command = [
        str(soffice), "--headless", "--nologo", "--nodefault", "--nofirststartwizard",
        f"-env:UserInstallation={profile.resolve().as_uri()}",
        "--convert-to", "pdf", "--outdir", str(directory), str(alias),
    ]
    completed = subprocess.run(
        command, stdout=subprocess.PIPE, stderr=subprocess.STDOUT, text=True,
        timeout=900, env=environment,
    )
    if completed.returncode or not destination.is_file():
        raise LosslessOCRError(f"LibreOffice 渲染失败：{completed.stdout.strip()}")
    return destination


def _font(size: int) -> ImageFont.FreeTypeFont | ImageFont.ImageFont:
    candidates = [
        Path("/usr/share/fonts/truetype/wqy/wqy-zenhei.ttc"),
        Path("/usr/share/fonts/opentype/noto/NotoSansCJK-Regular.ttc"),
    ]
    for candidate in candidates:
        if candidate.is_file():
            return ImageFont.truetype(str(candidate), size=size)
    return ImageFont.load_default()


def native_text_fallback(native: list[dict[str, Any]]) -> np.ndarray:
    """Render a missing text-only slide as auditable fallback evidence."""
    canvas = Image.new("RGB", (1920, 1080), "white")
    draw = ImageDraw.Draw(canvas)
    font = _font(48)
    x, y = 90, 90
    # Accessibility descriptions and speaker notes remain in native_lines but
    # must not be painted as if they had been visibly present on the slide.
    for item in (value for value in native if value.get("kind") == "slide_text"):
        text = item.get("text", "")
        if not text:
            continue
        # Preserve the exact native string in JSON; wrapping only affects pixels.
        chunks = [text[index:index + 30] for index in range(0, len(text), 30)] or [""]
        for chunk in chunks:
            draw.text((x, y), chunk, fill="black", font=font)
            y += 70
            if y > 990:
                break
        if y > 990:
            break
    return cv2.cvtColor(np.asarray(canvas), cv2.COLOR_RGB2BGR)


def validate_ppt_render_plan(source_name: str, source_pages: int, rendered_pages: int) -> list[str]:
    if rendered_pages == source_pages:
        return ["rendered"] * source_pages
    allowed = KNOWN_TRAILING_RENDER_FALLBACKS.get(source_name)
    if allowed == (source_pages, rendered_pages):
        return ["rendered"] * rendered_pages + ["native_text_fallback"]
    raise LosslessOCRError(
        f"PPT 源页数与渲染页数不一致：{source_name}: {source_pages}/{rendered_pages}"
    )


def _decode_member_frames(payload: bytes, extension: str) -> list[np.ndarray]:
    if extension == ".wdp":
        import imagecodecs

        decoded = np.asarray(imagecodecs.jpegxr_decode(payload))
        if decoded.ndim == 2:
            return [cv2.cvtColor(decoded, cv2.COLOR_GRAY2BGR)]
        if decoded.ndim == 3 and decoded.shape[2] == 3:
            return [cv2.cvtColor(decoded, cv2.COLOR_RGB2BGR)]
        if decoded.ndim == 3 and decoded.shape[2] == 4:
            return [cv2.cvtColor(decoded, cv2.COLOR_RGBA2BGRA)]
        raise LosslessOCRError(f"JPEG-XR/WDP 解码结果维度非法：{decoded.shape}")
    if extension in ANIMATED:
        image = Image.open(io.BytesIO(payload))
        return [
            cv2.cvtColor(np.asarray(frame.convert("RGB")), cv2.COLOR_RGB2BGR)
            for frame in ImageSequence.Iterator(image)
        ]
    decoded = cv2.imdecode(np.frombuffer(payload, np.uint8), cv2.IMREAD_COLOR)
    return [] if decoded is None else [decoded]


def _xlsx_native(payload: bytes) -> list[dict[str, Any]]:
    from openpyxl import load_workbook

    workbook = load_workbook(io.BytesIO(payload), read_only=True, data_only=False)
    sheets = []
    for sheet in workbook.worksheets:
        cells = []
        for row in sheet.iter_rows():
            for cell in row:
                if cell.value is not None:
                    cells.append({"coordinate": cell.coordinate, "value": str(cell.value)})
        sheets.append({"name": sheet.title, "cells": cells})
    workbook.close()
    return sheets


def _xls_native(payload: bytes) -> list[dict[str, Any]]:
    import xlrd

    workbook = xlrd.open_workbook(file_contents=payload, on_demand=True)
    sheets = []
    for sheet in workbook.sheets():
        cells = []
        for row in range(sheet.nrows):
            for column in range(sheet.ncols):
                value = sheet.cell_value(row, column)
                if value != "":
                    cells.append({"row": row + 1, "column": column + 1, "value": str(value)})
        sheets.append({"name": sheet.name, "cells": cells})
    workbook.release_resources()
    return sheets


def process_embedded_objects(
    source: Path,
    inventory: dict[str, Any],
    engine: Any,
    work_dir: Path,
    soffice: Path,
    zoom: float,
) -> tuple[list[dict[str, Any]], list[dict[str, Any]], list[dict[str, Any]]]:
    results, errors, risks = [], [], []
    if not inventory.get("objects"):
        return results, errors, risks
    source_sha = sha256_file(source)
    with ZipFile(source) as archive:
        for object_index, item in enumerate(inventory["objects"], 1):
            member = item["member"]
            payload = archive.read(member)
            extension = item["extension"]
            record = {**item, "raw_ocr_frames": [], "native_content": []}
            object_root = work_dir / "embedded" / source_sha / f"{object_index:04d}"
            try:
                if not payload:
                    record["status"] = "source_member_empty_unreadable"
                    risks.append({
                        "code": "empty_embedded_member", "source": str(source),
                        "member": member, "source_pages": item["source_pages"],
                        "member_sha256": item["member_sha256"],
                        "message": "源 OOXML 包中的内嵌成员为零字节，无法解码或 OCR；保留空成员字节身份，不伪造内容。",
                    })
                elif extension in RASTER | ANIMATED:
                    frames = _decode_member_frames(payload, extension)
                    if not frames:
                        raise LosslessOCRError(f"无法解码嵌入图片：{member}")
                    for frame_index, frame in enumerate(frames, 1):
                        record["raw_ocr_frames"].append(evidence_record(
                            engine, frame, object_root / f"frame-{frame_index:04d}.png",
                            page=item["source_pages"][0] if len(item["source_pages"]) == 1 else None,
                            part=frame_index, part_count=len(frames),
                            render_sha256=item["member_sha256"], render_mode="embedded_raster_frame",
                        ))
                    record["status"] = "ocr_complete"
                elif extension in VECTOR:
                    record["status"] = "covered_by_full_slide_render"
                    risks.append({
                        "code": "vector_not_separately_rasterized", "source": str(source),
                        "member": member, "source_pages": item["source_pages"],
                        "message": "矢量对象由整页渲染 OCR 覆盖，未独立栅格化。",
                    })
                elif extension in AUDIO_VIDEO:
                    manifest_path = ROOT / ".work/embedded-office/manifest.json"
                    if not manifest_path.is_file():
                        raise LosslessOCRError("嵌入音视频清单不存在。")
                    manifest = json.loads(manifest_path.read_text(encoding="utf-8"))
                    matches = [entry for entry in manifest.get("files", [])
                               if entry.get("sha256") == item["member_sha256"]]
                    if len(matches) != 1:
                        raise LosslessOCRError(f"嵌入音视频无法唯一对应既有证据：{member}")
                    extracted = Path(matches[0]["output_path"])
                    media_work = ROOT / ".work/single-video" / item["member_sha256"][:16]
                    if (not extracted.is_file()
                            or sha256_file(extracted) != item["member_sha256"]
                            or not (media_work / "asr.json").is_file()
                            or not (media_work / "ocr.json").is_file()):
                        raise LosslessOCRError(f"嵌入音视频 ASR/OCR 证据不完整：{member}")
                    record["status"] = "delegated_evidence_verified"
                    record["evidence_reference"] = {
                        "manifest": str(manifest_path),
                        "extracted_path": str(extracted),
                        "media_work_dir": str(media_work),
                    }
                elif extension == ".xlsx":
                    # Native extraction is the authoritative cell-level path and
                    # remains a hard requirement.  LibreOffice rendering is only
                    # supplemental: the workbook's visible chart/object surface
                    # is already present in the source page's full-page OCR.
                    record["native_content"] = _xlsx_native(payload)
                    attachment = object_root / (Path(member).stem + ".xlsx")
                    atomic_bytes(attachment, payload)
                    try:
                        rendered = run_soffice(attachment, object_root / "rendered", soffice)
                    except Exception as exc:
                        record["status"] = XLSX_RENDER_DEGRADED_STATUS
                        record["render_error"] = {
                            "exception_type": type(exc).__name__,
                            "message": str(exc),
                        }
                        risks.append({
                            "code": XLSX_RENDER_DEGRADED_RISK,
                            "source": str(source),
                            "member": member,
                            "source_pages": item["source_pages"],
                            "exception_type": type(exc).__name__,
                            "exception_message": str(exc),
                            "coverage": "source_full_page_ocr",
                            "message": (
                                "XLSX 原生单元格已完整提取；LibreOffice 独立渲染失败。"
                                "工作簿在来源页中的可见面由整页 OCR 覆盖。"
                            ),
                        })
                    else:
                        render_sha = sha256_file(rendered)
                        for page, part, part_count, image in pdf_page_images(rendered, zoom):
                            record["raw_ocr_frames"].append(evidence_record(
                                engine, image,
                                object_root / "evidence" / f"page-{page:04d}-{part:02d}.png",
                                page=page, part=part, part_count=part_count,
                                render_sha256=render_sha, render_mode="libreoffice_embedded_xlsx",
                            ))
                        record["render_sha256"] = render_sha
                        record["status"] = "native_and_ocr_complete"
                elif extension == ".bin":
                    try:
                        record["native_content"] = _xls_native(payload)
                    except Exception as exc:
                        record["status"] = "ole_preview_only"
                        risks.append({
                            "code": "unsupported_non_excel_ole", "source": str(source),
                            "member": member, "source_pages": item["source_pages"],
                            "message": f"OLE 不是可解析的 Excel；其预览仅由整页 OCR 覆盖：{exc}",
                        })
                    else:
                        attachment = object_root / "embedded.xls"
                        atomic_bytes(attachment, payload)
                        rendered = run_soffice(attachment, object_root / "rendered", soffice)
                        render_sha = sha256_file(rendered)
                        for page, part, part_count, image in pdf_page_images(rendered, zoom):
                            record["raw_ocr_frames"].append(evidence_record(
                                engine, image, object_root / "evidence" / f"page-{page:04d}-{part:02d}.png",
                                page=page, part=part, part_count=part_count,
                                render_sha256=render_sha, render_mode="libreoffice_embedded_xls_ole",
                            ))
                        record["render_sha256"] = render_sha
                        record["status"] = "native_and_ocr_complete"
                else:
                    record["status"] = "unsupported_member"
                    risks.append({
                        "code": "unsupported_embedded_member", "source": str(source),
                        "member": member, "source_pages": item["source_pages"],
                        "message": f"未支持的内嵌成员类型：{extension}",
                    })
                if not item["source_pages"]:
                    is_docx = source.suffix.lower() == ".docx"
                    risks.append({
                        "code": ("docx_object_page_unmapped" if is_docx
                                 else "orphan_embedded_member"),
                        "source": str(source),
                        "member": member,
                        "message": (
                            "DOCX 不提供稳定物理分页；对象保留部件定位，但不能伪造来源页。"
                            if is_docx else
                            "PPT 包含未被幻灯片、版式或母版引用的孤立成员；内容已处理但无页码。"
                        ),
                    })
            except Exception as exc:
                record["status"] = "error"
                errors.append({
                    "code": "embedded_object_processing_failed", "source": str(source),
                    "member": member, "source_pages": item["source_pages"],
                    "message": str(exc),
                })
            results.append(record)
    return results, errors, risks


def _embedded_inventory(source: Path) -> dict[str, Any]:
    extension = source.suffix.lower()
    if extension == ".pptx":
        return pptx_inventory(source)
    if extension == ".docx":
        return docx_inventory(source)
    return {"objects": []}


def embedded_inventory_sha256(objects: list[dict[str, Any]]) -> str:
    """Bind a cache to the exact package-member inventory and page locators."""
    canonical = [{
        "member": item["member"],
        "member_sha256": item["member_sha256"],
        "size_bytes": item["size_bytes"],
        "extension": item["extension"],
        "source_pages": item["source_pages"],
        "relationship_types": item["relationship_types"],
    } for item in sorted(objects, key=lambda value: value["member"])]
    return sha256_bytes(json.dumps(
        canonical, ensure_ascii=False, sort_keys=True, separators=(",", ":"),
    ).encode("utf-8"))


def refresh_cached_embedded_inventory(
    record: dict[str, Any], source: Path, engine: Any, work_dir: Path,
    soffice: Path, zoom: float,
) -> tuple[dict[str, Any], int]:
    """Upgrade stale caches without re-running already verified full-page OCR."""
    inventory = _embedded_inventory(source)
    expected = {item["member"]: item for item in inventory["objects"]}
    actual_values = record.get("embedded_objects")
    if not isinstance(actual_values, list):
        raise LosslessOCRError(f"缓存 embedded_objects 非数组：{source}")
    actual: dict[str, dict[str, Any]] = {}
    for item in actual_values:
        member = item.get("member") if isinstance(item, dict) else None
        if not isinstance(member, str) or member in actual:
            raise LosslessOCRError(f"缓存内嵌成员缺名或重复：{source}: {member}")
        actual[member] = item
    extras = sorted(set(actual) - set(expected))
    if extras:
        raise LosslessOCRError(f"缓存含源包不存在的内嵌成员：{source}: {extras}")

    refresh_names = set(expected) - set(actual)
    # WDP/JPEG-XR was explicitly unsupported in the older cache.  Once the
    # byte-exact JPEG-XR decoder is available it must be independently OCRed;
    # an old risk record is not accepted as permanent coverage.
    refresh_names.update(
        member for member, item in actual.items()
        if expected[member]["extension"] == ".wdp"
        and item.get("status") == "unsupported_member"
    )
    missing_items = [expected[name] for name in sorted(refresh_names)]
    new_risks: list[dict[str, Any]] = []
    if missing_items:
        # A separate namespace prevents subset indices from overwriting evidence
        # paths generated during the original full-file pass.
        added, errors, new_risks = process_embedded_objects(
            source, {"objects": missing_items}, engine,
            work_dir / "inventory-refresh-v2", soffice, zoom,
        )
        if errors:
            raise LosslessOCRError(
                f"缓存缺失内嵌成员补处理失败：{source}: "
                + json.dumps(errors, ensure_ascii=False)
            )
        actual.update({item["member"]: item for item in added})

    # Preserve OCR/native evidence and status, but bind provenance metadata to
    # the freshly derived source-package facts.
    provenance_keys = {
        "member", "member_sha256", "size_bytes", "extension",
        "source_pages", "relationship_types",
    }
    reconciled = []
    for member, package_item in sorted(expected.items()):
        cached_item = actual[member]
        if (cached_item.get("member_sha256") != package_item["member_sha256"]
                or cached_item.get("size_bytes") != package_item["size_bytes"]
                or cached_item.get("extension") != package_item["extension"]):
            raise LosslessOCRError(f"缓存内嵌成员字节身份不匹配：{source}: {member}")
        reconciled.append({
            **{key: value for key, value in cached_item.items() if key not in provenance_keys},
            **package_item,
        })

    risks = record.get("risks")
    if not isinstance(risks, list):
        raise LosslessOCRError(f"缓存 risks 非数组：{source}")
    mapping_codes = {
        "orphan_embedded_member", "docx_object_page_unmapped",
        # Pre-contract caches used this ambiguous name for DOCX media.  Remove
        # it during reconciliation so one object has one explicit locator risk.
        "embedded_object_page_unmapped",
    }
    refreshed_risks = [
        dict(risk) for risk in risks
        if not (isinstance(risk, dict) and risk.get("code") in mapping_codes)
        and not (
            isinstance(risk, dict)
            and risk.get("code") == "unsupported_embedded_member"
            and risk.get("member") in refresh_names
        )
    ]
    refreshed_risks.extend(new_risks)
    for risk in refreshed_risks:
        member = risk.get("member") if isinstance(risk, dict) else None
        if member in expected and "source_pages" in risk:
            risk["source_pages"] = expected[member]["source_pages"]
    for member, item in sorted(expected.items()):
        if item["source_pages"]:
            continue
        is_docx = source.suffix.lower() == ".docx"
        refreshed_risks.append({
            "code": "docx_object_page_unmapped" if is_docx else "orphan_embedded_member",
            "source": str(source.resolve()), "member": member,
            "message": (
                "DOCX 不提供稳定物理分页；对象保留部件定位，但不能伪造来源页。"
                if is_docx else
                "PPT 包含未被幻灯片、版式或母版引用的孤立成员；内容已处理但无页码。"
            ),
        })

    record = dict(record)
    record["embedded_objects"] = reconciled
    record["risks"] = refreshed_risks
    record["embedded_inventory_version"] = EMBEDDED_INVENTORY_VERSION
    record["embedded_inventory_sha256"] = embedded_inventory_sha256(inventory["objects"])
    return record, len(missing_items)


def external_risks(source: Path, values: list[dict[str, Any]]) -> list[dict[str, Any]]:
    output = []
    for item in values:
        if not item.get("file_like"):
            continue
        output.append({
            "code": "missing_external_asset",
            "source": str(source),
            "source_page": item.get("slide"),
            "relationship_part": item.get("relationship_part") or item.get("source_part"),
            "relationship_id": item.get("relationship_id"),
            "relationship_type": item.get("relationship_type"),
            "target": item.get("target"),
            "message": "外链文件字节不在 OOXML 包或 data 数据集中，无法 OCR。",
        })
    return output


def process_file(
    source: Path,
    engine: Any,
    work_dir: Path,
    render_roots: list[Path],
    soffice: Path,
    zoom: float,
    progress: Callable[[str], None] = print,
) -> dict[str, Any]:
    started = time.perf_counter()
    source = source.resolve()
    source_sha = sha256_file(source)
    extension = source.suffix.lower()
    file_root = work_dir / "files" / source_sha
    errors: list[dict[str, Any]] = []
    risks: list[dict[str, Any]] = []
    pages: list[dict[str, Any]] = []
    objects: list[dict[str, Any]] = []
    external: list[dict[str, Any]] = []
    render_info: dict[str, Any] | None = None
    native_document: list[dict[str, Any]] = []

    if extension == ".pptx":
        inventory = pptx_inventory(source)
        external = inventory["external_relationships"]
        expected_pages = len(inventory["slide_parts"])
        render = find_cached_render(source, render_roots)
        reused = render is not None
        if render is None:
            render = run_soffice(source, file_root / "rendered", soffice)
        render_sha = sha256_file(render)
        with pymupdf.open(render) as document:
            rendered_pages = len(document)
        plan = validate_ppt_render_plan(source.name, expected_pages, rendered_pages)
        render_info = {
            "path": str(render), "sha256": render_sha, "reused": reused,
            "page_count": rendered_pages, "source_slide_count": expected_pages,
        }
        rendered_images = dict(pdf_full_page_images(render, zoom))
        for page_number, mode in enumerate(plan, 1):
            native = inventory["native_by_page"].get(page_number, [])
            if mode == "rendered":
                image = rendered_images.get(page_number)
                if image is None:
                    raise LosslessOCRError(f"PPT 第 {page_number} 页没有唯一整页渲染图。")
                record = evidence_record(
                    engine, image, file_root / "evidence" / f"page-{page_number:04d}.png",
                    page=page_number, part=1, part_count=1,
                    render_sha256=render_sha, render_mode="libreoffice_pptx_page",
                )
            else:
                image = native_text_fallback(native)
                record = evidence_record(
                    engine, image, file_root / "evidence" / f"page-{page_number:04d}-fallback.png",
                    page=page_number, part=1, part_count=1,
                    render_sha256=None, render_mode="ooxml_native_text_fallback",
                )
                risks.append({
                    "code": "ppt_page_render_fallback", "source": str(source),
                    "source_page": page_number,
                    "message": "LibreOffice 漏页；本页以 OOXML 原生文字单独生成像素证据并 OCR。",
                })
            record["native_lines"] = native
            pages.append(record)
            if page_number == expected_pages or page_number % 10 == 0:
                progress(f"    页面 {page_number}/{expected_pages}")
        objects, object_errors, object_risks = process_embedded_objects(
            source, inventory, engine, work_dir, soffice, zoom,
        )
        errors.extend(object_errors)
        risks.extend(object_risks)
        risks.extend(external_risks(source, external))
        if len(pages) != expected_pages:
            raise LosslessOCRError(f"PPT 页面覆盖失败：{len(pages)}/{expected_pages}")
    elif extension == ".docx":
        inventory = docx_inventory(source)
        native_document = inventory["native_document"]
        external = inventory["external_relationships"]
        render = find_cached_render(source, render_roots)
        reused = render is not None
        if render is None:
            render = run_soffice(source, file_root / "rendered", soffice)
        render_sha = sha256_file(render)
        with pymupdf.open(render) as document:
            rendered_pages = len(document)
        render_info = {"path": str(render), "sha256": render_sha, "reused": reused,
                       "page_count": rendered_pages, "source_slide_count": None}
        for page, image in pdf_full_page_images(render, zoom):
            pages.append(evidence_record(
                engine, image, file_root / "evidence" / f"page-{page:04d}-01.png",
                page=page, part=1, part_count=1,
                render_sha256=render_sha, render_mode="libreoffice_docx_page",
            ))
        objects, object_errors, object_risks = process_embedded_objects(
            source, inventory, engine, work_dir, soffice, zoom,
        )
        errors.extend(object_errors)
        risks.extend(object_risks)
        risks.extend(external_risks(source, external))
    elif extension == ".pdf":
        native = pdf_native_pages(source)
        render_sha = source_sha
        with pymupdf.open(source) as document:
            physical_pages = len(document)
        render_info = {"path": str(source), "sha256": render_sha, "reused": True,
                       "page_count": physical_pages, "source_slide_count": None}
        for page, part, part_count, image in pdf_page_images(source, zoom):
            record = evidence_record(
                engine, image, file_root / "evidence" / f"page-{page:04d}-{part:02d}.png",
                page=page, part=part, part_count=part_count,
                render_sha256=render_sha, render_mode="pdf_page",
            )
            record["native_lines"] = native.get(page, [])
            pages.append(record)
    elif extension in {".png", ".jpg", ".jpeg"}:
        render_info = {"path": str(source), "sha256": source_sha, "reused": True,
                       "page_count": 1, "source_slide_count": None}
        for page, part, part_count, image in image_slices(source):
            pages.append(evidence_record(
                engine, image, file_root / "evidence" / f"page-{page:04d}-{part:02d}.png",
                page=page, part=part, part_count=part_count,
                render_sha256=source_sha, render_mode="source_image",
            ))
    else:
        raise LosslessOCRError(f"不支持的课件类型：{source}")

    inventory_objects = inventory.get("objects", []) if extension in {".pptx", ".docx"} else []
    return {
        "source": source.name,
        "source_path": str(source),
        "source_name": source.name,
        "source_sha256": source_sha,
        "size_bytes": source.stat().st_size,
        "file_type": extension.lstrip(".").upper(),
        "render": render_info,
        "page_record_count": len(pages),
        "pages": pages,
        "native_document_text": native_document,
        "embedded_objects": objects,
        "embedded_inventory_version": EMBEDDED_INVENTORY_VERSION,
        "embedded_inventory_sha256": embedded_inventory_sha256(inventory_objects),
        "external_relationships": external,
        "errors": errors,
        "risks": risks,
        "processing_complete": not errors,
        "processing_seconds": round(time.perf_counter() - started, 3),
        "render_zoom": zoom,
        "processor_version": PROCESSOR_VERSION,
    }


def collect_sources(roots: Iterable[Path]) -> list[Path]:
    output: dict[Path, Path] = {}
    for root in roots:
        root = root.resolve(strict=True)
        for path in root.rglob("*"):
            if path.is_file() and path.suffix.lower() in SUPPORTED:
                output[path.resolve()] = path.resolve()
    return sorted(output.values(), key=lambda path: (path.name.casefold(), str(path)))


def cache_valid(record: dict[str, Any], source: Path, engine_version: str, zoom: float) -> bool:
    if (record.get("processor_version") != PROCESSOR_VERSION
            or record.get("source_sha256") != sha256_file(source)
            or record.get("engine_version") != engine_version
            or record.get("render_zoom") != zoom
            or not record.get("processing_complete")):
        return False
    try:
        evidence = [page for page in record.get("pages", [])]
        for obj in record.get("embedded_objects", []):
            evidence.extend(obj.get("raw_ocr_frames", []))
        for page in evidence:
            path = Path(page["evidence_path"])
            if not path.is_file() or sha256_file(path) != page["image_sha256"]:
                return False
    except (KeyError, OSError, TypeError):
        return False
    return True


def parse_args(argv: list[str] | None = None) -> argparse.Namespace:
    parser = argparse.ArgumentParser(description="课件逐页、逐对象无损 OCR")
    parser.add_argument("--source-dir", type=Path, action="append")
    parser.add_argument(
        "--work-dir", type=Path,
        default=ROOT / ".work/first-principles-2026/courseware",
    )
    parser.add_argument(
        "--output", type=Path,
        default=ROOT / ".work/first-principles-2026/courseware_ocr_lossless.json",
    )
    parser.add_argument("--render-cache-dir", type=Path, action="append")
    parser.add_argument("--soffice", type=Path, default=ROOT / "scripts/soffice_local.sh")
    parser.add_argument("--zoom", type=float, default=2.2)
    parser.add_argument("--expected-files", type=int, default=29)
    parser.add_argument("--limit", type=int, default=0, help="仅处理前 N 个，供验证；0 为全部")
    parser.add_argument("--only", help="仅处理文件名包含此字符串的课件")
    parser.add_argument("--no-cache", action="store_true")
    return parser.parse_args(argv)


def main(argv: list[str] | None = None) -> int:
    args = parse_args(argv)
    source_roots = args.source_dir or [
        ROOT / "data/2025",
        ROOT / "data/2026",
        ROOT / ".work/extracted-2026/项目管理1-9",
    ]
    render_roots = args.render_cache_dir or [
        ROOT / ".work/courseware-2025-v5/rendered",
        ROOT / ".work/courseware-2026-v5/rendered",
    ]
    sources = collect_sources(source_roots)
    if len(sources) != args.expected_files:
        raise SystemExit(f"课件清单数量不符：{len(sources)}/{args.expected_files}")
    selected = [path for path in sources if not args.only or args.only in path.name]
    if args.limit:
        selected = selected[:args.limit]

    from rapidocr import EngineType, ModelType, OCRVersion, RapidOCR

    engine_version = importlib.metadata.version("rapidocr")
    engine_config = {
        "Global.text_score": 0.0,
        "Global.log_level": "warning",
        "Global.max_side_len": 4096,
        "Global.use_det": True,
        "Global.use_cls": True,
        "Global.use_rec": True,
        "Det.limit_side_len": 736,
        "Det.limit_type": "min",
        "Det.engine_type": EngineType.ONNXRUNTIME,
        "Det.ocr_version": OCRVersion.PPOCRV6,
        "Det.model_type": ModelType.SMALL,
        "Rec.engine_type": EngineType.ONNXRUNTIME,
        "Rec.ocr_version": OCRVersion.PPOCRV6,
        "Rec.model_type": ModelType.SMALL,
        "EngineConfig.onnxruntime.intra_op_num_threads": 2,
        "EngineConfig.onnxruntime.inter_op_num_threads": 1,
        "EngineConfig.onnxruntime.use_cuda": False,
    }
    engine = RapidOCR(params=engine_config)
    args.work_dir.mkdir(parents=True, exist_ok=True)
    cache_dir = args.work_dir / "cache"
    cache_dir.mkdir(parents=True, exist_ok=True)
    records = []
    run_errors = []
    print(f"课件无损 OCR：选择 {len(selected)}/{len(sources)} 个文件", flush=True)
    for index, source in enumerate(selected, 1):
        source_sha = sha256_file(source)
        cache_path = cache_dir / f"{source_sha}.json"
        print(f"[{index}/{len(selected)}] {source.name}", flush=True)
        if cache_path.is_file() and not args.no_cache:
            try:
                cached = json.loads(cache_path.read_text(encoding="utf-8"))
            except (OSError, ValueError):
                cached = None
            if cached and cache_valid(cached, source, engine_version, args.zoom):
                try:
                    cached, added_count = refresh_cached_embedded_inventory(
                        cached, source, engine, args.work_dir,
                        args.soffice.resolve(), args.zoom,
                    )
                except Exception as exc:
                    print(f"    缓存内嵌清单升级失败，转为完整重跑：{exc}", flush=True)
                else:
                    cached["engine_version"] = engine_version
                    atomic_json(cache_path, cached)
                    records.append(cached)
                    if added_count:
                        print(f"    复用页面 OCR；补处理内嵌成员 {added_count} 个", flush=True)
                    else:
                        print("    已验证并复用无损缓存", flush=True)
                    continue
        try:
            record = process_file(
                source, engine, args.work_dir, [Path(root) for root in render_roots],
                args.soffice.resolve(), args.zoom,
            )
            record["engine_version"] = engine_version
            atomic_json(cache_path, record)
            records.append(record)
        except Exception as exc:
            error = {"code": "file_processing_failed", "source": str(source), "message": str(exc)}
            run_errors.append(error)
            records.append({
                "source": source.name, "source_path": str(source), "source_name": source.name,
                "source_sha256": source_sha, "processing_complete": False,
                "errors": [error], "risks": [], "processor_version": PROCESSOR_VERSION,
                "engine_version": engine_version,
            })
            print(f"    失败：{exc}", flush=True)
        partial = {
            "schema_version": SCHEMA_VERSION,
            "processor_version": PROCESSOR_VERSION,
            "generated_at": datetime.now().astimezone().isoformat(timespec="seconds"),
            "engine": {"name": "RapidOCR", "version": engine_version,
                       "det_model": "PP-OCRv6/ch/small/onnxruntime",
                       "rec_model": "PP-OCRv6/ch/small/onnxruntime",
                       "cls_model": "PP-OCRv4/ch/mobile/onnxruntime",
                       "text_score": 0.0, "raw_output_preserved": True},
            "expected_file_count": args.expected_files,
            "selected_file_count": len(selected),
            "processed_file_count": len(records),
            "processing_complete": False,
            "files": records,
            "errors": [error for item in records for error in item.get("errors", [])],
            "risks": [risk for item in records for risk in item.get("risks", [])],
        }
        atomic_json(args.output, partial)

    errors = [error for item in records for error in item.get("errors", [])]
    result = {
        "schema_version": SCHEMA_VERSION,
        "processor_version": PROCESSOR_VERSION,
        "generated_at": datetime.now().astimezone().isoformat(timespec="seconds"),
        "engine": {"name": "RapidOCR", "version": engine_version,
                   "det_model": "PP-OCRv6/ch/small/onnxruntime",
                   "rec_model": "PP-OCRv6/ch/small/onnxruntime",
                   "cls_model": "PP-OCRv4/ch/mobile/onnxruntime",
                   "text_score": 0.0, "raw_output_preserved": True},
        "source_roots": [str(path.resolve()) for path in source_roots],
        "expected_file_count": args.expected_files,
        "selected_file_count": len(selected),
        "processed_file_count": len(records),
        "page_record_count": sum(item.get("page_record_count", 0) for item in records),
        "processing_complete": len(selected) == len(sources) and not errors
                               and all(item.get("processing_complete") for item in records),
        "source_complete": len(selected) == len(sources) and not errors and not any(
            risk.get("code") == "missing_external_asset"
            for item in records for risk in item.get("risks", [])
        ),
        "files": records,
        "errors": errors,
        "risks": [risk for item in records for risk in item.get("risks", [])],
    }
    atomic_json(args.output, result)
    print(
        f"完成 {len(records)}/{len(selected)}；页面/分片 {result['page_record_count']}；"
        f"错误 {len(errors)}；风险 {len(result['risks'])}；输出 {args.output}",
        flush=True,
    )
    return 0 if not errors else 2


if __name__ == "__main__":
    sys.exit(main())
